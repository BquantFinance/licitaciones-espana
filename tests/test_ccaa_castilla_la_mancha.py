"""Tests offline de scripts/ccaa_castilla_la_mancha.py.

El portal de contratación de la Junta (páginas de ficheros de transparencia y
sus ficheros) y el perfil de contratante de la UCLM (formulario ASP.NET con
__VIEWSTATE y un postback por ejercicio) se simulan con un ``requests.get`` y
un ``requests.post`` falsos; los XLSX, XLS, ZIP y RAR se generan en el test.
"""

import importlib.util
import io
import json
import runpy
import struct
import sys
import time
import zipfile
import zlib
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
import pytest
import requests

REPO_ROOT = Path(__file__).resolve().parents[1]
SCRIPT = REPO_ROOT / "scripts" / "ccaa_castilla_la_mancha.py"
SLEEP_REAL = time.sleep


def _cargar():
    spec = importlib.util.spec_from_file_location("ccaa_castilla_la_mancha", SCRIPT)
    modulo = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(modulo)
    return modulo


M = _cargar()
ANIO = datetime.now().year          # el script decide qué refrescar con el año real
VIEJO = ANIO - 3                    # un año cerrado (solo se vuelve a pedir con --comprobar-todo)
FICHEROS = "https://contratacion.castillalamancha.es/sites/default/files/"
CONCEPTOS = {"113": "Contratación menor: JCCM y SESCAM",
             "114": "Contratos formalizados y sobre aquellos más importantes (de valor igual o superior a 500.000€)",
             "117": "Datos estadísticos"}
HAY_LECTOR_RAR = M._modulo("libarchive") is not None or M._modulo("rarfile") is not None


# ---------------------------------------------------------------------------
# Ficheros de prueba
# ---------------------------------------------------------------------------

def _xlsx(hojas):
    libro = openpyxl.Workbook()
    libro.remove(libro.active)
    for nombre, filas in hojas.items():
        hoja = libro.create_sheet(nombre)
        for fila in filas:
            hoja.append(fila)
    datos = io.BytesIO()
    libro.save(datos)
    return datos.getvalue()


def _zip(miembros, utf8=True):
    """ZIP en memoria. utf8=False: como uno hecho en Windows, con los nombres en
    cp850 y sin la marca UTF-8 (zipfile siempre escribe UTF-8: se escribe un
    nombre ASCII del mismo largo y después se cambian sus bytes)."""
    datos, cambios = io.BytesIO(), {}
    with zipfile.ZipFile(datos, "w", zipfile.ZIP_DEFLATED) as archivo:
        for nombre, contenido in miembros.items():
            if not utf8:
                crudo = nombre.encode("cp850")
                provisional = bytes(b if b < 0x80 else 0x23 for b in crudo)
                cambios[provisional] = crudo
                nombre = provisional.decode("ascii")
            archivo.writestr(zipfile.ZipInfo(nombre, date_time=(2024, 5, 20, 7, 19, 28)), contenido,
                             zipfile.ZIP_DEFLATED)
    resultado = datos.getvalue()
    for provisional, crudo in sorted(cambios.items(), key=lambda c: -len(c[0])):
        resultado = resultado.replace(provisional, crudo)
    return resultado


def _rar4(nombre, contenido):
    """RAR 4 con un fichero sin comprimir (método 'store'), como los del SESCAM
    pero sin necesitar el programa rar para crearlo."""
    def bloque(tipo, flags, cuerpo):
        sin_crc = struct.pack("<BHH", tipo, flags, 7 + len(cuerpo)) + cuerpo
        return struct.pack("<H", zlib.crc32(sin_crc) & 0xFFFF) + sin_crc

    n = nombre.encode("ascii")
    fecha = ((2024 - 1980) << 25) | (5 << 21) | (20 << 16)
    cuerpo = struct.pack("<IIBIIBBHI", len(contenido), len(contenido), 2, zlib.crc32(contenido) & 0xFFFFFFFF,
                         fecha, 20, 0x30, len(n), 0x20) + n
    return (b"Rar!\x1a\x07\x00" + bloque(0x73, 0, struct.pack("<HI", 0, 0))
            + bloque(0x74, 0x8000, cuerpo) + contenido + bloque(0x7B, 0x4000, b""))


MHTML = ("From: <Saved by Blink>\r\nSnapshot-Content-Location: https://contratacion.castillalamancha.es/ano-2016/a-2\r\n"
         "MIME-Version: 1.0\r\nContent-Type: multipart/related;\r\n\ttype=\"text/html\";\r\n\tboundary=\"----x\"\r\n\r\n"
         "------x\r\nContent-Type: text/html\r\n\r\n<html><body><a href=\"a.xls\">a</a></body></html>\r\n"
         "------x--\r\n").encode("utf-8")

SPREADSHEETML = """<?xml version="1.0"?>
<?mso-application progid="Excel.Sheet"?>
<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet" xmlns:o="urn:schemas-microsoft-com:office:office"
 xmlns:x="urn:schemas-microsoft-com:office:excel" xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet"
 xmlns:html="http://www.w3.org/TR/REC-html40">
 <Worksheet ss:Name="Contratos Suministros-Servicios">
  <Table>
   <Row><Cell><Data ss:Type="String">Gerencia</Data></Cell><Cell><Data ss:Type="String">Artículo</Data></Cell>
    <Cell><Data ss:Type="String">Importe</Data></Cell><Cell><Data ss:Type="String">Fecha</Data></Cell></Row>
   <Row><Cell><Data ss:Type="String">61037000 GUETS</Data></Cell><Cell><Data ss:Type="String">LIMPIEZA</Data></Cell>
    <Cell><Data ss:Type="Number">28.880800000000001</Data></Cell>
    <Cell><Data ss:Type="DateTime">2016-07-06T00:00:00.000</Data></Cell></Row>
   <Row ss:Index="4"><Cell ss:Index="2"><ss:Data ss:Type="String" xmlns="http://www.w3.org/TR/REC-html40">GASAS <B>estériles</B></ss:Data></Cell>
    <Cell><Data ss:Type="Number">7</Data></Cell></Row>
  </Table>
 </Worksheet>
 <Worksheet ss:Name="Contratos Farmacia">
  <Table>
   <Row><Cell ss:MergeAcross="1"><Data ss:Type="String">Gerencia</Data></Cell>
    <Cell><Data ss:Type="String">Fecha</Data></Cell><Cell ss:Index="5"><Data ss:Type="String">Artículo</Data></Cell></Row>
   <Row><Cell><Data ss:Type="String">61035000 GAE Toledo</Data></Cell><Cell><Data ss:Type="String">sobra</Data></Cell>
    <Cell><Data ss:Type="DateTime">2016-07-06T10:30:00.000</Data></Cell><Cell ss:Index="5"><Data ss:Type="String">U-005</Data></Cell></Row>
  </Table>
 </Worksheet>
</Workbook>
""".encode("utf-8")

CABECERA_JUNTA = ["Nº expediente PICOS", "Organismo", "Descripción", "Importe adjudicación IVA incluido",
                  "CIF Adjudicatario(s)", "Adjudicatario(s)", "Fecha de adjudicación"]


def _junta(*expedientes):
    return _xlsx({"Menores": [CABECERA_JUNTA] + [
        [e, "Consejería de Fomento", f"Objeto {e}", 1210.5, "B13198536", "DOBLE C SERIGRAFIA S.L.",
         datetime(2023, 2, 15)] for e in expedientes]})


def _sescam_xlsx(tipo, articulo):
    return _xlsx({"Hoja1": [["Gerencia", "Tipo compra", "Artículo", "Proveedor", "Nº de factura",
                             f"Importe con IVA 1 TRIM {tipo}"],
                            ["61031300 GAI Villarrobledo", "SUMINISTRO MENOR", articulo, "JANSSEN-CILAG S.A.",
                             "0941044807", 15523.93]]})


# ---------------------------------------------------------------------------
# Portal simulado
# ---------------------------------------------------------------------------

class FakeResponse:
    def __init__(self, status=200, body=b"", headers=None):
        self.status_code = status
        self._body = body
        self.headers = headers or {}

    @property
    def text(self):
        return self._body.decode("utf-8", "replace")

    def iter_content(self, chunk_size=8192):
        yield self._body

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False


def pagina_jccm(enlaces, conceptos=CONCEPTOS, anios=None):
    """Como la página real (Drupal): filtros concepto/año y un views-row por fichero."""
    anios = anios if anios is not None else range(2015, ANIO + 1)
    opciones_c = '<option value="All">- Todos -</option>' + "".join(
        f'<option value="{c}">{n}</option>' for c, n in conceptos.items())
    opciones_a = '<option value="All">- Todos -</option>' + "".join(f'<option value="{a}">{a}</option>' for a in anios)
    filas = "".join('<div class="views-row">\n<span class="views-field views-field-title"><span class="field-content">'
                    f'<a href="{href}">{texto}</a></span></span>\n</div>' for href, texto in enlaces)
    return ('<!DOCTYPE html><html lang="es"><head><title>Ficheros de datos de transparencia</title></head><body>'
            '<a href="/servicios">Servicios</a>'
            '<img src="https://www.castillalamancha.es/sites/default/files/logonuevoazul_1_0.png" />'
            '<a href="https://www.castillalamancha.es/sites/default/files/folleto.pdf">Otro portal</a>'
            '<form action="/ficheros-transparencia" method="get">'
            f'<select data-drupal-selector="edit-concepto" name="concepto" class="form-select">{opciones_c}</select>'
            f'<select data-drupal-selector="edit-year" id="edit-year" name="year" class="form-select">{opciones_a}'
            '</select></form><div class="agrupador" role="heading">Contratación menor: JCCM y SESCAM</div>'
            f'{filas}</body></html>').encode("utf-8")


CABECERA_UCLM = ["Nif", "Proveedor", "Expediente", "Objeto", "Duraci&#243;n", "Adjudicado"]


def tabla_uclm(filas, cabecera=CABECERA_UCLM):
    th = "".join(f'<th class="t-auto text-center" scope="col">{c}</th>' for c in cabecera)
    trs = "".join("<tr>\n\t\t\t\t\t" + "".join(f'<td class="t-auto">{v}</td>' for v in f) + "\n\t\t\t\t</tr>"
                  for f in filas)
    return ('<table class="table table-striped gv_expediente" cellspacing="0" rules="all" border="1" '
            f'id="cph_Contenidos_gv_expediente" style="border-collapse:collapse;">\n<thead>\n<tr class="thead-light">'
            f'{th}</tr>\n</thead><tbody>{trs}</tbody>\n</table>')


def pagina_uclm(anios, viewstate, validacion, fecha, seleccionado=None, tabla=""):
    """Como la página real (ASP.NET): campos ocultos, desplegable de ejercicios,
    fecha del día en la cabecera y la tabla de resultados."""
    opciones = "".join(('<option selected="selected" ' if a == seleccionado else "<option ") + f'value="{a}">{a}</option>'
                       for a in sorted(anios, reverse=True))
    ocultos = "".join(f'<input type="hidden" name="{n}" id="{n}" value="{v}" />' for n, v in [
        ("__EVENTTARGET", ""), ("__EVENTARGUMENT", ""), ("__VIEWSTATE", viewstate),
        ("__VIEWSTATEGENERATOR", "0B7177FB"), ("__VIEWSTATEENCRYPTED", ""), ("__EVENTVALIDATION", validacion)])
    return ('<!DOCTYPE html>\n<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Perfil</title></head><body>'
            f'<span class="uclm_texto_cabecera">{fecha}</span>'
            '<form method="post" action="./contratosMenoresAnteriores.aspx" id="form1">'
            f'<div class="aspNetHidden">{ocultos}</div>'
            f'<select name="{M.CAMPO_EJERCICIO_UCLM}" id="cph_Contenidos_ddl_contrato_anterior">{opciones}</select>'
            '<a id="cph_Contenidos_lbtn_buscar" href="javascript:__doPostBack(&#39;ctl00$cph_Contenidos$lbtn_buscar'
            '&#39;,&#39;&#39;)"><span>Buscar</span></a>'
            f'<div class="table-responsive-sm"><div>{tabla}</div></div></form></body></html>').encode("utf-8")


class FakePortal:
    """urls: url -> bytes | código HTTP | lista de respuestas (una por petición);
    paginas: (concepto, año) -> [(href, texto)] | código; uclm: año -> filas."""

    def __init__(self):
        self.urls = {}
        self.paginas = {}
        self.portada = None
        self.conceptos = dict(CONCEPTOS)
        self.uclm = {}
        self.uclm_actuales = []
        self.uclm_get = 200
        self.llamadas = []
        self.posts = []
        self.peticiones = 0
        self.viewstate_get = None

    def _uclm(self, anios, seleccionado=None, tabla=""):
        # Cada respuesta con otro __VIEWSTATE, otra __EVENTVALIDATION y otra fecha
        self.peticiones += 1
        return pagina_uclm(anios, f"VS{self.peticiones}+/=", f"EV{self.peticiones}", f"día {self.peticiones}",
                           seleccionado, tabla)

    def get(self, url, params=None, headers=None, timeout=None, stream=False):
        self.llamadas.append((url, dict(params or {})))
        if url == M.URL_FICHEROS:
            if not params:
                if isinstance(self.portada, int):
                    return FakeResponse(status=self.portada)
                return FakeResponse(body=self.portada or pagina_jccm([], self.conceptos))
            enlaces = self.paginas.get((str(params["concepto"]), int(params["year"])), [])
            if isinstance(enlaces, int):
                return FakeResponse(status=enlaces)
            return FakeResponse(body=pagina_jccm(enlaces, self.conceptos))
        if url == M.URL_UCLM_ANTERIORES:
            if self.uclm_get != 200:
                return FakeResponse(status=self.uclm_get)
            cuerpo = self._uclm(self.uclm)
            self.viewstate_get = (f"VS{self.peticiones}+/=", f"EV{self.peticiones}")
            return FakeResponse(body=cuerpo)
        if url == M.URL_UCLM_ACTUALES:
            return FakeResponse(body=self._uclm([], tabla=tabla_uclm(self.uclm_actuales, ["Un.Funcional"]
                                                                     + CABECERA_UCLM)))
        cuerpo = self.urls.get(url, 404)
        if isinstance(cuerpo, list):
            cuerpo = cuerpo.pop(0) if len(cuerpo) > 1 else cuerpo[0]
        if isinstance(cuerpo, int):
            return FakeResponse(status=cuerpo, body=b"<html><body>Not Found</body></html>")
        if isinstance(cuerpo, tuple):                  # (bytes, cabeceras)
            return FakeResponse(body=cuerpo[0], headers=cuerpo[1])
        return FakeResponse(body=cuerpo)

    def post(self, url, data=None, headers=None, timeout=None):
        self.posts.append(dict(data or {}))
        assert url == M.URL_UCLM_ANTERIORES
        # El postback tiene que llevar los campos ocultos del formulario leído
        if (data.get("__VIEWSTATE"), data.get("__EVENTVALIDATION")) != self.viewstate_get:
            return FakeResponse(status=500, body=b"<html><body>Validation of viewstate MAC failed</body></html>")
        anio = int(data[M.CAMPO_EJERCICIO_UCLM])
        if data.get("__EVENTTARGET") != M.BOTON_BUSCAR_UCLM or anio not in self.uclm:
            return FakeResponse(body=self._uclm(self.uclm, anio))           # sin tabla
        return FakeResponse(body=self._uclm(self.uclm, anio, tabla_uclm(self.uclm[anio])))

    def pedidas(self, url, **params):
        return sum(1 for u, p in self.llamadas
                   if u == url and all(str(p.get(k)) == str(v) for k, v in params.items()))


def publicar(portal, concepto, anio, carpeta, nombre, texto, contenido):
    portal.paginas.setdefault((concepto, anio), []).append((f"/sites/default/files/{carpeta}/{nombre}", texto))
    portal.urls[f"{FICHEROS}{carpeta}/{nombre}"] = contenido


def _publicar_todo(portal):
    publicar(portal, "113", 2018, "2022-06", "sector_publico_4o_trimestre_2018_0.zip",
             "SECTOR PÚBLICO REGIONAL 4º trimestre 2018", _zip({
                 "0_seccion_18_contratos_menores_oct-dic_2018.xlsx": _xlsx({"oct-dic 2018": [
                     ["CONSEJERÍA DE EDUCACIÓN, CULTURA Y DEPORTES."], [],
                     ["Consejería", "Objeto del contrato", "Importe (con impuestos)", "Adjudicatario"],
                     ["Consejería de Educación", "MATERIAL REPARACION", 16.82, "COMERCIAL GALAN, S.A."]]}),
                 "sescam_4o_trimestre_2018.zip": _zip({"4º trimestre 2018.xlsx": _sescam_xlsx("2018", "GASAS")},
                                                     utf8=False)}))
    publicar(portal, "113", VIEJO, "2023-06", "CM_PRIMER_TRIMESTRE_JCCM.xlsx",
             f"JCCM 1º trimestre {VIEJO} gestor PICOS", _junta(f"{VIEJO}/000001", f"{VIEJO}/000002"))
    publicar(portal, "113", VIEJO, "2024-02", "Menores_completo_JCCM.xlsx",
             f"JCCM año {VIEJO} COMPLETO gestor PICOS", _junta(f"{VIEJO}/000001", f"{VIEJO}/000002", f"{VIEJO}/9"))
    publicar(portal, "113", VIEJO, "2024-05", "CM_PRIMER_TRIMESTRE_SESCAM.zip", f"SESCAM 1º Trimestre {VIEJO}",
             _zip({"CM_PRIMER_TRIMESTRE_FARMACIA_SESCAM.xlsx": _sescam_xlsx(VIEJO, "DARATUMUMAB"),
                   "CM_PRIMER_TRIMESTRE_SUMINISTROS_SESCAM.xlsx": _xlsx({"Hoja1": [
                       ["Gerencia", "Artículo", "Proveedor", "Importe  sin IVA 1 TRIM"],
                       ["61035000 GAE Toledo", "CATETER", "(0100035984) MERCE V. ELECTROMEDICINA, S", 14000]]})}))
    publicar(portal, "113", VIEJO, "2024-06", "Relacion_menores_organismo.xlsx",
             "Relación de contratos menores del organismo", _junta("R-1"))
    publicar(portal, "113", ANIO, "2026-05", f"Contratos_menores_primer_trimestre_{ANIO}_JCCM.xlsx",
             f"JCCM 1º TRIMESTRE {ANIO} gestor PICOS", _junta(f"{ANIO}/000001"))
    publicar(portal, "113", ANIO, "2026-05", "Contratos_menores_primer_trimestre_caja_pagadora.XLSX",
             f"JCCM 1º TRIMESTRE {ANIO} CAJA PAGADORA", _xlsx({"Menores 1T C. Pagadora": [
                 ["Organismo", "Descripción", "Importe IVA incluído", "NIF Adjudicatario", "Adjudicatario",
                  "Fecha Factura"],
                 ["AGENCIA DE TRANSFORMACIÓN DIGITAL", "GASTOS AGUA REUNIONES", 25.72, "B45047537", "ANJOFEVI, S.L.",
                  datetime(2026, 6, 5)]]}))
    publicar(portal, "117", 2021, "2022-07", "consejo_gobierno_menores_2021.xlsx",
             "Informe de Contratación Administrativa del Sector Público 2021 (contratos menores-fichero de datos)",
             _xlsx({"MODELO 2021": [["Número de contrato (NRC)", "IA"], ["000001/2021", 428.65]]}))
    # En "Datos estadísticos" hay más ficheros que no son de menores: no se bajan
    portal.paginas[("117", 2021)].append(("/sites/default/files/2022-07/INFORME_PYMES_2021.xlsx",
                                          "Información estadística contratos adjudicados a PYMES 2021"))
    portal.uclm = {ANIO - 2: [["80162019G", " CELIA VELASCO LOPEZ", "SE43324000210",
                               "TRADUCCIONES                   ", "30", "130,00 €"]],
                   ANIO - 1: [["B42768705", "ZONA SOLDADURA SL", "SE30125004226",
                               "REPARACIONES, MANTENIMIENTO Y CONSERVACI&#211;N", "1", "3.630,00 €"],
                              ["US000611648780", "ZOOM VIDEO COMMUNICATIONS INC", "SE10125005011",
                               "APLICACIONES", "1", "181,38 €"]]}
    portal.uclm_actuales = [["01110C0021", "05679241N", "FELIX BARINGO SERRANO", "SE34026000292", "ESTUDIOS", "1",
                             "594,11 €"]]


@pytest.fixture
def portal(monkeypatch):
    fake = FakePortal()
    monkeypatch.setattr(requests, "get", fake.get)
    monkeypatch.setattr(requests, "post", fake.post)
    monkeypatch.setattr(time, "sleep", lambda s: None)
    _publicar_todo(fake)
    return fake


def _ejecutar(salida, *args):
    return M.main(["--salida", str(salida), *args])


def _log(salida):
    return (salida / "raw" / "descarga_log.txt").read_text(encoding="utf-8")


def _manifiesto(salida):
    return json.loads((salida / "raw" / "_manifiesto.json").read_text(encoding="utf-8"))


def _parquet(salida, nombre):
    return pd.read_parquet(salida / f"{nombre}.parquet")


def _v(serie):
    return [None if pd.isna(v) else v for v in serie]


# ---------------------------------------------------------------------------
# Portal de la Junta: descubrimiento y clasificación
# ---------------------------------------------------------------------------

def test_descubre_clasifica_y_guarda_los_originales(portal, tmp_path):
    assert _ejecutar(tmp_path) == 0

    raw = tmp_path / "raw"
    esperados = {
        "sector_publico/2018/sector_publico_4o_trimestre_2018_0.zip": ("2022-06", "4T", "4"),
        f"menores_junta/{VIEJO}/CM_PRIMER_TRIMESTRE_JCCM.xlsx": ("2023-06", "1T", "1"),
        f"menores_junta/{VIEJO}/Menores_completo_JCCM.xlsx": ("2024-02", "anual", None),
        f"sescam/{VIEJO}/CM_PRIMER_TRIMESTRE_SESCAM.zip": ("2024-05", "1T", "1"),
        f"otros/{VIEJO}/Relacion_menores_organismo.xlsx": ("2024-06", None, None),
        f"menores_junta/{ANIO}/Contratos_menores_primer_trimestre_{ANIO}_JCCM.xlsx": ("2026-05", "1T", "1"),
        f"caja_pagadora/{ANIO}/Contratos_menores_primer_trimestre_caja_pagadora.XLSX": ("2026-05", "1T", "1"),
        "informe_menores/2021/consejo_gobierno_menores_2021.xlsx": ("2022-07", None, None),
    }
    man = _manifiesto(tmp_path)
    for rel, (carpeta, periodo, trimestre) in esperados.items():
        url = f"{FICHEROS}{carpeta}/{Path(rel).name}"
        assert (raw / rel).read_bytes() == portal.urls[url]            # original tal cual
        assert man[rel]["url"] == url and man[rel]["publicado"] is True
        assert (man[rel]["periodo"], man[rel]["trimestre"]) == (periodo, trimestre)
        assert man[rel]["dataset"] == rel.split("/")[0] and man[rel]["anio"] == int(rel.split("/")[1])
    assert man["informe_menores/2021/consejo_gobierno_menores_2021.xlsx"]["concepto"] == "117"
    # "Datos estadísticos": solo los ficheros de menores; los demás conceptos ni se piden
    assert portal.pedidas(f"{FICHEROS}2022-07/INFORME_PYMES_2021.xlsx") == 0
    assert portal.pedidas(M.URL_FICHEROS, concepto="114") == 0
    for anio in range(2015, ANIO + 1):
        assert portal.pedidas(M.URL_FICHEROS, concepto="113", year=anio) == 1
    # Lo que no casa con ninguna serie se guarda igual, con un aviso
    assert "'Relación de contratos menores del organismo'" in _log(tmp_path) and "va a 'otros'" in _log(tmp_path)
    inventario = json.loads((raw / "inventario_jccm.json").read_text(encoding="utf-8"))
    assert len(inventario["enlaces"]) == len(esperados)
    assert inventario["paginas"][f"concepto=113&year={VIEJO}"] == 4


def test_concepto_nuevo_de_menores_se_consulta_entero(portal, tmp_path):
    portal.conceptos["200"] = "Contratación menor: caja pagadora"
    publicar(portal, "200", ANIO, "2026-08", "caja_2T.xlsx", f"2º trimestre {ANIO}", _junta("C-1"))
    assert _ejecutar(tmp_path) == 0
    assert portal.pedidas(M.URL_FICHEROS, concepto="200", year=ANIO) == 1
    assert (tmp_path / "raw" / "otros" / str(ANIO) / "caja_2T.xlsx").exists()


@pytest.mark.parametrize("texto, nombre, esperado", [
    ("JCCM 1º trimestre 2019 gestor PICOS", "PUBLIC~1.XLS", ("menores_junta", "1T", "1")),
    ("JCCM 3º Trimestre 2023 gestor PICOS", "Contratos_menores_tercer_trimestre_2023_JCCM.zip",
     ("menores_junta", "3T", "3")),
    ("JCCM año 2020 COMPLETO gestor PICOS", "contratos_menores_jccm_ano_2020_completo-16-02-2021.xls",
     ("menores_junta", "anual", None)),
    ("JCCM 2º trimestre 2026 CAJA PAGADORA", "Contratos_menores_segundo_trimestre_2026_caja_pagadora.xlsx",
     ("caja_pagadora", "2T", "2")),
    ("SESCAM 4º Trimestre 2025", "CM_CUARTO_TRIMESTRE_2025_SESCAM.rar", ("sescam", "4T", "4")),
    ("SECTOR PÚBLICO REGIONAL 2º trimestre 2015", "sector_publico_regional_2o_semestre_2015.zip",
     ("sector_publico", "2T", "2")),
    ("Informe de Contratación Administrativa del Sector Público 2020 (contratos menores-fichero de datos)",
     "consejo_gobierno_menores_20.xlsx", ("informe_menores", None, None)),
    ("Descargar", "sescam-contratos_menores_trimestre_3-2019_0.rar", ("sescam", "3T", "3")),
    ("Descargar", "contratos_menores_cuarto_trimestre_2019_0.rar", ("otros", "4T", "4")),
    ("Descargar", "sector_publico_regional_2o_semestre_2015.zip", ("sector_publico", "2S", None)),
])
def test_clasificacion_de_los_enlaces_publicados(texto, nombre, esperado):
    assert (M.clasificar(texto, nombre),) + M.periodo_de(texto, nombre) == esperado


# ---------------------------------------------------------------------------
# Parquet por conjunto de datos
# ---------------------------------------------------------------------------

def test_parquet_por_dataset_como_texto_con_metadatos(portal, tmp_path):
    assert _ejecutar(tmp_path) == 0

    assert sorted(p.name for p in tmp_path.glob("*.parquet")) == sorted(
        f"{d}.parquet" for d in ["menores_junta", "caja_pagadora", "sector_publico", "sescam", "informe_menores",
                                 "uclm", "otros"])
    tabla = pq.read_table(tmp_path / "menores_junta.parquet")
    assert all(pa.types.is_string(f.type) for f in tabla.schema if f.name != "_en_ultima_descarga")
    junta = tabla.to_pandas()
    assert list(junta.columns[:len(CABECERA_JUNTA)]) == CABECERA_JUNTA
    assert list(junta.columns[len(CABECERA_JUNTA):]) == [c for c in M.ORDEN_METADATOS if c != "_miembro"]
    viejo = junta[junta["_anio"] == str(VIEJO)]
    # El anual repite los trimestres: se conservan los dos, distinguidos por _periodo
    assert viejo.groupby("_periodo").size().to_dict() == {"1T": 2, "anual": 3}
    assert set(viejo["Importe adjudicación IVA incluido"]) == {"1210.5"}
    assert set(viejo["Fecha de adjudicación"]) == {"2023-02-15"}
    assert set(viejo["_titulo"]) == {f"JCCM 1º trimestre {VIEJO} gestor PICOS", f"JCCM año {VIEJO} COMPLETO gestor PICOS"}
    assert set(junta["_unidad"]) == {"contrato menor"} and junta["_en_ultima_descarga"].all()
    assert _v(junta.loc[junta["_periodo"] == "anual", "_trimestre"]) == [None] * 3

    sescam = _parquet(tmp_path, "sescam")
    # Unión de esquemas de los dos XLSX del ZIP, sin perder columnas
    assert {"Nº de factura", f"Importe con IVA 1 TRIM {VIEJO}", "Importe  sin IVA 1 TRIM"} <= set(sescam.columns)
    assert sescam["_miembro"].tolist() == ["CM_PRIMER_TRIMESTRE_FARMACIA_SESCAM.xlsx",
                                           "CM_PRIMER_TRIMESTRE_SUMINISTROS_SESCAM.xlsx"]
    assert sescam[f"Importe con IVA 1 TRIM {VIEJO}"].tolist()[0] == "15523.93"
    assert set(sescam["_unidad"]) == {"línea de factura por artículo y gerencia (no es un contrato)"}

    publico = _parquet(tmp_path, "sector_publico")
    # ZIP con un ZIP dentro (nombre en cp850, sin marca UTF-8) y título encima de la cabecera
    assert publico["_miembro"].tolist() == ["0_seccion_18_contratos_menores_oct-dic_2018.xlsx",
                                            "sescam_4o_trimestre_2018.zip/4º trimestre 2018.xlsx"]
    assert _v(publico["Adjudicatario"]) == ["COMERCIAL GALAN, S.A.", None]
    assert _v(publico["Artículo"]) == [None, "GASAS"]
    assert publico["_anio"].tolist() == ["2018", "2018"] and set(publico["_periodo"]) == {"4T"}
    assert "2 filas antes de la cabecera" in _log(tmp_path)

    caja = _parquet(tmp_path, "caja_pagadora")
    assert caja[["NIF Adjudicatario", "Fecha Factura", "_hoja"]].values.tolist() == [
        ["B45047537", "2026-06-05", "Menores 1T C. Pagadora"]]
    informe = _parquet(tmp_path, "informe_menores")
    assert informe["IA"].tolist() == ["428.65"] and informe["_fuente"].tolist() == [
        f"{FICHEROS}2022-07/consejo_gobierno_menores_2021.xlsx"]


# ---------------------------------------------------------------------------
# UCLM: formulario ASP.NET
# ---------------------------------------------------------------------------

def test_uclm_un_postback_por_ejercicio_con_los_campos_del_formulario(portal, tmp_path):
    assert _ejecutar(tmp_path, "--fuente", "uclm") == 0

    assert sorted(int(p[M.CAMPO_EJERCICIO_UCLM]) for p in portal.posts) == [ANIO - 2, ANIO - 1]
    for datos in portal.posts:
        assert datos["__EVENTTARGET"] == M.BOTON_BUSCAR_UCLM and datos["__EVENTARGUMENT"] == ""
        assert (datos["__VIEWSTATE"], datos["__EVENTVALIDATION"]) == portal.viewstate_get
        assert datos["__VIEWSTATEGENERATOR"] == "0B7177FB"
    assert portal.pedidas(M.URL_UCLM_ACTUALES) == 1
    assert portal.pedidas(M.URL_FICHEROS) == 0                    # --fuente uclm: la Junta no se toca
    raw = tmp_path / "raw" / "uclm"
    guardado = (raw / str(ANIO - 1) / f"contratosMenoresAnteriores_{ANIO - 1}.html").read_text(encoding="utf-8")
    assert tabla_uclm(portal.uclm[ANIO - 1]) in guardado             # la tabla, tal cual
    assert "VS" not in guardado and "día" not in guardado            # sin el estado cifrado ni la fecha
    df = _parquet(tmp_path, "uclm")
    anteriores = df[df["_periodo"] == "anual"]
    assert anteriores["Nif"].tolist() == ["80162019G", "B42768705", "US000611648780"]
    assert anteriores["Proveedor"].tolist()[0] == " CELIA VELASCO LOPEZ"      # sin recortar
    assert anteriores["Objeto"].tolist()[1] == "REPARACIONES, MANTENIMIENTO Y CONSERVACIÓN"
    assert anteriores["Duración"].tolist() == ["30", "1", "1"] and anteriores["Adjudicado"].tolist()[1] == "3.630,00 €"
    # Menores actuales: el año siguiente al último de los anteriores, con la unidad funcional
    actuales = df[df["_periodo"] == "en curso"]
    assert actuales[["Un.Funcional", "Nif", "_anio"]].values.tolist() == [["01110C0021", "05679241N", str(ANIO)]]
    assert _v(anteriores["Un.Funcional"]) == [None] * 3
    # Unión de esquemas entre ficheros: la columna que solo trae el último no se pierde
    assert list(df.columns[:7]) == ["Nif", "Proveedor", "Expediente", "Objeto", "Duración", "Adjudicado",
                                    "Un.Funcional"]
    assert df["_fuente"].tolist() == [M.URL_UCLM_ANTERIORES] * 3 + [M.URL_UCLM_ACTUALES]
    man = _manifiesto(tmp_path)
    assert man[f"uclm/{ANIO - 1}/contratosMenoresAnteriores_{ANIO - 1}.html"]["filas"] == 2
    assert man[f"uclm/{ANIO}/contratosMenoresActuales_{ANIO}.html"]["metodo"] == "GET"


def test_uclm_actuales_es_el_ejercicio_siguiente_al_ultimo_cerrado(portal, tmp_path):
    # En enero la UCLM aún no ha pasado el año anterior a "anteriores": la página
    # de actuales sigue siendo la de ese año, no la del calendario
    portal.uclm = {ANIO - 3: portal.uclm[ANIO - 2], ANIO - 2: portal.uclm[ANIO - 1]}
    assert _ejecutar(tmp_path, "--fuente", "uclm") == 0
    assert (tmp_path / "raw" / "uclm" / str(ANIO - 1) / f"contratosMenoresActuales_{ANIO - 1}.html").exists()
    assert json.loads((tmp_path / "raw" / "inventario_uclm.json").read_text(encoding="utf-8")) == {
        "anteriores": [ANIO - 3, ANIO - 2], "actuales": ANIO - 1}


def test_uclm_no_crea_versiones_si_solo_cambia_el_viewstate(portal, tmp_path):
    assert _ejecutar(tmp_path, "--fuente", "uclm") == 0
    SLEEP_REAL(1.1)
    assert _ejecutar(tmp_path, "--fuente", "uclm") == 0
    # Cada respuesta trae otro __VIEWSTATE y otra fecha, pero los datos son los mismos
    assert not list((tmp_path / "raw" / "uclm").rglob(M.HISTORICO))
    assert f"uclm/{ANIO - 1}/contratosMenoresAnteriores_{ANIO - 1}.html (2 filas)" in _log(tmp_path).split(
        "SIN CAMBIOS")[-1]
    df = _parquet(tmp_path, "uclm")
    assert len(df) == 4 and df["_en_ultima_descarga"].all()


def test_uclm_ejercicio_retirado_conserva_sus_filas(portal, tmp_path):
    assert _ejecutar(tmp_path, "--fuente", "uclm") == 0
    SLEEP_REAL(1.1)
    del portal.uclm[ANIO - 2]                     # el desplegable ya no lo ofrece
    assert _ejecutar(tmp_path, "--fuente", "uclm") == 0
    df = _parquet(tmp_path, "uclm")
    retirado = df["_anio"] == str(ANIO - 2)
    assert df.loc[retirado, "Nif"].tolist() == ["80162019G"]
    assert not df.loc[retirado, "_en_ultima_descarga"].any() and df.loc[~retirado, "_en_ultima_descarga"].all()
    assert _manifiesto(tmp_path)[f"uclm/{ANIO - 2}/contratosMenoresAnteriores_{ANIO - 2}.html"]["publicado"] is False
    assert "la UCLM ya no lo ofrece" in _log(tmp_path)


@pytest.mark.parametrize("fallo", ["get", "post"])
def test_uclm_caida_no_retira_nada(portal, tmp_path, monkeypatch, fallo):
    assert _ejecutar(tmp_path, "--fuente", "uclm") == 0
    antes = _parquet(tmp_path, "uclm")
    SLEEP_REAL(1.1)
    if fallo == "get":
        portal.uclm_get = 503
    else:
        # El servidor rechaza todos los postback (HTTP 500): la tabla no llega
        monkeypatch.setattr(requests, "post", lambda url, **kw: FakeResponse(status=500))
    assert _ejecutar(tmp_path, "--fuente", "uclm", "--comprobar-todo") == 1
    despues = _parquet(tmp_path, "uclm")
    assert len(despues) == len(antes) and despues["_en_ultima_descarga"].all()
    assert all(i["publicado"] for i in _manifiesto(tmp_path).values())
    assert "uclm" in _log(tmp_path).split("ERRORES")[-1]


# ---------------------------------------------------------------------------
# Sesgo del superviviente: retirado, fallos del portal y versiones
# ---------------------------------------------------------------------------

def test_fichero_que_deja_de_enlazarse_queda_retirado(portal, tmp_path):
    assert _ejecutar(tmp_path, "--fuente", "jccm") == 0
    SLEEP_REAL(1.1)
    portal.paginas[("113", VIEJO)] = [e for e in portal.paginas[("113", VIEJO)] if "SESCAM" not in e[1]]
    assert _ejecutar(tmp_path, "--fuente", "jccm") == 0

    rel = f"sescam/{VIEJO}/CM_PRIMER_TRIMESTRE_SESCAM.zip"
    assert (tmp_path / "raw" / rel).exists()
    man = _manifiesto(tmp_path)
    assert man[rel]["publicado"] is False and "ya no lo enlaza" in man[rel]["detalle"]
    sescam = _parquet(tmp_path, "sescam")
    assert len(sescam) == 2 and not sescam["_en_ultima_descarga"].any()
    assert _parquet(tmp_path, "menores_junta")["_en_ultima_descarga"].all()
    assert f"{rel}: la página ya no lo enlaza" in _log(tmp_path)


@pytest.mark.parametrize("pagina", [503, []])
def test_pagina_caida_o_sin_enlaces_no_retira_nada(portal, tmp_path, pagina):
    assert _ejecutar(tmp_path, "--fuente", "jccm") == 0
    SLEEP_REAL(1.1)
    portal.paginas[("113", VIEJO)] = pagina
    assert _ejecutar(tmp_path, "--fuente", "jccm") == 1
    assert all(i["publicado"] for i in _manifiesto(tmp_path).values())
    assert _parquet(tmp_path, "sescam")["_en_ultima_descarga"].all()
    assert f"año {VIEJO}" in _log(tmp_path).split("ERRORES")[-1]


def test_portada_caida_no_retira_nada_y_sigue_con_la_uclm(portal, tmp_path):
    assert _ejecutar(tmp_path) == 0
    SLEEP_REAL(1.1)
    portal.portada = 503
    assert _ejecutar(tmp_path) == 1
    assert all(i["publicado"] for i in _manifiesto(tmp_path).values())
    assert "jccm: no se pudo leer" in _log(tmp_path)
    assert portal.pedidas(M.URL_UCLM_ANTERIORES) == 2


@pytest.mark.parametrize("fallo", [503, b"<!DOCTYPE html><html><body>Mantenimiento</body></html>"])
def test_descarga_fallida_no_pierde_la_copia(portal, tmp_path, fallo):
    url = f"{FICHEROS}2026-05/Contratos_menores_primer_trimestre_{ANIO}_JCCM.xlsx"
    original = portal.urls[url]
    assert _ejecutar(tmp_path, "--fuente", "jccm") == 0
    SLEEP_REAL(1.1)
    portal.urls[url] = fallo
    assert _ejecutar(tmp_path, "--fuente", "jccm") == 1
    raw = tmp_path / "raw" / "menores_junta" / str(ANIO) / Path(url).name
    assert raw.read_bytes() == original and len(M.versiones(raw)) == 1
    assert _parquet(tmp_path, "menores_junta")["_en_ultima_descarga"].all()
    assert _manifiesto(tmp_path)[f"menores_junta/{ANIO}/{Path(url).name}"]["publicado"] is True


def test_zip_cortado_se_vuelve_a_pedir(portal, tmp_path):
    url = f"{FICHEROS}2024-05/CM_PRIMER_TRIMESTRE_SESCAM.zip"
    bueno = portal.urls[url]
    portal.urls[url] = [bueno[: len(bueno) // 2], bueno]           # sin su directorio central
    assert _ejecutar(tmp_path, "--fuente", "jccm") == 0
    assert portal.pedidas(url) == 2
    assert (tmp_path / "raw" / "sescam" / str(VIEJO) / Path(url).name).read_bytes() == bueno


def test_descarga_mas_corta_que_su_content_length_se_vuelve_a_pedir(portal, tmp_path):
    # Un CSV cortado sigue pareciendo un CSV: solo lo delata el Content-Length
    bueno = "Expediente;Importe\nA;10\nB;20\n".encode("utf-8")
    publicar(portal, "113", ANIO, "2026-08", "menores_JCCM_2T.csv", f"JCCM 2º trimestre {ANIO} gestor PICOS",
             [(bueno[:-6], {"Content-Length": str(len(bueno))}), bueno])
    assert _ejecutar(tmp_path, "--fuente", "jccm") == 0
    assert portal.pedidas(f"{FICHEROS}2026-08/menores_JCCM_2T.csv") == 2
    df = _parquet(tmp_path, "menores_junta")
    assert df.loc[df["_titulo"] == f"JCCM 2º trimestre {ANIO} gestor PICOS", "Expediente"].tolist() == ["A", "B"]


def test_dos_ficheros_con_el_mismo_nombre_no_se_pisan(portal, tmp_path):
    for carpeta, articulo in (("2025-07", "PRIMERO"), ("2025-12", "CORREGIDO")):
        publicar(portal, "113", ANIO, carpeta, "CM_SESCAM.zip", f"SESCAM 3º Trimestre {ANIO} ({carpeta})",
                 _zip({"a.xlsx": _sescam_xlsx(ANIO, articulo)}))
    assert _ejecutar(tmp_path, "--fuente", "jccm") == 0
    carpeta = tmp_path / "raw" / "sescam" / str(ANIO)
    assert sorted(p.name for p in carpeta.iterdir()) == ["2025-07_CM_SESCAM.zip", "2025-12_CM_SESCAM.zip"]
    df = _parquet(tmp_path, "sescam")
    assert sorted(df.loc[df["_anio"] == str(ANIO), "Artículo"]) == ["CORREGIDO", "PRIMERO"]


def test_fichero_modificado_conserva_las_filas_anteriores(portal, tmp_path):
    url = f"{FICHEROS}2026-05/Contratos_menores_primer_trimestre_{ANIO}_JCCM.xlsx"
    portal.urls[url] = _junta("A", "B", "C")
    assert _ejecutar(tmp_path, "--fuente", "jccm") == 0
    SLEEP_REAL(1.1)          # la versión nueva tiene que llevar otra fecha (resolución: 1 s)
    portal.urls[url] = _junta("A", "C", "D")
    assert _ejecutar(tmp_path, "--fuente", "jccm") == 0

    df = _parquet(tmp_path, "menores_junta")
    actual = df[df["_anio"] == str(ANIO)]
    assert actual[["Nº expediente PICOS", "_en_ultima_descarga"]].values.tolist() == [
        ["A", True], ["B", False], ["C", True], ["D", True]]
    assert df.loc[df["_anio"] == str(VIEJO), "_en_ultima_descarga"].all()
    raw = tmp_path / "raw" / "menores_junta" / str(ANIO) / Path(url).name
    assert len(M.versiones(raw)) == 2


def test_segunda_ejecucion_solo_lee_lo_que_se_ha_vuelto_a_comprobar(portal, tmp_path, monkeypatch):
    assert _ejecutar(tmp_path) == 0
    SLEEP_REAL(1.1)
    leidos = []
    leer = M.leer_tabla
    monkeypatch.setattr(M, "leer_tabla", lambda ruta: leidos.append(Path(ruta).name) or leer(ruta))
    assert _ejecutar(tmp_path) == 0
    # Los ficheros de años cerrados no se vuelven a pedir ni a leer (millones de
    # filas del SESCAM): sus filas salen del Parquet anterior
    assert sorted(leidos) == sorted([f"Contratos_menores_primer_trimestre_{ANIO}_JCCM.xlsx",
                                     "Contratos_menores_primer_trimestre_caja_pagadora.XLSX",
                                     f"contratosMenoresAnteriores_{ANIO - 1}.html",
                                     f"contratosMenoresActuales_{ANIO}.html"])
    sescam = _parquet(tmp_path, "sescam")
    assert len(sescam) == 2 and sescam["_en_ultima_descarga"].all()


def test_fichero_que_ya_no_esta_en_raw_conserva_sus_filas(portal, tmp_path):
    assert _ejecutar(tmp_path) == 0
    antes = _parquet(tmp_path, "sescam")
    (tmp_path / "raw" / "sescam" / str(VIEJO) / "CM_PRIMER_TRIMESTRE_SESCAM.zip").unlink()
    assert _ejecutar(tmp_path, "--solo-parquet") == 0
    despues = _parquet(tmp_path, "sescam")
    pd.testing.assert_frame_equal(despues, antes)
    assert f"sescam/{VIEJO}/CM_PRIMER_TRIMESTRE_SESCAM.zip ya no está en raw/" in _log(tmp_path)


def test_solo_se_vuelven_a_pedir_el_anio_actual_y_el_anterior(portal, tmp_path):
    viejo = f"{FICHEROS}2023-06/CM_PRIMER_TRIMESTRE_JCCM.xlsx"
    actual = f"{FICHEROS}2026-05/Contratos_menores_primer_trimestre_{ANIO}_JCCM.xlsx"
    _ejecutar(tmp_path)
    _ejecutar(tmp_path)
    assert (portal.pedidas(viejo), portal.pedidas(actual)) == (1, 2)
    assert len(portal.posts) == 2 + 1              # la UCLM de hace dos años no se vuelve a pedir
    _ejecutar(tmp_path, "--comprobar-todo")
    assert portal.pedidas(viejo) == 2
    assert not list((tmp_path / "raw").rglob(M.HISTORICO))    # nada cambió


# ---------------------------------------------------------------------------
# ZIP, RAR, XLS y tablas HTML
# ---------------------------------------------------------------------------

def test_zip_con_anidados_restos_y_ficheros_que_no_son_tablas(tmp_path):
    ruta = tmp_path / "sector_publico_1o_trimestre_2016.zip"
    ruta.write_bytes(_zip({
        "EDUCACIÓN/contratos menores 1T 2016.csv": "Objeto;Importe\nSeñalización – 5 €;1.000,00\n".encode("cp1252"),
        "__MACOSX/EDUCACIÓN/._contratos menores 1T 2016.csv": b"\x00\x05\x16\x07basura",
        "nota.pdf": b"%PDF-1.4 relacion firmada",
        "LEEME.txt": "Relación remitida por las consejerías; importes sin IVA\n".encode("utf-8"),
        "anexo.zip": _zip({"Fomento.xlsx": _xlsx({"Hoja1": [["Objeto", "Importe"], ["Bacheo", 7.0]]})}),
        # Página del portal guardada desde el navegador: no es una tabla aunque se llame .html
        "Año 2016 (cuarto trimestre) _ Portal.html": MHTML,
    }, utf8=False))
    df, avisos = M.leer_tabla(ruta)
    # Nombres en cp850 sin la marca UTF-8 (como los ZIP hechos en Windows): 'Ó' no es de cp437
    assert df["_miembro"].tolist() == ["EDUCACIÓN/contratos menores 1T 2016.csv", "anexo.zip/Fomento.xlsx"]
    assert df["Objeto"].tolist() == ["Señalización – 5 €", "Bacheo"]
    assert df["Importe"].tolist() == ["1.000,00", "7"]
    assert any("nota.pdf: no es una tabla (PDF)" in a for a in avisos)
    assert any("LEEME.txt: no es una tabla (CSV)" in a for a in avisos)       # texto, pero no una tabla
    assert any("Portal.html: no es una tabla (MHTML)" in a for a in avisos)
    assert any("__MACOSX" in a and "se ignora" in a for a in avisos)


def test_hoja_xml_de_excel_2003(tmp_path):
    # Como el RAR del SESCAM de 2016 "en formato xml": SpreadsheetML con dos hojas
    ruta = tmp_path / "menores.zip"
    ruta.write_bytes(_zip({"Menores Sescam 2 trimestre 2016.xml": SPREADSHEETML}))
    df, avisos = M.leer_tabla(ruta)
    assert df["_hoja"].tolist() == ["Contratos Suministros-Servicios"] * 2 + ["Contratos Farmacia"]
    assert _v(df["Gerencia"]) == ["61037000 GUETS", None, "61035000 GAE Toledo"]
    assert _v(df["Artículo"]) == ["LIMPIEZA", "GASAS estériles", "U-005"]    # texto con formato: su texto
    assert _v(df["Importe"]) == ["28.8808", "7", None]
    assert _v(df["Fecha"]) == ["2016-07-06", None, "2016-07-06 10:30:00"]
    assert _v(df["Unnamed: 1"]) == [None, None, "sobra"]          # la celda combinada ocupa dos columnas
    assert "2 hojas con datos" in " ".join(avisos)


def test_si_un_extractor_falla_a_mitad_se_usa_el_siguiente_sin_repetir(tmp_path, monkeypatch):
    ruta = tmp_path / "x.zip"
    ruta.write_bytes(_zip({"a.csv": b"A\n1\n", "b.csv": b"A\n2\n"}))

    def a_medias(ruta, carpeta):
        yield next(M._extraer_zip(ruta, carpeta))
        raise OSError("File CRC error")

    monkeypatch.setattr(M, "_extractores", lambda formato: [("roto", a_medias), ("bueno", M._extraer_zip)])
    df, _ = M.leer_tabla(ruta)
    assert df[["A", "_miembro"]].values.tolist() == [["1", "a.csv"], ["2", "b.csv"]]


@pytest.mark.parametrize("disponibles, lista", [(["libarchive"], "pendientes"),
                                                (["rarfile", "libarchive"], "fallidos")])
def test_rar_que_no_se_puede_descomprimir(portal, tmp_path, monkeypatch, disponibles, lista):
    # libarchive falla con algún RAR válido (el del SESCAM 2016-2T): se guarda igual
    # y queda pendiente de leerlo con unrar; si también falla unrar, es un error
    def falla(ruta, carpeta):
        raise OSError("File CRC error")
        yield

    real = M._extractores
    monkeypatch.setattr(M, "_extractores", lambda formato: [(n, falla) for n in disponibles]
                        if formato == "rar" else real(formato))
    _publicar_rar(portal)
    assert _ejecutar(tmp_path, "--fuente", "jccm") == 1
    rar = tmp_path / "raw" / "sescam" / str(ANIO) / f"CM_SEGUNDO_TRIMESTRE_{ANIO}_SESCAM.rar"
    assert rar.read_bytes() == portal.urls[f"{FICHEROS}2026-09/{rar.name}"]
    bloques = _log(tmp_path).split("PENDIENTES")[-1].split("ERRORES")
    bloque = bloques[0] if lista == "pendientes" else bloques[-1]
    assert f"sescam/{ANIO}/{rar.name}" in bloque and "File CRC error" in bloque


def test_rar_danado_se_guarda_igual(portal, tmp_path):
    # Un RAR que el descompresor de aquí no sabe leer (CRC que no casa) no se
    # rechaza en la descarga: el original se guarda y el fallo sale al leerlo
    rar = bytearray(_rar4("datos.csv", b"A\n1\n"))
    rar[7 + 13 + 7 + 9:7 + 13 + 7 + 13] = b"\x00\x00\x00\x00"          # FILE_CRC
    publicar(portal, "113", ANIO, "2026-09", "CM_TERCER_TRIMESTRE_SESCAM.rar", f"SESCAM 3º Trimestre {ANIO}",
             bytes(rar))
    assert _ejecutar(tmp_path, "--fuente", "jccm") == 1
    guardado = tmp_path / "raw" / "sescam" / str(ANIO) / "CM_TERCER_TRIMESTRE_SESCAM.rar"
    assert guardado.read_bytes() == bytes(rar)
    assert _manifiesto(tmp_path)[f"sescam/{ANIO}/{guardado.name}"]["publicado"] is True
    assert f"sescam/{ANIO}/{guardado.name}" in _log(tmp_path).split("PENDIENTES")[-1]


def test_hoja_xls_llena_puede_estar_truncada(tmp_path, monkeypatch):
    xlwt = pytest.importorskip("xlwt")
    monkeypatch.setattr(M, "FILAS_MAXIMAS_XLS", 3)
    libro = xlwt.Workbook()
    hoja = libro.add_sheet("menores")
    for i, fila in enumerate([["Expediente"], ["A"], ["B"]]):
        hoja.write(i, 0, fila[0])
    ruta = tmp_path / "MENORES_2015 Sescam sin farmacia.xls"
    libro.save(str(ruta))
    df, avisos = M.leer_tabla(ruta)
    assert df["Expediente"].tolist() == ["A", "B"]
    assert any("[menores]: 3 filas, el máximo de un .xls" in a for a in avisos)


def test_zip_sin_tablas_es_un_error(tmp_path):
    ruta = tmp_path / "x.zip"
    ruta.write_bytes(_zip({"nota.pdf": b"%PDF-1.4"}))
    with pytest.raises(ValueError, match="ninguna tabla"):
        M.leer_tabla(ruta)


def _publicar_rar(portal):
    publicar(portal, "113", ANIO, "2026-09", f"CM_SEGUNDO_TRIMESTRE_{ANIO}_SESCAM.rar", f"SESCAM 2º Trimestre {ANIO}",
             _rar4("02. contratos menores FARMACIA.xlsx", _sescam_xlsx(ANIO, "BENRALIZUMAB")))


@pytest.mark.skipif(not HAY_LECTOR_RAR, reason="hace falta libarchive-c o rarfile")
def test_rar_del_sescam_se_descomprime_y_se_lee(portal, tmp_path):
    _publicar_rar(portal)
    assert _ejecutar(tmp_path, "--fuente", "jccm") == 0
    sescam = _parquet(tmp_path, "sescam")
    rar = sescam[sescam["_anio"] == str(ANIO)]
    assert rar[["Artículo", "_miembro", "_periodo"]].values.tolist() == [
        ["BENRALIZUMAB", "02. contratos menores FARMACIA.xlsx", "2T"]]
    assert (tmp_path / "raw" / "sescam" / str(ANIO) / f"CM_SEGUNDO_TRIMESTRE_{ANIO}_SESCAM.rar").read_bytes()[:4] == b"Rar!"


def test_rar_sin_lector_queda_pendiente_sin_perder_el_original(portal, tmp_path, monkeypatch):
    real = M._modulo
    monkeypatch.setattr(M, "_modulo", lambda nombre: None if nombre in ("libarchive", "rarfile") else real(nombre))
    _publicar_rar(portal)
    assert _ejecutar(tmp_path, "--fuente", "jccm") == 1
    rar = tmp_path / "raw" / "sescam" / str(ANIO) / f"CM_SEGUNDO_TRIMESTRE_{ANIO}_SESCAM.rar"
    assert rar.read_bytes() == portal.urls[f"{FICHEROS}2026-09/{rar.name}"]
    log = _log(tmp_path)
    assert "PENDIENTES" in log and f"sescam/{ANIO}/{rar.name}" in log.split("PENDIENTES")[-1]
    assert "libarchive-c" in log
    # Las demás filas del SESCAM sí están; las del RAR, no
    assert _parquet(tmp_path, "sescam")["_anio"].tolist() == [str(VIEJO)] * 2

    if not HAY_LECTOR_RAR:
        return
    # En cuanto hay con qué leerlo, sus filas entran sin volver a descargar nada
    monkeypatch.setattr(M, "_modulo", real)
    assert _ejecutar(tmp_path, "--solo-parquet") == 0
    assert _parquet(tmp_path, "sescam")["_anio"].tolist() == [str(VIEJO)] * 2 + [str(ANIO)]


def test_xls_binario(tmp_path):
    xlwt = pytest.importorskip("xlwt")
    libro = xlwt.Workbook()
    hoja = libro.add_sheet("Contratos")
    fecha = xlwt.easyxf(num_format_str="DD/MM/YYYY")
    for c, v in enumerate(["Nº expediente PICOS", "Importe", "Fecha"]):
        hoja.write(0, c, v)
    hoja.write(1, 0, "2019/000123")
    hoja.write(1, 1, 1234.5)
    hoja.write(1, 2, datetime(2019, 3, 4), fecha)
    hoja.write(2, 0, "X")
    hoja.write(2, 1, 10)
    ruta = tmp_path / "PUBLIC~1.XLS"
    libro.save(str(ruta))
    df, _ = M.leer_tabla(ruta)
    assert df["Nº expediente PICOS"].tolist() == ["2019/000123", "X"]
    assert df["Importe"].tolist() == ["1234.5", "10"]
    assert _v(df["Fecha"]) == ["2019-03-04", None]


def test_tabla_html_conserva_el_texto_de_cada_celda(tmp_path):
    ruta = tmp_path / "tabla.html"
    ruta.write_bytes(M.documento_tabla("t", M.URL_UCLM_ANTERIORES, tabla_uclm([
        ["B1", "  ACME &amp; CIA ", "SE1", "L&#205;NEA 1<br/>L&#205;NEA 2", "&nbsp;", ""],
        ["B2", "OTRO", "SE2", "<span>OBJETO</span> <b>X</b>", "1", "5,00 €", "sobra"],
    ])))
    df, avisos = M.leer_tabla(ruta)
    assert list(df.columns) == ["Nif", "Proveedor", "Expediente", "Objeto", "Duración", "Adjudicado",
                                "_columna_extra_1"]
    assert df["Proveedor"].tolist() == ["  ACME & CIA ", "OTRO"]
    assert df["Objeto"].tolist() == ["LÍNEA 1\nLÍNEA 2", "OBJETO X"]
    assert _v(df["Duración"]) == ["\xa0", "1"] and _v(df["Adjudicado"]) == [None, "5,00 €"]
    assert _v(df["_columna_extra_1"]) == [None, "sobra"]
    assert any("1 filas con más celdas" in a for a in avisos)


def test_tabla_html_de_la_pagina(tmp_path):
    pagina = pagina_uclm([2025], "VS", "EV", "hoy", 2025, tabla_uclm([["B1", "P", "E", "O", "1", "1,00 €"]]))
    fragmento = M.tabla_html(pagina.decode("utf-8"), M.TABLA_UCLM)
    assert fragmento.startswith("<table") and fragmento.endswith("</table>") and "VS" not in fragmento
    assert M.tabla_html("<html><table id='otra'></table></html>", M.TABLA_UCLM) is None
    campos = M.campos_ocultos(pagina.decode("utf-8"))
    assert campos["__VIEWSTATE"] == "VS" and campos["__EVENTVALIDATION"] == "EV" and campos["__EVENTTARGET"] == ""


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def test_cli(portal, tmp_path, monkeypatch):
    monkeypatch.setattr(sys, "argv", [str(SCRIPT), "--salida", str(tmp_path), "--desde", str(ANIO)])
    with pytest.raises(SystemExit) as salida:
        runpy.run_path(str(SCRIPT), run_name="__main__")
    assert salida.value.code == 0
    # --desde: solo se consultan los años desde ese
    assert portal.pedidas(M.URL_FICHEROS, concepto="113", year=VIEJO) == 0
    assert (tmp_path / "caja_pagadora.parquet").exists() and not (tmp_path / "sescam.parquet").exists()


def test_salida_por_defecto_en_el_repo():
    assert M.SALIDA == REPO_ROOT / "ccaa_castilla_la_mancha"
