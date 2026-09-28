"""Tests offline de scripts/municipios_menores.py.

Los portales (datos abiertos de Gijón y Vigo, páginas de Valladolid,
Fuenlabrada, Leganés y Santa Cruz de Tenerife, y los CKAN de Málaga, Córdoba y
Vigo) se simulan con un ``requests.get`` falso. Páginas y ficheros son recortes
de las muestras reales descargadas el 2026-09-27 (mismas cabeceras, primeras
filas, misma estructura de la página); los Excel y ODS se generan en el test.
"""

import importlib.util
import io
import json
import runpy
import sys
import time
import zipfile
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
import pytest
import requests

REPO_ROOT = Path(__file__).resolve().parents[1]
SCRIPT = REPO_ROOT / "scripts" / "municipios_menores.py"
SLEEP_REAL = time.sleep


def _cargar():
    spec = importlib.util.spec_from_file_location("municipios_menores", SCRIPT)
    modulo = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(modulo)
    return modulo


M = _cargar()
ANIO = datetime.now().year          # el script decide qué refrescar con el año real


# ---------------------------------------------------------------------------
# Ficheros (recortes de las muestras reales)
# ---------------------------------------------------------------------------

def _gijon(*registros):
    """JSON con la forma del de Gijón: espacio inicial, contratos.contrato y ' : '."""
    return (' {"contratos":{"contrato":[' + ",".join(registros) + "]}}").encode("utf-8")


def _registro_gijon(codigo, ejercicio, presupuesto, valor, precio, cif):
    return ('{"codigo_contrato" : "%s","lote" : "","numero_expediente" : "29843J/%s",'
            '"objeto_del_contrato" : "Obsequios protocolarios intercambio jias-caea","ejercicio" : %s,'
            '"poder_adjudicador" : "Divertia Gijón SA","organo_de_contrataciÓn" : "Gerencia de Divertia Gijón",'
            '"fase_del_contrato" : "ADJUDICADO","duraciÓn_del_contrato" : " ",'
            '"presupuesto_base_de_licitaciÓn" : "%s","valor_estimado_sin_iva" : %s,"precio_de_adjudicaciÓn" : %s,'
            '"importe_del_iva" : "10,41","procedimiento_de_adjudicaciÓn" : "Contrato menor",'
            '"fecha_adjudicacion" : "09/07/18","tipo_de_contrato" : "Suministro","cif_adjudicatario" : "%s",'
            '"identidad_del_adjudicatario" : "SUAREZ*CUERVO,JORDAN","concurrencia" : 1}'
            % (codigo, ejercicio, ejercicio, presupuesto, valor, precio, cif))


# Primer registro real; el valor estimado 14750.50 del tercero es inventado (el
# JSON real no trae decimales con punto) para ver que no pasa por float
REG_A = _registro_gijon("09.GERENTEDIVERTIAMEN2018002081", 2018, "49,59", 0, 60, "**35**63**")
REG_B = _registro_gijon("02.DIRECCIONFMCEUPMEN2024003484", 2024, "169,4", 140, '"169,4"', "0***882***")
REG_C = _registro_gijon("11.URBANISMOMEN2026000001", 2026, "17847,5", "14750.50", '"15427,5"', "A28242808")

VIGO_2019 = ('unidade,asunto,tipo,adxudicatario,importe,num_licitadores,fin_contrato,data\n'
             '"Área de Recursos Humanos e Formación",ACCION FORMATIVA S.E.I.S. RESCATE ACUÁTICO SUPERFICIAL,'
             'SERVIZOS,"2. Seguridad y Rescate Formación, S.L",4513.0,2,2019/12/31,2019/04/04\n'
             'ESTATÍSTICA,"CONTRATO MENOR DE SUBMINISTRO\nPAPELETAS ELECCIONES LOCALESD",SUBMINISTROS,JADFEL,'
             '13157.3,1,2019/05/26,2019/04/30\n'
             '"ÁREA DE CULTURA","CONTRATACIÓN DA INSERCIÓN PUBLICITARIA DA PROGRAMACIÓN VIGOCULTURA NOS NÚMEROS '
             'DE OUTUBRO E NOVEMBRO DA REVISTA ""A MOVIDA""",SERVIZOS,A MOVIDA CULTURAL S.L.L.,1452.0,1,'
             '2019/12/31,2019/07/17\n').encode("utf-8")


def _vigo(anio):
    return ('unidade,asunto,tipo,adxudicatario,importe,num_licitadores,fin_contrato,data,expediente\n'
            'DESENVOLVEMENTO LOCAL E EMPREGO,Oe vigo capacita x: subministro manuais didácticos,Subministro,'
            f'"DISFERP, S.C.",841.3,2,{anio}/08/20,{anio}/02/03,21362/77\n').encode("utf-8")


def _xlsx(hojas):
    libro = openpyxl.Workbook()
    libro.remove(libro.active)
    for nombre, filas in hojas.items():
        hoja = libro.create_sheet(nombre)
        for fila in filas:
            hoja.append(fila)
    datos = io.BytesIO()
    libro.save(datos)
    return datos.getvalue()


CABECERA_SICALWIN = ["Nº Operación", "Fase", "Fecha", "Referencia", "Aplicación", "Importe", "Nombre Ter.",
                     "Texto Libre", "Oper.", "Localidad", "Provincia", "Tipo Contrato", "Procedim. Contrato",
                     "Criterio Contrato", "PROCEDIMIENTO", "TRIMESTRE", "CAPITULO", "CAPITULO EN TEXTO", "AREA",
                     "AREA EN TEXTO", "PROGRAMA", "PROGRAMA EN TEXTO", "ARTICULO", "CONCEPTO"]


def _valladolid(anio):
    """XLSX de SICALWIN: lista de códigos, operaciones y una tabla dinámica."""
    return _xlsx({
        "Hoja1": [["01", "01. Alcaldía"], ["02", "02. Urbanismo y vivienda"]],
        "OPERACIONES SICALWIN": [
            CABECERA_SICALWIN,
            [220250001196, "AD", datetime(anio, 2, 11), 22025001011, f"{anio} 11 1361 16200", 2178,
             "PETZL ESPAÑA.S.L.", "MATRÍCULA PARA 2 BOMBEROS CURSO FORMACIÓN VAX003", 220, "BARCELONA", "BARCELONA",
             "Servicios - De servicios", "AdDirec - Adjudicación Directa", "SinC - Sin Criterio",
             "Contratación menor", "Primero", "1", "1. Gastos de personal", "11",
             "11. Salud pública y seguridad ciudadana", "1361", "1361. Prevención y extinción de incencios",
             16, 16200],
            [220250005020, "ADO", datetime(anio, 3, 21), 22025001224, f"{anio} 08 1341 22699", 726.5,
             "AC CAMERFIRMA,S.A", "CE-SELLO.ELECTRONICO NIVEL MEDIO (PSC) SW - 3 AÑOS", 240, "MADRID", "MADRID",
             "Servicios - De servicios", "AdDirec - Adjudicación Directa", "SinC - Sin Criterio",
             "Contratación menor", "Primero", "2", "2. Gastos corrientes en bienes y servicios", "08",
             "08. Tráfico y movilidad", "1341", "1341. Movilidad", 22, 22699]],
        "MENOR-AREA": [["TRIMESTRE", "(Todas)"], [], ["Suma de Importe", "Etiquetas de columna"],
                       ["TERCEROS", "01. Alcaldía", "Total general"], ["1A INGENIEROS,  S.L.P.", 0, 1573]],
    })


def _fuenlabrada_xlsx():
    """XLSX de Fuenlabrada: una hoja por organismo con una fila de título, filas
    de subtotal y una hoja que solo dice que no hay contratos."""
    return _xlsx({
        "Menores_ayto_2trim_2026": [
            ["CONTRATOS MENORES. AYUNTAMIENTO DE FUENLABRADA\nSEGUNDO TRIMESTRE DE 2026"],
            ["Num. Expe.", "Organismo", "Denominación", "Fecha Aprobación", "Duración", "Nº Licitadores", "CIF",
             "Adjudicatario", "Importe"],
            ["2026/MSV/000911", "CONCEJALIA DE CULTURA", "Actuación grupo La Pedrá tributo a Extremoduro en el "
             "recinto ferial el 14 de septiembre", "24/06/2026", "1 DIA", 1, "B83154740",
             "A VALLEKAS EDICIONES MUSICALES S. L    ", 11495],
            [None, None, None, None, None, None, None, "Total A VALLEKAS EDICIONES MUSICALES S. L    ", 11495]],
        "OTAF 4ºTR 2018": [["NINGÚN  CONTRATO EN EL ÚLTIMO TRIMESTRE DE 2018"]],
    })


_NS_ODS = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
           'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
           'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"')


def _ods(tablas_xml):
    contenido = (f'<?xml version="1.0" encoding="UTF-8"?><office:document-content {_NS_ODS}><office:body>'
                 f'<office:spreadsheet>{tablas_xml}</office:spreadsheet></office:body></office:document-content>')
    datos = io.BytesIO()
    with zipfile.ZipFile(datos, "w") as archivo:
        archivo.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet")
        archivo.writestr("content.xml", contenido)
    return datos.getvalue()


def _texto_ods(texto):
    return f'<table:table-cell office:value-type="string"><text:p>{texto}</text:p></table:table-cell>'


# AYTO-1T-2019.ods de Fuenlabrada: columna A vacía, fila de título, una celda con
# dos párrafos, <text:s> (espacios), fecha, números, celdas repetidas y combinadas
# y el millón de filas vacías del final
FUENLABRADA_ODS = _ods(
    '<table:table table:name="Sumario">'
    '<table:table-row><table:table-cell/>' + _texto_ods("AYUNTAMIENTO DE FUENLABRADA") +
    '<table:table-cell table:number-columns-repeated="1000"/></table:table-row>'
    '<table:table-row><table:table-cell/>'
    + "".join(_texto_ods(t) for t in ["Num. Expe.", "Título", "Fecha Aprobación", "Plazo Ejec.", "Importe",
                                      "Num. Licitadores", "CIF", "Adjudicatario"]) + '</table:table-row>'
    '<table:table-row><table:table-cell/>' + _texto_ods("2019/000001") +
    '<table:table-cell office:value-type="string"><text:p>Taller de Pintacaras + Talleres de Circo</text:p>'
    '<text:p>Cabalgata Parque Miraflores. 5 de enero de 2019</text:p></table:table-cell>'
    '<table:table-cell office:value-type="date" office:date-value="2019-01-11T00:00:00"><text:p>11/01/19</text:p>'
    '</table:table-cell>' + _texto_ods("1 día") +
    '<table:table-cell office:value-type="float" office:value="605"><text:p>605,00 €</text:p></table:table-cell>'
    '<table:table-cell office:value-type="float" office:value="1"><text:p>1</text:p></table:table-cell>'
    + _texto_ods("G87606521") + _texto_ods("ASOCIACION DE KOLORES") + '</table:table-row>'
    '<table:table-row><table:table-cell/>' + _texto_ods("2018/002872") + _texto_ods("Reparación Cortalunas") +
    '<table:table-cell office:value-type="date" office:date-value="2019-01-11"/>'
    '<table:table-cell table:number-columns-spanned="2" office:value-type="string"><text:p>15 días</text:p>'
    '</table:table-cell><table:covered-table-cell/>'
    '<table:table-cell office:value-type="float" office:value="1"/>' + _texto_ods("A36633600") +
    '<table:table-cell office:value-type="string"><text:p>INCIPRESA<text:s text:c="2"/>S.A</text:p>'
    '</table:table-cell></table:table-row>'
    '<table:table-row table:number-rows-repeated="1048000"><table:table-cell table:number-columns-repeated="1024"/>'
    '</table:table-row></table:table>')

CABECERA_LEGANES = ["NIF\t", "Adjudicatario\t", "Expediente\t", "Objeto del contrato\t", "Fecha contrato\t",
                    "Precio sin impuestos\t", "Precio con impuestos", "Duración contrato"]


def _leganes(periodo, anio):
    return _xlsx({"Datos": [
        ["INFORME MENSUAL DE CONTRATOS MENORES"], ["ENTIDAD:", "Junta de Gobierno del Ayuntamiento de Leganés"],
        ["AÑO:", anio], ["MENSUAL", periodo], [f"Informe generado el 31-07-{anio}"], CABECERA_LEGANES,
        ["B82281395", "DESCALZOS PRODUCCIONES SL", f"{anio}/0183", "Representación del espectáculo “Pijama para seis”",
         f"30-JUL-{anio}", "2.560", "3.097,6 €", 1],
        ["47542054L", "GONZALO CARDONE DOMINGUEZ", f"{anio}/0172", "Representación del espectáculo “Un mundo perfecto”",
         f"20-JUL-{anio}", "1.400", "1.694 €", 1]]})


LEGANES_2015 = ('NIF/CIF Tercero;Razón Social Tercero;Procedimiento de Contratación;Num. Expediente;'
                'Titulo Expediente;Importe Total Presup.Licit.;Organismo\r\n'
                'G84682830;ASOCIACION CULTURAL FLAMENCA JONDO;Contrato Menor;0001/2015;Concurso Cante Silla de Oro;'
                '18.000,00 €;JMD La Fortuna\r\n'
                '52795437H;CARMEN SERRANO AGUILELLA;Contrato Menor;0002/2015;'
                '"Teatro Inf ""La fábrica juguetes defect""-10ene14";1.700,00 €;Cultura\r\n').encode("cp1252")

SANTA_CRUZ_2024 = (
    'LISTADO DE CONTRATOS ENTRE EL 01/01/2024 Y EL 31/12/2024;;;;;;;;;;;;\r\n'
    'CONTRATO;EXPEDIENTE;FECHA ALTA;UNIDAD DE CONTRATACIÓN;OBJETO DEL CONTRATO;TIPO;SUBTIPO;PROCEDIMIENTO;'
    'DESCRIPCION_ESTADO;PRESUPUESTO BASE SIN IMPUESTOS;IMPORTE DE ADJUDICACION;IMPORTE DE ADJ SIN IMPUESTOS;\r\n'
    'MEN2024000002;1/2024/MC-MEN;02/01/2024 00:00;O.A. de Fiestas y Actividades Recreativas;Suministro mediante '
    'arrendamiento plataforma giratoria para la gala de elección de la Reina del Carnaval 2024;Suministro;Alquiler;'
    'Contrato menor;Adjudicado;2225;2380,75;2225;\r\n'
    'MEN2024000003;2/2024/MC-MEN;02/01/2024 00:00;Organismo Autónomo de Cultura;Adquisición de diferentes alfombras '
    'para el Teatro Guimerá.;Suministro;Adquisición;Contrato menor;Adjudicado;11214,95;6067,51;6067,51;\r\n'
).encode("cp1252")

CORDOBA_CSV = ('nid;NIF;Adjudicatario;Objeto;Duracion;ImporteLicitacionSinIVA;ImporteDeAdjudicacionSinIVA;'
               'ImporteDeAdjudicacionIVA;InstrumDePublic;NLicitad;Fecha;Decreto;Delegacion\r\n'
               '166;B95604021;CORREO INTELIGENTE POSTAL, S.L.;Correspondencia Ordinaria, notificaciones, etc;'
               '2 meses;14999,17;14999,17;18149;PCE;1;14/01/2021 00:00:00;171-21;GMU\r\n').encode("utf-8")

CORDOBA_ODS = _ods(
    '<table:table table:name="TERCER_TRIMESTRE"><table:table-row>'
    + _texto_ods("RELACIÓN CONTRATOS MENORES TERCER TRIMESTRE 2024") + '</table:table-row><table:table-row>'
    + "".join(_texto_ods(t) for t in ["NIF", "Adjudicatario", "Objeto", "Duración", "Importe licitación sin IVA",
                                      "Importe de adjudicac."]) + '</table:table-row><table:table-row>'
    + _texto_ods("B19665173") + _texto_ods("PROVISIONA INGENIERIA SL") + _texto_ods("SERVICIO DE ASISTENCIA")
    + _texto_ods("26 de junio al  18 de agosto")
    + '<table:table-cell office:value-type="float" office:value="13112.04"/>'
      '<table:table-cell office:value-type="currency" office:value="12213.87"/></table:table-row></table:table>')


# ---------------------------------------------------------------------------
# Portales simulados
# ---------------------------------------------------------------------------

V = "https://www.valladolid.gob.es"
RUTA_VLL = "/es/perfil-contratante/contratos-menores-volumen-contratacion-tipo-procedimiento"
VLL_1T = (f"{V}{RUTA_VLL}/ano-2025/ayuntamiento-valladolid.ficheros/"
          "1067400-CONTRATACION%20PRIMER%20TRIMESTE%20EJERCICIO%202025%20AYUNTAMIENTO%20VALLADOLID.xlsx")
VLL_3T = (f"{V}{RUTA_VLL}/ano-2025/ayuntamiento-valladolid.ficheros/1135791-CONTRATACION%20PRIMER%2CSEGUNDO"
          "%20Y%20TERCER%20TRIMES%20EJERCICIO%202025%20AYUNTAMIENTO%20VALLADOLID.xlsx")
VLL_2024 = f"{V}{RUTA_VLL}/ano-2024.ficheros/1051832-CONTRATACION%20EJERCICIO%202024%20AYUNTAMIENTO%20VALLADOLID.xlsx"
VLL_PDF = f"{V}{RUTA_VLL}/ano-2025/fundacion-municipal-deportes.ficheros/1093627-SEGUNDO%20TRIMESTRE%20%202025.pdf"

UPLOADS = "https://transparencia.ayto-fuenlabrada.es/wp-content/uploads/"
FUE_2026 = UPLOADS + "2026/08/Contratos-menores-AYTO-Y-OO.AA_.-2T-2026.xls"
FUE_NG = UPLOADS + "2026/02/NEXT-GENERATION-2021-2025-AYTO.xlsx"
FUE_2019 = UPLOADS + "2019/05/AYTO-1T-2019.ods"

L = "https://www.leganes.org"
LEG_JULIO = f"{L}/documents/113177/417685/Informe_menores_julio_2026.xlsx/c62bff07?t=1790159490114"
LEG_JULIO_AGOSTO = f"{L}/documents/113177/250601/Informe+julio+y+agosto+menores.xlsx/7b9d53f2?t=1774353904774"
LEG_2015 = f"{L}/documents/113177/250360/0_60474_1.csv/400e0946?t=1774353945469"
LEG_PDF = f"{L}/documents/113177/250360/0_53439_1.pdf/32700a77?t=1774353943302"
LEG_AGOSTO = f"{L}/documents/131847/231195/Informe_menores_agosto_2026.xlsx/d1f5265f?t=1790159511328"

SC = "https://www.santacruzdetenerife.es/gobiernoabierto/transparencia/fileadmin/user_upload/web/Transparencia/"
SC_2024 = SC + "contratos/CONTRATOS_MENORES_2024_CSV.csv"
SC_RESUMEN = SC + "contratos/Resumen_menores_2024_CSV.csv"
SC_PDF = SC + "contratos/CONTRATOS_MENORES_2024.pdf"

MAL = "https://datosabiertos.malaga.eu/recursos/presidencia/"
COR = "https://datosabiertos.cordoba.es/dataset/x/resource/"


def _a(href, texto):
    return f'<a target="_blank" href="{href}" class="link-icon w-100" download> <i class="fa"></i> {texto} </a>'


def pagina_valladolid(enlaces):
    return ("<html><body>" + "".join(_a(h, t) for h, t in enlaces) + "</body></html>").encode("utf-8")


def pagina_fuenlabrada(filas):
    """Tabla del listado de Fuenlabrada (con el enlace 'Relación' comentado, como la real)."""
    cuerpo = "".join(
        f'<tr><td data-head="Fecha">\n  {fecha}  </td><td data-head="Nombre">{nombre}</td>'
        '<!--\n<td data-head="Documento">\n-->\n<!--\n<a href="" class="btn" download> Relación</a>\n-->'
        f'<td data-head="Documento">\n<a href="{url}" class="btn"> Descargar</a>\n</td></tr>'
        for fecha, nombre, url in filas)
    return ("<html><body><table><thead><tr><th>Fecha</th><th>Nombre</th><th>Documento</th></tr></thead>"
            f"<tbody>{cuerpo}</tbody></table></body></html>").encode("utf-8")


FILAS_FUENLABRADA_1 = [("25-08-2026", "2º Trimestre 2026", FUE_2026),
                       ("31-12-2025", "Información complementaria financiación UE-Next Generation 2021-2025", FUE_NG)]
FILAS_FUENLABRADA_2 = [("31-12-2025", "Información complementaria financiación UE-Next Generation 2021-2025", FUE_NG),
                       ("23-05-2019", "1er Trimestre 2019 – Ayuntamiento", FUE_2019)]

ENLACES_LEGANES = [(LEG_AGOSTO, "Menores agosto 2026"), (LEG_JULIO, "Menores julio 2026"),
                   (LEG_JULIO_AGOSTO, "Menores julio y agosto 2024"), (LEG_PDF, "Menores Septiembre 2023"),
                   (LEG_PDF, "09 Menores Septiembre 2016"), (LEG_2015, "Listado Contratos Menores 2015")]


def pagina_leganes(enlaces):
    return ('<html><body><div class="d-flex flex-wrap">'
            + "".join(_a(url.replace(L, ""), texto) for url, texto in enlaces)
            + '</div><a href="/web/transparencia/buscador">Buscador</a></body></html>').encode("utf-8")


def _recurso_ckan(nombre, formato, url, identificador="r"):
    return {"id": identificador, "name": nombre, "format": formato, "url": url}


def paquetes_malaga():
    return [
        {"id": "1", "name": "contratos-menores-2o-trimestre-2026-ayuntamiento-de-malaga",
         "title": "Contratos Menores 2o Trimestre 2026 - Ayuntamiento de Málaga", "resources": [
             _recurso_ckan("pdf", "PDF", MAL + "contratacion2026/2TRMENORES.pdf"),
             _recurso_ckan("xlsx", "XLSX", MAL + "contratacion2026/2TRMENORES.xlsx"),
             _recurso_ckan("ods", "ODS", MAL + "contratacion2026/2TRMENORES.ods")]},
        {"id": "2", "name": "contratos-menores-4-trimestre-2016-cemi",
         "title": "Contratos menores 4 trimestre 2016 - CEMI", "resources": [
             _recurso_ckan("pdf", "PDF", MAL + "cemicontratacion2016/CONTRATOS_MENORES_4TRIMESTRE_2016.pdf"),
             _recurso_ckan("xls", "XLS", MAL + "cemicontratacion2016/CONTRATOS_MENORES_4TRIMESTRE_2016.xls")]},
        {"id": "3", "name": "contratos-menores-3-trimestre-2016-ayuntamiento-de-malaga",
         "title": "Contratos menores 3 trimestre 2016 - Ayuntamiento de Málaga", "resources": [
             _recurso_ckan("pdf", "PDF", MAL + "contratacion2016/CONTRATOS_MENORES_3TRIMESTRE_2016.pdf")]},
        {"id": "4", "name": "informacion-sig-mapa-estrategico-de-ruido-de-malaga-zonas-tranquilas-indice-lnoche",
         "title": "Información SIG Mapa Estratégico de Ruido de Málaga - Zonas tranquilas Índice Lnoche",
         "resources": [_recurso_ckan("csv", "CSV", MAL + "ruido/lnoche.csv")]},
    ]


def paquetes_cordoba():
    return [
        {"id": "1", "name": "contratacion-administrativa-contratos-menores",
         "title": "Contratación administrativa - Contratos menores", "resources": [
             _recurso_ckan("AYUNCORDOBA_2025_CONTRATOS_MENORES_4T", "PDF",
                           COR + "a/download/ayuncordoba_2025_contratos_menores_4t.pdf"),
             _recurso_ckan("AYUNCORDOBA_2024_CONTRATOS_MENORES_3T", "PDF",
                           COR + "b/download/ayuncordoba_2024_contratos_menores_3t.pdf"),
             _recurso_ckan("AYUNCORDOBA_2024_CONTRATOS_MENORES_3T", "XLS",
                           COR + "c/download/ayuncordoba_2024_contratos_menores_3t.ods")]},
        {"id": "2", "name": "contratos-menores", "title": "Contratos menores", "resources": [
            _recurso_ckan("Contratos menores (VISUALIZACIÓN)", "CSV", COR + "d/download/contratos_menores_2020.csv")]},
        {"id": "3", "name": "jgl-acuerdos-2021",
         "title": "Junta de Gobierno Local del Ayuntamiento de Córdoba - Acuerdos 2021", "resources": [
             _recurso_ckan("a.-jgl-n-744-21.-toma-conoc.-contratos-menores", "PDF", COR + "e/download/jgl.pdf")]},
    ]


class FakeResponse:
    def __init__(self, status=200, body=b"", json_data=None):
        self.status_code = status
        self._body = body
        self._json = json_data
        self.headers = {}

    def json(self):
        return self._json if self._json is not None else json.loads(self._body.decode("utf-8"))

    @property
    def text(self):
        return self._body.decode("utf-8", "replace")

    def iter_content(self, chunk_size=8192):
        yield self._body

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False


class FakePortal:
    """urls: url completa -> bytes | código HTTP; ckan: url de package_search ->
    lista de paquetes | código HTTP."""

    def __init__(self):
        self.urls = {}
        self.ckan = {}
        self.llamadas = []

    def get(self, url, params=None, headers=None, timeout=None, stream=False):
        self.llamadas.append((url, dict(params or {})))
        if url in self.ckan:
            paquetes = self.ckan[url]
            if isinstance(paquetes, int):
                return FakeResponse(status=paquetes)
            ini, n = int(params["start"]), int(params["rows"])
            return FakeResponse(json_data={"success": True, "result": {"count": len(paquetes),
                                                                       "results": paquetes[ini:ini + n]}})
        cuerpo = self.urls.get(url, 404)
        if isinstance(cuerpo, int):
            return FakeResponse(status=cuerpo, body=b"<html><body>Not Found</body></html>")
        return FakeResponse(body=cuerpo)

    def pedidas(self, url):
        return sum(1 for u, _ in self.llamadas if u == url)

    def pedidas_de(self, prefijo):
        return [u for u, _ in self.llamadas if u.startswith(prefijo)]


def _publicar_todo(portal):
    portal.urls[M.URL_GIJON] = _gijon(REG_A, REG_B, REG_C)
    for anio in range(M.VIGO_DESDE, ANIO + 1):
        portal.urls[M.URL_VIGO.format(aa=f"{anio % 100:02d}")] = VIGO_2019 if anio == 2019 else _vigo(anio)
    portal.urls["https://datos.vigo.org/data/sector-publico/contratos-menores-2018.csv"] = _vigo(2018)
    portal.ckan[M.URL_CKAN_VIGO] = [{"id": "v", "name": "contratos_menores-18", "title": "Contratos menores do 2018",
                                     "resources": [
                                         _recurso_ckan("csv", "CSV", "https://datos.vigo.org/data/sector-publico/"
                                                                     "contratos-menores-2018.csv"),
                                         _recurso_ckan("csv", "CSV", M.URL_VIGO.format(aa="25"))]}]
    portal.urls[M.URL_VALLADOLID] = pagina_valladolid(
        [(f"{RUTA_VLL}/ano-2024", "Año 2024"), (f"{RUTA_VLL}/ano-2025", "Año 2025"),
         (f"{RUTA_VLL}.nodos,10,10", "Siguiente")])
    portal.urls[f"{V}{RUTA_VLL}/ano-2024"] = pagina_valladolid([(VLL_2024.replace(V, ""), "xlsx")])
    portal.urls[f"{V}{RUTA_VLL}/ano-2025"] = pagina_valladolid(
        [(f"{RUTA_VLL}/ano-2025", "Año 2025"), (f"{RUTA_VLL}/ano-2025/ayuntamiento-valladolid", "Ayuntamiento"),
         (f"{RUTA_VLL}/ano-2025/fundacion-municipal-deportes", "FMD")])
    portal.urls[f"{V}{RUTA_VLL}/ano-2025/ayuntamiento-valladolid"] = pagina_valladolid(
        [(VLL_1T.replace(V, ""), "xlsx"), (VLL_3T.replace(V, ""), "xlsx")])
    portal.urls[f"{V}{RUTA_VLL}/ano-2025/fundacion-municipal-deportes"] = pagina_valladolid(
        [(VLL_PDF.replace(V, ""), "pdf")])
    portal.urls[VLL_1T] = portal.urls[VLL_3T] = _valladolid(2025)
    portal.urls[VLL_2024] = _valladolid(2024)
    portal.urls[VLL_PDF] = b"%PDF-1.5"
    portal.urls[M.URL_FUENLABRADA] = pagina_fuenlabrada(FILAS_FUENLABRADA_1)
    portal.urls[M.URL_FUENLABRADA + "page/2/"] = pagina_fuenlabrada(FILAS_FUENLABRADA_2)
    portal.urls[M.URL_FUENLABRADA + "page/3/"] = pagina_fuenlabrada([])
    portal.urls[FUE_2026] = _fuenlabrada_xlsx()
    portal.urls[FUE_NG] = _xlsx({"2021": [["Num. Expe.", "Importe"], ["2021/001072", 15125]]})
    portal.urls[FUE_2019] = FUENLABRADA_ODS
    portal.urls[M.URL_LEGANES] = pagina_leganes(ENLACES_LEGANES)
    portal.urls[LEG_JULIO] = _leganes("JULIO", 2026)
    portal.urls[LEG_JULIO_AGOSTO] = _leganes("JULIO Y AGOSTO", 2024)
    portal.urls[LEG_2015] = LEGANES_2015
    portal.urls[LEG_PDF] = b"%PDF-1.5"
    portal.urls[LEG_AGOSTO] = b"<!DOCTYPE html><html><head><title>Home - Ayuntamiento de Legan\xc3\xa9s</title></html>"
    portal.ckan[M.URL_CKAN_MALAGA] = paquetes_malaga()
    portal.urls[MAL + "contratacion2026/2TRMENORES.xlsx"] = _xlsx({"Informe de Contratos": [
        ["AYUNTAMIENTO DE MÁLAGA"], ["LISTADO CONTRATOS"],
        ["Nº de contrato", "Objeto del contrato", "Importe adjudicación  (IVA incluido)", "CIF", "Adjudicatario"],
        ["2026000024/MEN/AYTO/SERV", "OBJETO DEL CONTRATO _x000D_\nBiblioparque en parques infantiles", "1.050,00 €",
         "G93338226", "ASOCIACION CULTURAMA"]]})
    portal.urls[MAL + "cemicontratacion2016/CONTRATOS_MENORES_4TRIMESTRE_2016.xls"] = _xlsx({"Hoja1": [
        ["CONTRATOS MENORES TRAMITADOS  POR EL CENTRO MUNICIPAL DE INFORMÁTICA  EN EL CUARTO TRIMESTRE 2016"],
        ["EXPTE", "DESCRIPCION", "TOTAL", "TERCERO"],
        [62, "Seguros anuales vehículos R. Kangoo 4x4 1025DBF", 1549.71, "Axa Seguros Generales, S.A."]]})
    portal.ckan[M.URL_CKAN_CORDOBA] = paquetes_cordoba()
    portal.urls[COR + "c/download/ayuncordoba_2024_contratos_menores_3t.ods"] = CORDOBA_ODS
    portal.urls[COR + "d/download/contratos_menores_2020.csv"] = CORDOBA_CSV
    portal.urls[M.URL_SANTA_CRUZ] = ("<html><body>"
                                     + _a(SC_2024.replace("https://www.santacruzdetenerife.es", ""),
                                          "Contratos menores 2024 (CSV)")
                                     + _a(SC_RESUMEN, ".csv") + _a(SC_PDF, "Contratos menores 2024")
                                     + _a(SC + "contratos/CONTRATACION_PLAN_ANUAL_2026.pdf", ".pdf")
                                     + "</body></html>").encode("utf-8")
    portal.urls[SC_2024] = SANTA_CRUZ_2024
    portal.urls[SC_RESUMEN] = b"NUMERO DE CONTRATOS MENORES EN EL EJERCICIO 2024;;;;;;;;;1.267\r\n"


@pytest.fixture
def portal(monkeypatch):
    fake = FakePortal()
    monkeypatch.setattr(requests, "get", fake.get)
    monkeypatch.setattr(time, "sleep", lambda s: None)
    _publicar_todo(fake)
    return fake


def _ejecutar(salida, *args):
    return M.main(["--salida", str(salida), *args])


def _parquet(salida, municipio):
    return pd.read_parquet(salida / f"{municipio}_menores.parquet")


def _log(salida):
    return (salida / "raw" / "descarga_log.txt").read_text(encoding="utf-8")


def _manifiesto(salida):
    return json.loads((salida / "raw" / "_manifiesto.json").read_text(encoding="utf-8"))


def _v(serie):
    return [None if pd.isna(v) else v for v in serie]


# ---------------------------------------------------------------------------
# Un municipio cada vez: lo que se lee y cómo
# ---------------------------------------------------------------------------

def test_gijon_json_con_los_numeros_tal_cual(portal, tmp_path):
    assert _ejecutar(tmp_path, "--municipio", "gijon") == 0

    tabla = pq.read_table(tmp_path / "gijon_menores.parquet")
    assert all(pa.types.is_string(f.type) for f in tabla.schema if f.name != "_en_ultima_descarga")
    df = tabla.to_pandas()
    assert list(df.columns[:3]) == ["codigo_contrato", "lote", "numero_expediente"]
    assert "organo_de_contrataciÓn" in df.columns            # nombres tal como vienen
    assert df["presupuesto_base_de_licitaciÓn"].tolist() == ["49,59", "169,4", "17847,5"]
    assert df["precio_de_adjudicaciÓn"].tolist() == ["60", "169,4", "15427,5"]
    assert df["valor_estimado_sin_iva"].tolist() == ["0", "140", "14750.50"]   # no pasa por float
    assert _v(df["lote"]) == [None] * 3 and set(df["duraciÓn_del_contrato"]) == {" "}
    assert df["cif_adjudicatario"].tolist() == ["**35**63**", "0***882***", "A28242808"]
    assert df["_anio"].tolist() == ["2018", "2024", "2026"]                    # columna ejercicio
    assert set(df["_municipio"]) == {"Gijón"} and set(df["_codigo_ine"]) == {"33024"}
    assert set(df["_fuente"]) == {M.URL_GIJON} and df["_en_ultima_descarga"].all()
    raw = tmp_path / "raw" / "gijon" / "contratos_menores_adjudicados.json"
    assert raw.read_bytes() == _gijon(REG_A, REG_B, REG_C)                    # original tal cual
    assert portal.pedidas_de("https://www.valladolid") == []                  # solo el municipio pedido


def test_vigo_csv_anual_ckan_y_anios_no_publicados(portal, tmp_path):
    del portal.urls[M.URL_VIGO.format(aa=f"{ANIO % 100:02d}")]               # año en curso sin publicar
    assert _ejecutar(tmp_path, "--municipio", "vigo") == 0

    for anio in range(M.VIGO_DESDE, ANIO + 1):
        assert portal.pedidas(M.URL_VIGO.format(aa=f"{anio % 100:02d}")) == 1
    assert portal.pedidas(M.URL_VIGO.format(aa="25")) == 1                    # el CKAN no la duplica
    df = _parquet(tmp_path, "vigo")
    d19 = df[df["_anio"] == "2019"]
    assert len(d19) == 3 and d19["expediente"].isna().all()                  # 2019 no trae expediente
    assert d19["asunto"].iloc[1] == "CONTRATO MENOR DE SUBMINISTRO\nPAPELETAS ELECCIONES LOCALESD"
    assert d19["asunto"].iloc[2].endswith('DA REVISTA "A MOVIDA"')
    assert d19["importe"].tolist() == ["4513.0", "13157.3", "1452.0"]          # tal cual, sin convertir
    assert df.loc[df["_anio"] == "2025", "expediente"].tolist() == ["21362/77"]
    assert df.loc[df["_anio"] == "2018", "_archivo_origen"].tolist() == ["vigo/contratos-menores-2018.csv"]
    log = _log(tmp_path)
    assert f"vigo: {ANIO}" in log and "4 registros en 5 líneas" in log


def test_vigo_catalogo_ckan_caido_no_impide_la_descarga(portal, tmp_path):
    portal.ckan[M.URL_CKAN_VIGO] = 503
    assert _ejecutar(tmp_path, "--municipio", "vigo") == 0
    assert len(portal.pedidas_de(M.URL_CKAN_VIGO)) == M.INTENTOS_OPCIONAL == 2   # opcional: no espera minutos
    assert len(_parquet(tmp_path, "vigo")) == 3 + (ANIO - 2019)                  # las URL conocidas
    assert "vigo: catálogo CKAN" in _log(tmp_path)


def test_vigo_anio_confirmado_que_no_esta_es_un_error(portal, tmp_path):
    del portal.urls[M.URL_VIGO.format(aa="21")]
    assert _ejecutar(tmp_path, "--municipio", "vigo") == 1
    assert "vigo 2021: año publicado según las fuentes" in _log(tmp_path)


def test_valladolid_solo_hojas_de_operaciones_y_trimestres_acumulados(portal, tmp_path):
    assert _ejecutar(tmp_path, "--municipio", "valladolid") == 0

    df = _parquet(tmp_path, "valladolid")
    assert len(df) == 6 and set(df["_hoja"]) == {"OPERACIONES SICALWIN"}      # 2 operaciones x 3 ficheros
    assert "TERCEROS" not in df.columns and "01. Alcaldía" not in df.columns
    assert list(df.columns[:len(CABECERA_SICALWIN)]) == CABECERA_SICALWIN
    por_fichero = df.groupby("_archivo_origen")[["_anio", "_trimestre"]].first()
    assert por_fichero.to_dict("index") == {
        "valladolid/2024/1051832-CONTRATACION EJERCICIO 2024 AYUNTAMIENTO VALLADOLID.xlsx":
            {"_anio": "2024", "_trimestre": "1-4"},
        "valladolid/2025/1067400-CONTRATACION PRIMER TRIMESTE EJERCICIO 2025 AYUNTAMIENTO VALLADOLID.xlsx":
            {"_anio": "2025", "_trimestre": "1"},
        "valladolid/2025/1135791-CONTRATACION PRIMER,SEGUNDO Y TERCER TRIMES EJERCICIO 2025 AYUNTAMIENTO "
        "VALLADOLID.xlsx": {"_anio": "2025", "_trimestre": "1-3"}}
    fila = df[df["_anio"] == "2025"].iloc[1]
    assert (fila["Nº Operación"], fila["Fecha"], fila["Importe"]) == ("220250005020", "2025-03-21", "726.5")
    assert portal.pedidas(VLL_PDF) == 0                                       # el PDF no se baja
    inventario = json.loads((tmp_path / "raw" / "valladolid" / "_inventario.json").read_text(encoding="utf-8"))
    assert [(r["url"], r["estructurado"], r["motivo"]) for r in inventario if r["formato"] == "pdf"] == [
        (VLL_PDF, False, "PDF")]
    log = _log(tmp_path)
    assert "no se cargan 2 hojas de resumen o auxiliares" in log and "Hoja1 (2 filas), MENOR-AREA (3 filas)" in log
    assert "valladolid: 1 ficheros (PDF: 1)" in log


def test_fuenlabrada_listado_paginado_titulos_y_ods(portal, tmp_path):
    assert _ejecutar(tmp_path, "--municipio", "fuenlabrada") == 0

    assert portal.pedidas(M.URL_FUENLABRADA + "page/3/") == 1                # página sin filas: fin
    assert portal.pedidas(M.URL_FUENLABRADA + "page/4/") == 0
    assert portal.pedidas(FUE_NG) == 1                                        # en dos páginas, una vez
    df = _parquet(tmp_path, "fuenlabrada")
    assert df["_archivo_origen"].str.contains("NEXT-GENERATION").sum() == 1
    ayto = df[df["_hoja"] == "Menores_ayto_2trim_2026"]
    assert _v(ayto["Num. Expe."]) == ["2026/MSV/000911", None]
    assert _v(ayto["Adjudicatario"])[1] == "Total A VALLEKAS EDICIONES MUSICALES S. L    "   # subtotal
    assert set(ayto["_titulo_tabla"]) == {"CONTRATOS MENORES. AYUNTAMIENTO DE FUENLABRADA\nSEGUNDO TRIMESTRE DE 2026"}
    assert set(ayto["_titulo"]) == {"2º Trimestre 2026"}
    assert (ayto["_anio"].iloc[0], ayto["_trimestre"].iloc[0]) == ("2026", "2")
    ng = df[df["_archivo_origen"].str.contains("NEXT-GENERATION")]
    assert ng["_anio"].tolist() == ["2021-2025"] and ng["_trimestre"].isna().all()
    # ODS: párrafos con salto de línea, <text:s>, fecha, números y celdas combinadas
    ods = df[df["_archivo_origen"] == "fuenlabrada/2019/AYTO-1T-2019.ods"]
    assert ods["Num. Expe."].tolist() == ["2019/000001", "2018/002872"]
    assert ods["Título"].iloc[0] == ("Taller de Pintacaras + Talleres de Circo\n"
                                     "Cabalgata Parque Miraflores. 5 de enero de 2019")
    assert ods["Fecha Aprobación"].tolist() == ["2019-01-11", "2019-01-11"]
    assert _v(ods["Importe"]) == ["605", None] and ods["Num. Licitadores"].tolist() == ["1", "1"]
    assert ods["Plazo Ejec."].tolist() == ["1 día", "15 días"]
    assert ods["Adjudicatario"].tolist() == ["ASOCIACION DE KOLORES", "INCIPRESA  S.A"]
    assert set(ods["_titulo_tabla"]) == {"AYUNTAMIENTO DE FUENLABRADA"}
    assert (ods["_anio"].iloc[0], ods["_trimestre"].iloc[0]) == ("2019", "1")
    assert "hojas sin filas de datos: OTAF 4ºTR 2018 (NINGÚN  CONTRATO" in _log(tmp_path)


def test_leganes_periodo_del_texto_del_enlace_y_enlace_roto(portal, tmp_path):
    assert _ejecutar(tmp_path, "--municipio", "leganes") == 1                # agosto 2026 da la portada

    df = _parquet(tmp_path, "leganes")
    julio = df[df["_titulo"] == "Menores julio 2026"]
    assert (julio["_anio"].iloc[0], julio["_mes"].iloc[0]) == ("2026", "07")
    assert julio["NIF\t"].tolist() == ["B82281395", "47542054L"]            # cabecera con tabulador
    assert "ENTIDAD: Junta de Gobierno del Ayuntamiento de Leganés" in julio["_titulo_tabla"].iloc[0]
    assert set(df.loc[df["_titulo"] == "Menores julio y agosto 2024", "_mes"]) == {"07-08"}
    d15 = df[df["_anio"] == "2015"]
    assert d15["Importe Total Presup.Licit."].tolist() == ["18.000,00 €", "1.700,00 €"]      # cp1252
    assert d15["Titulo Expediente"].iloc[1] == 'Teatro Inf "La fábrica juguetes defect"-10ene14'
    assert d15["_archivo_origen"].tolist() == ["leganes/2015/0_60474_1.csv"] * 2
    assert portal.pedidas(LEG_PDF) == 0
    log = _log(tmp_path)
    assert "leganes/2026/Informe_menores_agosto_2026.xlsx: la respuesta es HTML" in log
    assert "se enlaza como 'Menores Septiembre 2023' y como '09 Menores Septiembre 2016'" in log
    assert not (tmp_path / "raw" / "leganes" / "2026" / "Informe_menores_agosto_2026.xlsx").exists()


def test_leganes_url_nueva_del_mismo_fichero_es_una_version_nueva(portal, tmp_path):
    del portal.urls[LEG_AGOSTO]
    portal.urls[M.URL_LEGANES] = pagina_leganes([(LEG_JULIO_AGOSTO, "Menores julio y agosto 2024")])
    _ejecutar(tmp_path, "--municipio", "leganes")
    SLEEP_REAL(1.1)
    nueva = LEG_JULIO_AGOSTO.replace("t=1774353904774", "t=1790000000000")   # Liferay: otra ?t=
    portal.urls[M.URL_LEGANES] = pagina_leganes([(nueva, "Menores julio y agosto 2024")])
    portal.urls[nueva] = _leganes("JULIO Y AGOSTO (corregido)", 2024)
    assert _ejecutar(tmp_path, "--municipio", "leganes") == 0               # 2024 es año cerrado

    assert portal.pedidas(nueva) == 1
    raw = tmp_path / "raw" / "leganes" / "2024" / "Informe+julio+y+agosto+menores.xlsx"
    assert len(M.versiones(raw)) == 2
    assert _manifiesto(tmp_path)["leganes/2024/Informe+julio+y+agosto+menores.xlsx"]["url"] == nueva
    df = _parquet(tmp_path, "leganes")
    assert len(df) == 2 and df["_en_ultima_descarga"].all()                 # mismas filas: no se duplican


def test_cada_url_conserva_su_ruta_local(portal, tmp_path):
    """Una URL ya descargada sigue en su ruta aunque cambie el texto del enlace
    (y con él el periodo), y un fichero nuevo con el mismo nombre que se enlaza
    antes no la ocupa: lleva un hash de su URL."""
    rel = "leganes/2024/Informe+julio+y+agosto+menores.xlsx"
    portal.urls[M.URL_LEGANES] = pagina_leganes([(LEG_JULIO_AGOSTO, "Menores julio y agosto 2024")])
    assert _ejecutar(tmp_path, "--municipio", "leganes") == 0
    SLEEP_REAL(1.1)
    otro = f"{L}/documents/113177/999999/Informe+julio+y+agosto+menores.xlsx/aa11?t=1"
    portal.urls[otro] = _leganes("JULIO Y AGOSTO (OO.AA.)", 2024)
    portal.urls[M.URL_LEGANES] = pagina_leganes([(otro, "Menores julio y agosto 2024 OO.AA."),
                                                 (LEG_JULIO_AGOSTO, "Menores julio y agosto")])   # sin año
    assert _ejecutar(tmp_path, "--municipio", "leganes") == 0

    manifiesto = _manifiesto(tmp_path)
    assert manifiesto[rel]["url"] == LEG_JULIO_AGOSTO and manifiesto[rel]["publicado"]
    [nuevo] = [r for r in manifiesto if r != rel]
    assert nuevo.startswith("leganes/2024/Informe+julio+y+agosto+menores__") and manifiesto[nuevo]["url"] == otro
    df = _parquet(tmp_path, "leganes")
    assert len(df) == 4 and df["_en_ultima_descarga"].all()                 # nada duplicado ni retirado


def test_malaga_ckan_formato_preferido_y_periodo_del_titulo(portal, tmp_path, monkeypatch):
    monkeypatch.setattr(M, "FILAS_CKAN", 1)
    assert _ejecutar(tmp_path, "--municipio", "malaga") == 0

    assert len(portal.pedidas_de(M.URL_CKAN_MALAGA)) == 4                     # paginado
    assert portal.pedidas(MAL + "contratacion2026/2TRMENORES.xlsx") == 1
    for url in ["contratacion2026/2TRMENORES.pdf", "contratacion2026/2TRMENORES.ods", "ruido/lnoche.csv",
                "cemicontratacion2016/CONTRATOS_MENORES_4TRIMESTRE_2016.pdf"]:
        assert portal.pedidas(MAL + url) == 0
    df = _parquet(tmp_path, "malaga")
    ayto = df[df["_titulo"] == "Contratos Menores 2o Trimestre 2026 - Ayuntamiento de Málaga"]
    assert (ayto["_anio"].iloc[0], ayto["_trimestre"].iloc[0]) == ("2026", "2")
    assert ayto["Objeto del contrato"].iloc[0] == "OBJETO DEL CONTRATO \r\nBiblioparque en parques infantiles"
    assert ayto["_titulo_tabla"].iloc[0] == "AYUNTAMIENTO DE MÁLAGA | LISTADO CONTRATOS"
    cemi = df[df["_titulo"] == "Contratos menores 4 trimestre 2016 - CEMI"]
    assert cemi[["EXPTE", "TOTAL", "_anio", "_trimestre"]].values.tolist() == [["62", "1549.71", "2016", "4"]]
    log = _log(tmp_path)
    assert ("malaga: 4 ficheros (PDF: 1, mismo fichero en ODS (se baja el XLSX): 1, mismo fichero en PDF "
            "(se baja el XLS): 1, mismo fichero en PDF (se baja el XLSX): 1)") in log


def test_cordoba_ckan_grupos_por_nombre_de_recurso(portal, tmp_path):
    assert _ejecutar(tmp_path, "--municipio", "cordoba") == 0

    df = _parquet(tmp_path, "cordoba")
    ods = df[df["_archivo_origen"].str.endswith(".ods")]                     # declarado XLS: se lee como ODS
    assert ods[["NIF", "Importe licitación sin IVA", "Importe de adjudicac.", "_anio", "_trimestre"]].values.tolist() == [
        ["B19665173", "13112.04", "12213.87", "2024", "3"]]
    assert ods["Duración"].iloc[0] == "26 de junio al  18 de agosto"
    csv_ = df[df["_archivo_origen"] == "cordoba/contratos-menores/contratos_menores_2020.csv"]
    assert csv_["ImporteLicitacionSinIVA"].tolist() == ["14999,17"]
    assert csv_["_anio"].isna().all()               # el nombre dice 2020 pero el recurso no trae año
    assert portal.pedidas_de(COR + "a/") == portal.pedidas_de(COR + "b/") == portal.pedidas_de(COR + "e/") == []
    assert "cordoba: 2 ficheros (PDF: 1, mismo fichero en PDF (se baja el ODS): 1); años 2024-2025" in _log(tmp_path)


def test_santa_cruz_csv_con_fila_de_titulo(portal, tmp_path):
    assert _ejecutar(tmp_path, "--municipio", "santa_cruz_tenerife") == 0

    df = _parquet(tmp_path, "santa_cruz_tenerife")
    assert list(df.columns[:4]) == ["CONTRATO", "EXPEDIENTE", "FECHA ALTA", "UNIDAD DE CONTRATACIÓN"]
    assert "Unnamed: 12" not in df.columns                                   # separador final vacío
    assert df["IMPORTE DE ADJUDICACION"].tolist() == ["2380,75", "6067,51"]
    assert set(df["_titulo_tabla"]) == {"LISTADO DE CONTRATOS ENTRE EL 01/01/2024 Y EL 31/12/2024"}
    assert set(df["_anio"]) == {"2024"} and set(df["_codigo_ine"]) == {"38038"}
    assert portal.pedidas(SC_RESUMEN) == portal.pedidas(SC_PDF) == 0
    assert "santa_cruz_tenerife: 2 ficheros (PDF: 1, resumen (no son registros): 1)" in _log(tmp_path)


def test_xls_binario(tmp_path):
    xlwt = pytest.importorskip("xlwt")
    pytest.importorskip("xlrd")
    libro = xlwt.Workbook()
    hoja = libro.add_sheet("PMD 4ºTR 2018")
    hoja.write(0, 0, "PATRONATO DE DEPORTES")
    for c, v in enumerate(["Num. Expe.", "Título", "Fecha Aprobación", "Importe"]):
        hoja.write(1, c, v)
    hoja.write(2, 0, "2018/002198")
    hoja.write(2, 1, "Placas señalética para Polideportivos")
    hoja.write(2, 2, datetime(2018, 10, 1), xlwt.easyxf(num_format_str="DD/MM/YYYY"))
    hoja.write(2, 3, 10760.53)
    ruta = tmp_path / "4-CONTRATOS-MENORES-AYTO-Y-OOAA-4o-TRIMESTRE-2018.xls"
    libro.save(str(ruta))
    df, _ = M.leer_tabla(ruta)
    assert df[["Num. Expe.", "Fecha Aprobación", "Importe", "_hoja", "_titulo_tabla"]].values.tolist() == [
        ["2018/002198", "2018-10-01", "10760.53", "PMD 4ºTR 2018", "PATRONATO DE DEPORTES"]]


def test_cabecera_con_huecos_y_hoja_sin_cabecera_no_se_comen_el_primer_registro(tmp_path):
    """Málaga, 2T 2017: 4 rótulos sobre 8 columnas (celdas combinadas); 3T 2021,
    Hoja2: lista de códigos de área sin cabecera."""
    ruta = tmp_path / "CONTRATOS_MENORES_2TRIMESTRE_2017.xlsx"
    ruta.write_bytes(_xlsx({
        "C.M.2trim2017": [
            ["CONTRATOS MENORES TRAMITADOS EN EL SEGUNDO TRIMESTRE DE 2017"],
            ["Expte.", None, "Descripción", None, None, "Total", None, "Tercero"],
            [2017000395, 108, "SERVICIO DE DIFUSIÓN Y PUESTA EN VALOR DEL MUSEO DEL PATRIMONIO MUNICIPAL", "SV",
             17900, 21659, "01   3335 22609", "FACTORÍA DE ARTE Y DESARROLLO S.L.U."],
            [2017000416, 94, "SERVICIO ESTERILIZACIÓN PARA LOS ANIMALES EN ADOPCIÓN", "SV", 2999.99, 3629.99,
             "21   3115 22706", "EMILIO GARCIA-MONCLUS DEL CASTILLO"]],
        "Hoja2": [[19, "ALCALDÍA"], [22, "INFRAESTR. Y PROYECTOS"]]}))
    df, avisos = M.leer_tabla(ruta)
    cm = df[df["_hoja"] == "C.M.2trim2017"]
    assert cm["Expte."].tolist() == ["2017000395", "2017000416"]
    assert cm["Total"].tolist() == ["21659", "3629.99"] and cm["Unnamed: 4"].tolist() == ["17900", "2999.99"]
    assert set(cm["_titulo_tabla"]) == {"CONTRATOS MENORES TRAMITADOS EN EL SEGUNDO TRIMESTRE DE 2017"}
    codigos = df[df["_hoja"] == "Hoja2"]
    assert codigos[["columna_1", "columna_2"]].values.tolist() == [["19", "ALCALDÍA"], ["22", "INFRAESTR. Y PROYECTOS"]]
    assert any("[C.M.2trim2017]: 4 columnas con valores y sin nombre" in a for a in avisos)


# ---------------------------------------------------------------------------
# Sesgo del superviviente
# ---------------------------------------------------------------------------

def test_registro_retirado_y_modificado_se_conservan(portal, tmp_path):
    assert _ejecutar(tmp_path, "--municipio", "gijon") == 0
    SLEEP_REAL(1.1)          # la versión nueva tiene que llevar otra fecha (resolución: 1 s)
    reg_c2 = REG_C.replace('"15427,5"', '"15500"')
    portal.urls[M.URL_GIJON] = _gijon(REG_A, reg_c2)
    assert _ejecutar(tmp_path, "--municipio", "gijon") == 0

    df = _parquet(tmp_path, "gijon")
    assert df[["_anio", "precio_de_adjudicaciÓn", "_en_ultima_descarga"]].values.tolist() == [
        ["2018", "60", True], ["2024", "169,4", False], ["2026", "15427,5", False], ["2026", "15500", True]]
    assert len(M.versiones(tmp_path / "raw" / "gijon" / "contratos_menores_adjudicados.json")) == 2


@pytest.mark.parametrize("fallo", [503, b"<html><body>Mantenimiento</body></html>", _gijon()])
def test_descarga_fallida_no_pierde_nada(portal, tmp_path, fallo):
    assert _ejecutar(tmp_path, "--municipio", "gijon") == 0
    SLEEP_REAL(1.1)
    portal.urls[M.URL_GIJON] = fallo
    assert _ejecutar(tmp_path, "--municipio", "gijon") == (0 if isinstance(fallo, bytes) and b"contrato" in fallo
                                                           else 1)
    df = _parquet(tmp_path, "gijon")
    assert len(df) == 3 and df["_en_ultima_descarga"].all()                 # nada se marca como retirado
    raw = tmp_path / "raw" / "gijon" / "contratos_menores_adjudicados.json"
    if not (isinstance(fallo, bytes) and b"contrato" in fallo):
        assert raw.read_bytes() == _gijon(REG_A, REG_B, REG_C) and len(M.versiones(raw)) == 1
    else:                    # JSON sin registros: se guarda, pero no retira nada
        assert "no tiene registros que cargar; no se marca nada como retirado" in _log(tmp_path)


def test_fichero_que_deja_de_enlazarse_queda_retirado(portal, tmp_path):
    assert _ejecutar(tmp_path, "--municipio", "fuenlabrada") == 0
    SLEEP_REAL(1.1)
    portal.urls[M.URL_FUENLABRADA + "page/2/"] = pagina_fuenlabrada(FILAS_FUENLABRADA_2[:1])   # sin el de 2019
    assert _ejecutar(tmp_path, "--municipio", "fuenlabrada") == 0

    df = _parquet(tmp_path, "fuenlabrada")
    retirado = df["_archivo_origen"] == "fuenlabrada/2019/AYTO-1T-2019.ods"
    assert retirado.sum() == 2 and not df.loc[retirado, "_en_ultima_descarga"].any()
    assert df.loc[~retirado, "_en_ultima_descarga"].all()
    assert _manifiesto(tmp_path)["fuenlabrada/2019/AYTO-1T-2019.ods"]["publicado"] is False
    assert (tmp_path / "raw" / "fuenlabrada" / "2019" / "AYTO-1T-2019.ods").exists()
    assert "fuenlabrada/2019/AYTO-1T-2019.ods: el portal ya no lo enlaza" in _log(tmp_path)


@pytest.mark.parametrize("pagina", [503, pagina_fuenlabrada([])])
def test_listado_caido_o_vacio_no_retira_nada(portal, tmp_path, pagina):
    assert _ejecutar(tmp_path, "--municipio", "fuenlabrada") == 0
    SLEEP_REAL(1.1)
    portal.urls[M.URL_FUENLABRADA] = pagina
    assert _ejecutar(tmp_path, "--municipio", "fuenlabrada") == 1
    df = _parquet(tmp_path, "fuenlabrada")
    assert df["_en_ultima_descarga"].all()
    assert all(e["publicado"] for e in _manifiesto(tmp_path).values())
    assert "fuenlabrada:" in _log(tmp_path).split("ERRORES")[-1]


def test_anio_sondeado_que_pasa_a_dar_404_queda_retirado(portal, tmp_path):
    url = M.URL_VIGO.format(aa=f"{ANIO % 100:02d}")
    assert _ejecutar(tmp_path, "--municipio", "vigo") == 0
    SLEEP_REAL(1.1)
    del portal.urls[url]
    assert _ejecutar(tmp_path, "--municipio", "vigo") == 0
    df = _parquet(tmp_path, "vigo")
    assert df.loc[df["_anio"] == str(ANIO), "_en_ultima_descarga"].tolist() == [False]
    assert df.loc[df["_anio"] != str(ANIO), "_en_ultima_descarga"].all()
    assert f"vigo vigo/contratos-menores-{ANIO % 100:02d}.csv: el portal ya no lo sirve" in _log(tmp_path)


# ---------------------------------------------------------------------------
# Refresco, CLI y rutas
# ---------------------------------------------------------------------------

def test_solo_se_vuelven_a_pedir_los_anios_recientes(portal, tmp_path):
    _ejecutar(tmp_path, "--municipio", "vigo")
    _ejecutar(tmp_path, "--municipio", "vigo")
    assert portal.pedidas(M.URL_VIGO.format(aa="19")) == 1
    assert portal.pedidas(M.URL_VIGO.format(aa=f"{ANIO % 100:02d}")) == 2
    assert portal.pedidas(M.URL_VIGO.format(aa=f"{(ANIO - 1) % 100:02d}")) == 2
    _ejecutar(tmp_path, "--municipio", "vigo", "--comprobar-todo")
    assert portal.pedidas(M.URL_VIGO.format(aa="19")) == 2
    assert not (tmp_path / "raw" / "vigo" / "_historico").exists()           # nada cambió


def test_cli_varios_municipios_y_solo_procesar(portal, tmp_path):
    assert _ejecutar(tmp_path, "--municipio", "santa_cruz_tenerife", "--municipio", "cordoba") == 0
    assert portal.pedidas(M.URL_SANTA_CRUZ) == 1 and portal.pedidas_de(M.URL_CKAN_CORDOBA)
    assert portal.pedidas(M.URL_GIJON) == 0 and portal.pedidas(M.URL_FUENLABRADA) == 0
    assert sorted(p.name for p in tmp_path.glob("*.parquet")) == [
        "cordoba_menores.parquet", "santa_cruz_tenerife_menores.parquet"]
    (tmp_path / "cordoba_menores.parquet").unlink()
    llamadas = len(portal.llamadas)
    assert _ejecutar(tmp_path, "--solo-procesar") == 0
    assert len(portal.llamadas) == llamadas                                  # sin descargas
    assert len(_parquet(tmp_path, "cordoba")) == 2                           # se regenera desde raw/
    with pytest.raises(SystemExit):
        _ejecutar(tmp_path, "--municipio", "madrid")


def test_cli(portal, tmp_path, monkeypatch):
    monkeypatch.setattr(sys, "argv", [str(SCRIPT), "--salida", str(tmp_path), "--municipio", "gijon"])
    with pytest.raises(SystemExit) as salida:
        runpy.run_path(str(SCRIPT), run_name="__main__")
    assert salida.value.code == 0
    assert (tmp_path / "gijon_menores.parquet").exists()


def test_salida_por_defecto_en_el_repo():
    assert M.SALIDA == REPO_ROOT / "municipios_menores"


def test_periodo_de_textos_reales():
    casos = {"1er Trimestre 2018": ("2018", "1", None), "4-TR-2022-AYTO-Y-OO.AA_..xls": ("2022", "4", None),
             "Menores julio y agosto 2024": ("2024", None, "07-08"), "03-04 Menores Marzo y Abril 2016":
                 ("2016", None, "03-04"), "12 Menores Diciembre 2016": ("2016", None, "12"),
             "Menores 2Trimestre 2023": ("2023", "2", None), "contratos menores ii trimestre 2026": ("2026", "2", None),
             "Contratos-menores-AYTO-y-OOAA-1-ER-TRIMESTRE-2021.xls": ("2021", "1", None),
             "14511_15315320181404.xlsx": (None, None, None)}
    assert {t: M.periodo_de_texto(t) for t in casos} == casos
