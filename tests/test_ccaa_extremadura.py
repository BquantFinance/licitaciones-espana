"""Tests offline de scripts/ccaa_extremadura.py.

El portal juntaex.es (buscador, páginas trimestrales del Registro de Contratos
y documentos del gestor documental) se simula con un ``requests.get`` falso;
los XLSX (y un XLS con xlwt) se generan en el propio test con las cabeceras de
los tres esquemas reales de contratos menores.
"""

import importlib.util
import io
import json
import runpy
import sys
import time
import uuid
from datetime import datetime
from pathlib import Path
from urllib.parse import parse_qs, quote_plus, urlparse

import openpyxl
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
import pytest
import requests

REPO_ROOT = Path(__file__).resolve().parents[1]
SCRIPT = REPO_ROOT / "scripts" / "ccaa_extremadura.py"
SLEEP_REAL = time.sleep


def _cargar():
    spec = importlib.util.spec_from_file_location("ccaa_extremadura", SCRIPT)
    modulo = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(modulo)
    return modulo


M = _cargar()
ANIO = datetime.now().year          # el script decide qué refrescar con el año real
CONFIRMADOS_REALES = list(M.TRIMESTRES_CONFIRMADOS)
TITULO_PAGINA = "Contratos e incidencias inscritas en el Registro de Contratos"
CARPETA = "/documents/77055/621084/"

# ---------------------------------------------------------------------------
# Los tres esquemas de los listados de contratos menores (cabeceras reales)
# ---------------------------------------------------------------------------

# 2022 y 1T 2023 (.xls): las celdas empiezan por un espacio
CAB_A = ["Nº Contrato", "Nº Expediente", "Consejería", "Órgano", "Órgano delega", "Tipo contrato", "Objeto",
         "CIF/NIF Contratista", "Denominación contratista", "Código CPV", "Importe contrato (con IVA)",
         "Fecha asiento"]
FILA_A = [" CM000001/22", " CS/02/1122000032/22/CM", " 50 - SES (Servicio Extremeño de Salud)",
          " 5006 - 08 GERENTE AREA DE MERIDA SES", " 1 - 01 DIRECTOR GERENTE SERVICIO EXTREMEÑO DE SALUD (SES)",
          " Contrato de Suministros", " PAÑOS DESECHABLES", " B06032965", " SANEX, S.L.  (SANIDAD EXTREMEÑA)",
          " 33141000", 9447.3, " 21/01/2022"]
# 2T-3T 2023
CAB_B = ["Número de registro de contrato", "Código del expediente", "RECO-Órg.Gestor",
         "Consejería/organismo/entidad", "RECO-Órgano de contratación", "Denominación órgano de contratación",
         "Objeto", "NIF", "Denominación contratista", "Importe Adjudicación"]
FILA_B = ["CM005815/23", "PRM/2023/0000034929", "15", "Consejería de Cultura, Turismo y Deportes", "1500600",
          None, "INSERCIÓN PUBLICITARIA", "B06386122", "C.M. EXTREMADURA PUBLICIDAD MU", 2178]
# Desde 4T 2023
CAB_C = ["Número de registro de contrato", "RECO-Órg.Gestor", "Consejería/Organismo/Entidad",
         "RECO-Órgano de contratación", "Denominación órgano de contratación", "Código del expediente",
         "Objeto del contrato", "Código CPV", "NIF", "Denominación Adjudicatario", "Fecha adjudicación",
         "Importe de adjudicación con impuestos"]
FILA_C = ["CM0000000001/2024", "11100", "SERVICIO EXTREMEÑO DE SALUD", "1110006",
          "06 GERENTE AREA DE DON BENITO-VILLANUEVA SES", "DB66058", "SUMINISTRO VARIOS MEDICAMENTOS",
          "33600000", "B86418787", "ABBVIE SPAIN SL", datetime(2024, 1, 2), 9783.1]

CAB_INCIDENCIAS = ["Número de registro de contrato", "Número de anotación", "Incidencia", "Fecha de Anotación"]


def _xlsx(hojas):
    libro = openpyxl.Workbook()
    libro.remove(libro.active)
    for nombre, filas in hojas.items():
        hoja = libro.create_sheet(nombre)
        for fila in filas:
            hoja.append(fila)
    datos = io.BytesIO()
    libro.save(datos)
    return datos.getvalue()


def _menores_c(*registros):
    return _xlsx({"Sheet1": [CAB_C] + [[r] + FILA_C[1:] for r in registros]})


def _incidencias(registro):
    return _xlsx({"Sheet1": [CAB_INCIDENCIAS, [registro, "00000000000000000546/2024", "Prórroga del contrato",
                                                datetime(2024, 6, 10)]]})


# ---------------------------------------------------------------------------
# Portal simulado
# ---------------------------------------------------------------------------

class FakeResponse:
    def __init__(self, status=200, body=b""):
        self.status_code = status
        self._body = body
        self.headers = {}

    def json(self):
        return json.loads(self._body)

    @property
    def text(self):
        return self._body.decode("utf-8", "replace")

    def iter_content(self, chunk_size=8192):
        yield self._body

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False


NO_ENCONTRADO = b"<html><head><title>Contenido no encontrado - Juntaex.es</title></head><body></body></html>"


def html_pagina(descripcion, documentos, titulo=TITULO_PAGINA):
    """Como las páginas reales: og:title y og:description, enlaces relativos a
    /documents/..., un icono del gestor documental y otros enlaces."""
    enlaces = "".join(f'<li><a href="{d["href"]}" target="_self">\n  {d["texto"]}\n</a></li>' for d in documentos)
    return (f'<html><head><title>{titulo} - Juntaex.es</title>'
            f'<meta property="og:title" content="{titulo}"/>'
            f'<meta property="og:description" content="{descripcion}"/></head><body>'
            '<a href="/w/otra-pagina">Otra página</a><img src="/documents/77055/110341/ico_organizador.png">'
            '<a href="/documents/77055/110341/logoJuntaEx.JPG/0d940219-8fd9-c7fe-c537-7ed2fb7dbf3f?t=1617807334949">'
            'Junta de Extremadura</a>'
            f'<h2>{descripcion}</h2><ul>{enlaces}</ul></body></html>').encode("utf-8")


def html_buscador(resultados, total):
    elementos = "".join(
        f'<li class="result-container"><h3 role="heading"><a href="https://www.juntaex.es/w/{slug}'
        f'?inheritRedirect=true">\n{titulo}\n</a></h3><p class="excerpt">{descripcion}</p></li>'
        for slug, titulo, descripcion in resultados)
    return (f'<html><body><a href="https://www.juntaex.es/w/noticia">Una noticia</a>'
            f'<div id="buscador-resultados">Se han encontrado <strong>{total}</strong> resultados'
            f'<ul>{elementos}</ul>Mostrando el intervalo 1 - {len(resultados)} de {total} resultados.'
            '</div></body></html>').encode("utf-8")


class FakePortal:
    """paginas: slug -> {descripcion, documentos, titulo} o código HTTP;
    ficheros: ruta del documento (sin ?t=, como sirve Liferay) -> bytes o código;
    buscador: slugs en el orden del buscador, o un código HTTP."""

    def __init__(self):
        self.paginas = {}
        self.ficheros = {}
        self.buscador = []
        self.llamadas = []

    def documento(self, nombre, contenido, texto="Ver documento", t=1700000000000, clave=None):
        clave = clave or str(uuid.uuid5(uuid.NAMESPACE_URL, nombre))
        ruta = f"{CARPETA}{quote_plus(nombre)}/{clave}"
        self.ficheros[ruta] = contenido
        return {"href": f"{ruta}?t={t}", "texto": texto, "nombre": nombre, "clave": clave, "ruta": ruta}

    def publicar(self, slug, descripcion, documentos, titulo=TITULO_PAGINA):
        self.paginas[slug] = {"descripcion": descripcion, "documentos": list(documentos), "titulo": titulo}
        return self.paginas[slug]

    def quitar_enlace(self, slug, nombre):
        self.paginas[slug]["documentos"] = [d for d in self.paginas[slug]["documentos"] if d["nombre"] != nombre]

    def get(self, url, params=None, headers=None, timeout=None, stream=False):
        if params:
            url = requests.Request("GET", url, params=params).prepare().url
        self.llamadas.append(url)
        partes = urlparse(url)
        if partes.path == "/buscador":
            if isinstance(self.buscador, int):
                return FakeResponse(status=self.buscador, body=b"<html>Error</html>")
            consulta = parse_qs(partes.query)
            delta, numero = int(consulta.get("delta", ["20"])[0]), int(consulta.get("start", ["1"])[0])
            resultados = []
            for slug in self.buscador[(numero - 1) * delta:numero * delta]:
                pagina = self.paginas.get(slug)
                pagina = pagina if isinstance(pagina, dict) else {"descripcion": "", "titulo": TITULO_PAGINA}
                resultados.append((slug, pagina["titulo"], pagina["descripcion"]))
            return FakeResponse(body=html_buscador(resultados, len(self.buscador)))
        if partes.path.startswith("/w/"):
            pagina = self.paginas.get(partes.path[3:], 404)
            if isinstance(pagina, int):
                return FakeResponse(status=pagina, body=NO_ENCONTRADO)
            return FakeResponse(body=html_pagina(pagina["descripcion"], pagina["documentos"], pagina["titulo"]))
        cuerpo = self.ficheros.get(partes.path, 404)
        if isinstance(cuerpo, int):
            return FakeResponse(status=cuerpo, body=NO_ENCONTRADO)
        return FakeResponse(body=cuerpo)

    def pedidas(self, fragmento):
        return sum(1 for u in self.llamadas if fragmento in u)

    def pedidas_documento(self, nombre):
        return self.pedidas(CARPETA + quote_plus(nombre) + "/")


# Trimestres confirmados en el portal simulado
CONFIRMADOS = [(2022, 1), (2022, 3), (2023, 2), (2023, 3), (2024, 1), (2025, 1), (2026, 2)]
# Orden del buscador: con 2 resultados por página, las dos páginas de slug
# irregular (2T 2026 y el resumen) solo salen en la segunda página
BUSCADOR = ["registro-contratos-1t-2024", "registro-contratos-2t-2023",
            "contratos-incidencias-inscritas-registro-contratos", "resumen-estadistico-de-contratos-2024",
            "registro-contratos-1t-2022"]


def _publicar_base(portal):
    d = portal.documento
    portal.publicar("registro-contratos-1t-2022", "1er. Trimestre de 2022.", [
        d("LISTADO CONTRATOS MAYORES 1T 2022.xls", _xlsx({"Masivo": [["Nº Contrato", "Objeto"], [" 0745/22", " OBRA"]]})),
        d("LISTADO CONTRATOS MENORES 1T 2022.xls", _xlsx({"Masiva Contratos Menores": [CAB_A, FILA_A]}),
          texto="Listado de Contratos Menores"),
        d("LISTADO MODIFICACIONES 1T 2022.xls",
          _xlsx({"Masiva Modificaciones": [["Nº Contrato", "Nº Modificación"], [" 0697/18", " 5"]]})),
    ])
    # Página índice con slug irregular que el buscador no devuelve
    portal.publicar("publicacion-registro-contratos", "3er. Trimestre de 2022.", [
        d("ListadoMENORES_3T_2022_CM1666594182411.xlsx",
          _xlsx({"Transparencia Masiva Contratos": [CAB_A, [" CM028236/22"] + FILA_A[1:]]}),
          texto="2.- Contratos menores 3T 2022"),
    ])
    portal.publicar("registro-contratos-2t-2023", "2º y 3er. Trimestres de 2023.", [
        d("CONTRATOS MENORES 2 y 3T 2023.xlsx", _xlsx({"Hoja1": [CAB_B, FILA_B]}),
          texto="2.- CONTRATOS MENORES 2º y 3º T. 2023"),
        d("INCIDENCIAS 2 y 3T 2023.xlsx", _incidencias("0014/21")),
    ])
    portal.publicar("registro-contratos-1t-2024", "1er. Trimestre de 2024.", [
        d("Listado de contratos Menores 1T 2024.xlsx", _menores_c("CM0000000001/2024", "CM0000000002/2024")),
        d("Listado de Incidencias 1T 2024.xlsx", _incidencias("1596/21")),
    ])
    # Solo se encuentra por los slugs candidatos (año con dos cifras)
    portal.publicar("registro-contratos-1t-25", "1er. Trimestre de 2025", [
        d("LISTADO CONTRATOS MENORES 1T 2025.xlsx", _menores_c("CM0000000001/2025")),
    ])
    portal.publicar("contratos-incidencias-inscritas-registro-contratos", "2º Trimestre de 2026", [
        d("LISTADO CONTRATOS MENORES 2T 26.xlsx", _menores_c("CM0000010108/2026")),
    ])
    portal.publicar("resumen-estadistico-de-contratos-2024", "Resumen estadístico del ejercicio 2024", [
        d("LISTADO POR CONTRATISTAS.pdf", b"%PDF-1.7\n%resumen\n"),
    ])
    portal.buscador = list(BUSCADOR)


@pytest.fixture
def portal(monkeypatch):
    fake = FakePortal()
    monkeypatch.setattr(requests, "get", fake.get)
    monkeypatch.setattr(time, "sleep", lambda s: None)
    monkeypatch.setattr(M, "POR_PAGINA_BUSCADOR", 2)
    monkeypatch.setattr(M, "PAGINAS_CONOCIDAS", {"publicacion-registro-contratos": "3T 2022"})
    monkeypatch.setattr(M, "TRIMESTRES_CONFIRMADOS", list(CONFIRMADOS))
    _publicar_base(fake)
    return fake


def _ejecutar(salida, *args):
    return M.main(["--salida", str(salida), *args])


def _log(salida):
    return (salida / "raw" / "descarga_log.txt").read_text(encoding="utf-8")


def _manifiesto(salida):
    return json.loads((salida / "raw" / "_manifiesto.json").read_text(encoding="utf-8"))


def _menores(salida):
    return pd.read_parquet(salida / "registro_contratos_menores.parquet")


def _v(serie):
    return [None if pd.isna(v) else v for v in serie]


# ---------------------------------------------------------------------------
# Descubrimiento: buscador, páginas conocidas y slugs candidatos
# ---------------------------------------------------------------------------

def test_buscador_se_recorre_entero(portal, tmp_path):
    assert _ejecutar(tmp_path) == 0
    # 5 resultados de 2 en 2: tres páginas de resultados
    assert portal.pedidas("/buscador?") == 3
    assert portal.pedidas("start=2") == 1 and portal.pedidas("start=3") == 1
    # Las dos páginas de slug irregular solo estaban en la segunda
    raw = tmp_path / "raw"
    assert (raw / "menores" / "menores_2026_2T.xlsx").exists()
    assert (raw / "resumen_estadistico" / "resumen_estadistico_2024_contratistas.pdf").exists()
    inventario = json.loads((raw / "paginas.json").read_text(encoding="utf-8"))
    assert inventario["https://www.juntaex.es/w/contratos-incidencias-inscritas-registro-contratos"][
        "trimestres"] == [2]
    # El enlace a una noticia del buscador no es del registro: no se visita
    assert portal.pedidas("/w/noticia") == 0


def test_pagina_conocida_que_el_buscador_no_devuelve(portal, tmp_path):
    assert _ejecutar(tmp_path) == 0
    assert portal.pedidas("/w/publicacion-registro-contratos") == 1
    df = _menores(tmp_path)
    fila = df[df["_archivo_origen"] == "menores/menores_2022_3T.xlsx"]
    assert fila["Nº Contrato"].tolist() == [" CM028236/22"]
    assert fila["_nombre_publicado"].tolist() == ["ListadoMENORES_3T_2022_CM1666594182411.xlsx"]
    assert fila["_pagina"].tolist() == ["https://www.juntaex.es/w/publicacion-registro-contratos"]


def test_slugs_candidatos_de_los_trimestres_sin_pagina(portal, tmp_path):
    assert _ejecutar(tmp_path) == 0
    # 1T 2025 no está en el buscador: se prueba -1t-2025 (404) y después -1t-25
    assert portal.pedidas("/w/registro-contratos-1t-2025") == 1
    assert portal.pedidas("/w/registro-contratos-1t-25") == 1
    assert (tmp_path / "raw" / "menores" / "menores_2025_1T.xlsx").exists()
    # Un trimestre cubierto (3T 2023 va en la página de 2T y 3T) no se prueba...
    assert portal.pedidas("/w/registro-contratos-3t-2023") == 0
    assert portal.pedidas("/w/registro-contratos-4t-2022") == 1      # ...uno sin página, sí
    # ...salvo con --comprobar-todo
    assert _ejecutar(tmp_path, "--comprobar-todo") == 0
    assert portal.pedidas("/w/registro-contratos-3t-2023") == 1
    assert portal.pedidas("/w/registro-contratos-3t-23") == 1


@pytest.mark.parametrize("nombre, tipo, anio, trimestres", [
    ("LISTADO CONTRATOS MENORES 1T 2022.xls", "menores", 2022, (1,)),
    ("Listado de contratos menores 2T 2022.xls", "menores", 2022, (2,)),
    ("ListadoMENORES_3T_2022_CM1666594182411.xlsx", "menores", 2022, (3,)),
    ("LISTADO CONTRATOS MENORES 1T 2023.xls", "menores", 2023, (1,)),
    ("CONTRATOS MENORES 2 y 3T 2023.xlsx", "menores", 2023, (2, 3)),
    ("LISTADO CONTRATOS MENORES 4 T 2023.xlsx", "menores", 2023, (4,)),
    ("Listado de contratos Menores 1T 2024.xlsx", "menores", 2024, (1,)),
    ("LISTADOS CONTRATOS MENORES 3º T 2024.xlsx", "menores", 2024, (3,)),
    ("LISTADO CONTRATOS MENORES 2º TRIMESTRE 2025.xlsx", "menores", 2025, (2,)),
    ("CONTRATOS MENORES 3ºT 2025.xlsx", "menores", 2025, (3,)),
    ("MENORES 4º T 2025.xlsx", "menores", 2025, (4,)),
    ("LISTADO CONTRATOS MENORES 1º T 2026.xlsx", "menores", 2026, (1,)),
    ("LISTADO CONTRATOS MENORES 2T 26.xlsx", "menores", 2026, (2,)),
    ("ListadoMAYORES_3T_2022_C1666594135676.xls", "mayores", 2022, (3,)),
    ("LISTADOS CONTRATOS MAYORES 4T 2022.xls", "mayores", 2022, (4,)),
    ("LISTADO DE CONTRATOS MAYORES 2º TRIMESTRE 2025.xlsx", "mayores", 2025, (2,)),
    ("MAYORES 4º T 2025.xlsx", "mayores", 2025, (4,)),
    ("INCIDENCIAS 2 y 3T 2023.xlsx", "incidencias", 2023, (2, 3)),
    ("LISTADO DE INCIDENCIAS 2T 26.xlsx", "incidencias", 2026, (2,)),
    ("LISTADO MODIFICACIONES 1T 2022.xls", "modificaciones", 2022, (1,)),
    ("Listado de Modificacioens de contratos 2T 2022.xls", "modificaciones", 2022, (2,)),
    ("ListadoModificaciones_3T_2022_1666594216208.xls", "modificaciones", 2022, (3,)),
    ("LISTADO AMPLIACIONES PLAZO 1T 2022.xls", "ampliaciones_plazo", 2022, (1,)),
    ("ListadoAumentoPlazo_3T_2022_1666594216208.xls", "ampliaciones_plazo", 2022, (3,)),
    ("LISTADO DE AMPLIACIÓN DE PLAZO 1T 2023.xls", "ampliaciones_plazo", 2023, (1,)),
    ("LISTADO PRÓRROGAS EJECUCIÓN 1T 2022.xls", "prorrogas", 2022, (1,)),
    ("ListadoProrrogasContrato_3T_2022_1666594216208.xls", "prorrogas", 2022, (3,)),
    ("LISTADO DE PRÓRROGAS DE CONTRATOS 1T 2023.xls", "prorrogas", 2023, (1,)),
    ("LISTADO RESOLUCIONES DE CONTRATOS 4T 2022.xls", "resoluciones", 2022, (4,)),
    ("ListadoResoluciones_3T_2022_1666594216208.xls", "resoluciones", 2022, (3,)),
    ("LISTADO POR CONTRATISTAS (RANKING 25 CONTRATISTAS) 2025.pdf", "resumen_estadistico", 2025, ()),
    ("VOLUMEN PRESUPUESTARIO CONTRATOS 2023 POR ÓRGANO.pdf", "resumen_estadistico", 2023, ()),
    ("Listado  por Procedimientos de Adjudicación (1).pdf", "resumen_estadistico", None, ()),
])
def test_clasificacion_de_nombres_publicados(nombre, tipo, anio, trimestres):
    assert M.tipo_documento(nombre) == tipo
    assert M.periodo(nombre) == (anio, trimestres)


@pytest.mark.parametrize("texto, esperado", [
    ("1er. Trimestre de 2026", (2026, (1,))),
    ("2º y 3er. Trimestres de 2023.", (2023, (2, 3))),
    ("4º. Trimestre de 2024", (2024, (4,))),
    ("3º Trimestre de 2025", (2025, (3,))),
    ("Resumen estadístico del ejercicio 2025", (2025, ())),
    ("registro-contratos-1t-23", (2023, (1,))),
    ("registro-contratos-3t-2024", (2024, (3,))),
    ("registro-contratos", (None, ())),
    ("Listado de Contratos Menores 3er. T. 2024", (2024, (3,))),
    ("2.- Contratos menores 3T 2022", (2022, (3,))),
])
def test_periodo_de_paginas_y_enlaces(texto, esperado):
    assert M.periodo(texto) == esperado


def test_el_periodo_que_falta_en_el_nombre_sale_de_la_pagina(portal, tmp_path):
    portal.publicar("registro-contratos-4t-2023", "4º Trimestre de 2023.", [
        portal.documento("Listado de contratos menores.xlsx", _menores_c("CM008231/23"))])
    portal.buscador.append("registro-contratos-4t-2023")
    assert _ejecutar(tmp_path) == 0
    df = _menores(tmp_path)
    fila = df[df["Número de registro de contrato"] == "CM008231/23"]
    assert fila[["_anio", "_trimestre", "_archivo_origen"]].values.tolist() == [
        ["2023", "4T", "menores/menores_2023_4T.xlsx"]]


# ---------------------------------------------------------------------------
# Parquet: tres esquemas, texto, .xls y .xlsx, filas de título
# ---------------------------------------------------------------------------

def test_tres_esquemas_en_un_parquet_como_texto(portal, tmp_path):
    assert _ejecutar(tmp_path) == 0
    tabla = pq.read_table(tmp_path / "registro_contratos_menores.parquet")
    assert all(pa.types.is_string(f.type) for f in tabla.schema if f.name != "_en_ultima_descarga")
    df = tabla.to_pandas()
    # Ninguna columna de ningún esquema se pierde (tal cual: 'Consejería/organismo/entidad' de
    # 2T-3T 2023 y 'Consejería/Organismo/Entidad' de 2024 son columnas distintas)
    for columna in CAB_A + CAB_B + CAB_C:
        assert columna in df.columns
    assert list(df.columns[:len(CAB_A)]) == CAB_A
    assert len(df) == 7
    a = df[df["_archivo_origen"] == "menores/menores_2022_1T.xls"].iloc[0]
    assert a["Nº Contrato"] == " CM000001/22" and a["CIF/NIF Contratista"] == " B06032965"
    assert a["Importe contrato (con IVA)"] == "9447.3" and pd.isna(a["NIF"])
    assert a["_hoja"] == "Masiva Contratos Menores" and a["_tipo"] == "menores"
    b = df[df["_trimestre"] == "2T-3T"].iloc[0]
    assert b["NIF"] == "B06386122" and b["Importe Adjudicación"] == "2178" and b["_anio"] == "2023"
    assert pd.isna(b["Denominación órgano de contratación"]) and pd.isna(b["Nº Contrato"])
    c = df[df["_archivo_origen"] == "menores/menores_2024_1T.xlsx"]
    assert c["Número de registro de contrato"].tolist() == ["CM0000000001/2024", "CM0000000002/2024"]
    assert set(c["Fecha adjudicación"]) == {"2024-01-02"} and set(c["Código CPV"]) == {"33600000"}
    assert set(df["_dataset"]) == {"registro_contratos_menores"}
    assert df["_en_ultima_descarga"].all()
    # Cada fila lleva la URL publicada (con ?t=) y la página que la enlaza
    assert set(c["_fuente"]) == {"https://www.juntaex.es" + portal.paginas["registro-contratos-1t-2024"][
        "documentos"][0]["href"]}
    assert set(c["_pagina"]) == {"https://www.juntaex.es/w/registro-contratos-1t-2024"}
    # Incidencias: los ficheros por tipo de 2022 y los listados únicos desde 2023, en una tabla
    inc = pd.read_parquet(tmp_path / "registro_contratos_incidencias.parquet")
    assert sorted(inc["_tipo"]) == ["incidencias", "incidencias", "modificaciones"]
    assert set(pd.read_parquet(tmp_path / "registro_contratos_mayores.parquet")["Nº Contrato"]) == {" 0745/22"}
    # El PDF del resumen se guarda pero no es una tabla
    assert sorted(p.name for p in tmp_path.glob("*.parquet")) == [
        "registro_contratos_incidencias.parquet", "registro_contratos_mayores.parquet",
        "registro_contratos_menores.parquet"]


def test_contrato_publicado_otra_vez_se_marca_sin_quitar_filas(portal, tmp_path):
    # Como el listado real de 4T 2023, que vuelve a publicar contratos de 1T y 2T-3T 2023
    portal.publicar("registro-contratos-4t-2023", "4º Trimestre de 2023.", [
        portal.documento("LISTADO CONTRATOS MENORES 4 T 2023.xlsx", _menores_c(
            "CM005815/23", "CM000001/22", "CM0000005815/2023", "CM008231/23", "CM008231/23"))])
    portal.buscador.append("registro-contratos-4t-2023")
    assert _ejecutar(tmp_path) == 0
    df = _menores(tmp_path)
    assert len(df) == 12                                              # no se quita ninguna fila
    cuarto = df[df["_archivo_origen"] == "menores/menores_2023_4T.xlsx"]
    assert _v(cuarto["_repetido_de"]) == ["menores/menores_2023_2T-3T.xlsx", "menores/menores_2022_1T.xls",
                                          None, None, None]
    assert df.loc[df["_archivo_origen"] != "menores/menores_2023_4T.xlsx", "_repetido_de"].isna().all()
    assert "2023 4T: 5 filas (2 ya estaban en un listado anterior (_repetido_de))" in _log(tmp_path)
    SLEEP_REAL(1.1)
    assert _ejecutar(tmp_path, "--comprobar-todo") == 0
    otra = _menores(tmp_path)
    assert _v(otra["_repetido_de"]) == _v(df["_repetido_de"]) and len(otra) == 12
    # Si el listado anterior deja de publicarse, la copia posterior ya no cuenta como repetida
    SLEEP_REAL(1.1)
    portal.quitar_enlace("registro-contratos-2t-2023", "CONTRATOS MENORES 2 y 3T 2023.xlsx")
    assert _ejecutar(tmp_path) == 0
    df = _menores(tmp_path)
    cuarto = df[df["_archivo_origen"] == "menores/menores_2023_4T.xlsx"]
    assert _v(cuarto["_repetido_de"]) == [None, "menores/menores_2022_1T.xls", None, None, None]
    vigentes = df[df["_en_ultima_descarga"] & df["_repetido_de"].isna()]
    assert vigentes["Número de registro de contrato"].tolist().count("CM005815/23") == 1


def test_xls_binario_y_xlsx(portal, tmp_path):
    xlwt = pytest.importorskip("xlwt")
    pytest.importorskip("xlrd")
    libro = xlwt.Workbook()
    hoja = libro.add_sheet("Masiva Contratos Menores")
    fecha = xlwt.easyxf(num_format_str="DD/MM/YYYY")
    for c, v in enumerate(["Nº Contrato", "CIF/NIF Contratista", "Importe contrato (con IVA)", "Fecha"]):
        hoja.write(0, c, v)
    for f, (contrato, nif, importe) in enumerate([(" CM044386/22", " B10024941", 6364.6),
                                                   (" CM044387/22", " 28954758N", 100)], start=1):
        hoja.write(f, 0, contrato)
        hoja.write(f, 1, nif)
        hoja.write(f, 2, importe)
    hoja.write(1, 3, datetime(2022, 9, 5), fecha)
    datos = io.BytesIO()
    libro.save(datos)
    portal.publicar("registro-contratos-4t-2022", "4º Trimestre de 2022.", [
        portal.documento("LISTADO CONTRATOS MENORES 4T 2022.xls", datos.getvalue())])
    assert _ejecutar(tmp_path) == 0

    raw = tmp_path / "raw" / "menores" / "menores_2022_4T.xls"
    assert raw.read_bytes() == datos.getvalue()                       # el original, tal cual
    df = _menores(tmp_path)
    xls = df[df["_archivo_origen"] == "menores/menores_2022_4T.xls"]
    assert xls["Nº Contrato"].tolist() == [" CM044386/22", " CM044387/22"]
    assert xls["CIF/NIF Contratista"].tolist() == [" B10024941", " 28954758N"]
    assert xls["Importe contrato (con IVA)"].tolist() == ["6364.6", "100"]
    assert _v(xls["Fecha"]) == ["2022-09-05", None]
    assert xls["_hoja"].tolist() == ["Masiva Contratos Menores"] * 2
    xlsx = df[df["_archivo_origen"] == "menores/menores_2024_1T.xlsx"]
    assert xlsx["Importe de adjudicación con impuestos"].tolist() == ["9783.1", "9783.1"]


def test_filas_de_titulo_antes_de_la_cabecera(portal, tmp_path):
    portal.paginas["registro-contratos-1t-2024"]["documentos"][0] = portal.documento(
        "Listado de contratos Menores 1T 2024.xlsx",
        _xlsx({"Sheet1": [["REGISTRO DE CONTRATOS - CONTRATOS MENORES"], ["1er. Trimestre 2024"], [],
                          CAB_C, ["CM0000000001/2024"] + FILA_C[1:], ["CM0000000002/2024"] + FILA_C[1:]]}))
    assert _ejecutar(tmp_path) == 0
    df = _menores(tmp_path)
    filas = df[df["_archivo_origen"] == "menores/menores_2024_1T.xlsx"]
    assert filas["Número de registro de contrato"].tolist() == ["CM0000000001/2024", "CM0000000002/2024"]
    assert filas["NIF"].tolist() == ["B86418787"] * 2
    assert "REGISTRO DE CONTRATOS - CONTRATOS MENORES" not in df.columns
    log = _log(tmp_path)
    assert "3 filas antes de la cabecera" in log and "REGISTRO DE CONTRATOS - CONTRATOS MENORES" in log


# ---------------------------------------------------------------------------
# Ficheros crudos y manifiesto
# ---------------------------------------------------------------------------

def test_nombres_locales_normalizados_y_manifiesto(portal, tmp_path):
    assert _ejecutar(tmp_path) == 0
    raw = tmp_path / "raw"
    assert sorted(p.name for p in (raw / "menores").iterdir() if p.is_file()) == [
        "menores_2022_1T.xls", "menores_2022_3T.xlsx", "menores_2023_2T-3T.xlsx", "menores_2024_1T.xlsx",
        "menores_2025_1T.xlsx", "menores_2026_2T.xlsx"]
    assert (raw / "modificaciones" / "modificaciones_2022_1T.xls").exists()
    assert (raw / "incidencias" / "incidencias_2023_2T-3T.xlsx").exists()
    doc = portal.paginas["registro-contratos-2t-2023"]["documentos"][0]
    entrada = _manifiesto(tmp_path)["menores/menores_2023_2T-3T.xlsx"]
    assert entrada["url"] == "https://www.juntaex.es" + doc["href"]
    assert entrada["pagina"] == "https://www.juntaex.es/w/registro-contratos-2t-2023"
    assert entrada["nombre_publicado"] == "CONTRATOS MENORES 2 y 3T 2023.xlsx"
    assert entrada["texto_enlace"] == "2.- CONTRATOS MENORES 2º y 3º T. 2023"
    assert entrada["clave"] == doc["clave"] and entrada["publicado"] is True
    assert entrada["tipo"] == "menores" and entrada["anio"] == 2023 and entrada["trimestres"] == [2, 3]
    assert entrada["sha256"] == M.sha256(raw / "menores" / "menores_2023_2T-3T.xlsx")
    assert entrada["version_portal"] == "2023-11-14T22:13:20Z"               # ?t=1700000000000
    assert entrada["fecha_descarga"] and entrada["comprobado"]
    assert (raw / "menores" / "menores_2023_2T-3T.xlsx").read_bytes() == portal.ficheros[doc["ruta"]]


def test_dos_documentos_del_mismo_trimestre_no_se_pisan(portal, tmp_path):
    assert _ejecutar(tmp_path) == 0
    SLEEP_REAL(1.1)
    # Aparece un segundo listado de menores de 1T 2024: el primero conserva su nombre
    extra = portal.documento("Listado de contratos Menores 1T 2024 (complementario).xlsx",
                             _menores_c("CM0000000099/2024"))
    portal.paginas["registro-contratos-1t-2024"]["documentos"].append(extra)
    assert _ejecutar(tmp_path) == 0
    menores = tmp_path / "raw" / "menores"
    segundo = f"menores_2024_1T_{extra['clave'][:8]}.xlsx"
    assert (menores / segundo).exists() and (menores / "menores_2024_1T.xlsx").exists()
    man = _manifiesto(tmp_path)
    assert man["menores/menores_2024_1T.xlsx"]["nombre_publicado"] == "Listado de contratos Menores 1T 2024.xlsx"
    assert man[f"menores/{segundo}"]["nombre_publicado"] == extra["nombre"]
    df = _menores(tmp_path)
    assert df.loc[df["_trimestre"] == "1T", "Número de registro de contrato"].tolist().count("CM0000000099/2024") == 1
    assert df["_en_ultima_descarga"].all()


def test_solo_se_vuelven_a_pedir_los_recientes_o_con_otra_version(portal, tmp_path):
    reciente = portal.documento(f"LISTADO CONTRATOS MENORES 1T {ANIO}.xlsx", _menores_c(f"CM0000000001/{ANIO}"))
    portal.publicar(f"registro-contratos-1t-{ANIO}", f"1er. Trimestre de {ANIO}", [reciente])
    portal.buscador.append(f"registro-contratos-1t-{ANIO}")
    antiguo = "Listado de contratos Menores 1T 2024.xlsx"
    assert _ejecutar(tmp_path) == 0
    assert _ejecutar(tmp_path) == 0
    assert portal.pedidas_documento(antiguo) == 1
    assert portal.pedidas_documento(reciente["nombre"]) == 2
    assert "menores/menores_2024_1T.xlsx (ya descargado" in _log(tmp_path)
    # El portal lo enlaza con otra versión (?t= distinto): se vuelve a pedir
    documentos = portal.paginas["registro-contratos-1t-2024"]["documentos"]
    documentos[0] = portal.documento(antiguo, portal.ficheros[documentos[0]["ruta"]], t=1800000000000)
    assert _ejecutar(tmp_path) == 0
    assert portal.pedidas_documento(antiguo) == 2
    assert _ejecutar(tmp_path, "--comprobar-todo") == 0
    assert portal.pedidas_documento(antiguo) == 3
    assert portal.pedidas_documento("LISTADO CONTRATOS MENORES 1T 2022.xls") == 2
    assert not (tmp_path / "raw" / "menores" / "_historico").exists()       # nada cambió


# ---------------------------------------------------------------------------
# Sesgo del superviviente
# ---------------------------------------------------------------------------

def test_documento_sustituido_conserva_la_version_anterior(portal, tmp_path):
    assert _ejecutar(tmp_path) == 0
    SLEEP_REAL(1.1)          # la versión nueva tiene que llevar otra fecha (resolución: 1 s)
    documentos = portal.paginas["registro-contratos-1t-2024"]["documentos"]
    # Mismo documento (mismo uuid), otra versión: CM...02 se corrige a CM...03
    documentos[0] = portal.documento(documentos[0]["nombre"], _menores_c("CM0000000001/2024", "CM0000000003/2024"),
                                     t=1800000000000, clave=documentos[0]["clave"])
    assert _ejecutar(tmp_path) == 0
    raw = tmp_path / "raw" / "menores" / "menores_2024_1T.xlsx"
    assert len(M.versiones(raw)) == 2
    df = _menores(tmp_path)
    filas = df[df["_archivo_origen"] == "menores/menores_2024_1T.xlsx"]
    assert filas[["Número de registro de contrato", "_en_ultima_descarga"]].values.tolist() == [
        ["CM0000000001/2024", True], ["CM0000000002/2024", False], ["CM0000000003/2024", True]]


def test_documento_nuevo_que_sustituye_a_otro_retirado_ocupa_su_nombre(portal, tmp_path):
    assert _ejecutar(tmp_path) == 0
    SLEEP_REAL(1.1)
    # El portal borra el documento y sube otro (uuid nuevo) con el mismo nombre
    documentos = portal.paginas["registro-contratos-1t-2024"]["documentos"]
    documentos[0] = portal.documento(documentos[0]["nombre"], _menores_c("CM0000000001/2024"),
                                     clave="0a577059-cc9d-dc58-0780-be92e9b68054", t=1800000000000)
    assert _ejecutar(tmp_path) == 0
    raw = tmp_path / "raw" / "menores"
    assert sorted(p.name for p in raw.iterdir() if p.name.startswith("menores_2024")) == ["menores_2024_1T.xlsx"]
    assert len(M.versiones(raw / "menores_2024_1T.xlsx")) == 2
    assert _manifiesto(tmp_path)["menores/menores_2024_1T.xlsx"]["clave"] == "0a577059-cc9d-dc58-0780-be92e9b68054"
    df = _menores(tmp_path)
    filas = df[df["_archivo_origen"] == "menores/menores_2024_1T.xlsx"]
    assert filas[["Número de registro de contrato", "_en_ultima_descarga"]].values.tolist() == [
        ["CM0000000001/2024", True], ["CM0000000002/2024", False]]


def test_documento_que_deja_de_enlazarse_queda_retirado(portal, tmp_path):
    assert _ejecutar(tmp_path) == 0
    antes = _menores(tmp_path)
    SLEEP_REAL(1.1)
    portal.quitar_enlace("registro-contratos-1t-2024", "Listado de contratos Menores 1T 2024.xlsx")
    assert _ejecutar(tmp_path) == 0

    raw = tmp_path / "raw" / "menores" / "menores_2024_1T.xlsx"
    assert raw.exists()                                                  # no se borra nada
    entrada = _manifiesto(tmp_path)["menores/menores_2024_1T.xlsx"]
    assert entrada["publicado"] is False and entrada["retirado_desde"]
    df = _menores(tmp_path)
    assert len(df) == len(antes)
    retirado = df["_archivo_origen"] == "menores/menores_2024_1T.xlsx"
    assert retirado.sum() == 2 and not df.loc[retirado, "_en_ultima_descarga"].any()
    assert df.loc[~retirado, "_en_ultima_descarga"].all()
    assert "menores/menores_2024_1T.xlsx (Listado de contratos Menores 1T 2024.xlsx): ninguna página" in _log(tmp_path)

    # Si vuelve a enlazarse, sus filas vuelven a estar publicadas
    SLEEP_REAL(1.1)
    _publicar_base(portal)
    assert _ejecutar(tmp_path) == 0
    df = _menores(tmp_path)
    assert len(df) == len(antes) and df["_en_ultima_descarga"].all()
    assert _manifiesto(tmp_path)["menores/menores_2024_1T.xlsx"]["publicado"] is True


def test_pagina_que_desaparece_retira_lo_que_solo_ella_enlazaba(portal, tmp_path):
    assert _ejecutar(tmp_path) == 0
    SLEEP_REAL(1.1)
    portal.paginas["registro-contratos-1t-2024"] = 404
    portal.buscador.remove("registro-contratos-1t-2024")
    assert _ejecutar(tmp_path) == 0
    man = _manifiesto(tmp_path)
    assert man["menores/menores_2024_1T.xlsx"]["publicado"] is False
    assert man["incidencias/incidencias_2024_1T.xlsx"]["publicado"] is False
    assert man["menores/menores_2023_2T-3T.xlsx"]["publicado"] is True
    df = _menores(tmp_path)
    assert not df.loc[df["_anio"] == "2024", "_en_ultima_descarga"].any()
    assert df.loc[df["_anio"] != "2024", "_en_ultima_descarga"].all()
    assert "registro-contratos-1t-2024: la página ya no está" in _log(tmp_path)


def _sin_documentos(portal):
    portal.paginas["registro-contratos-2t-2023"]["documentos"] = []


def _buscador_sin_resultados(portal):
    portal.buscador = []


def _buscador_caido(portal):
    portal.buscador = 503


def _pagina_caida(portal):
    portal.paginas["registro-contratos-2t-2023"] = 503


def _pagina_del_buscador_con_404(portal):
    portal.paginas["registro-contratos-1t-2022"] = 404


@pytest.mark.parametrize("fallo, mensaje", [
    (_buscador_caido, "buscador del portal"),
    (_buscador_sin_resultados, "no devuelve ninguna página del registro"),
    (_pagina_caida, "registro-contratos-2t-2023: HTTP 503"),
    (_pagina_del_buscador_con_404, "el buscador la lista pero da HTTP 404"),
    (_sin_documentos, "la página ya no enlaza ningún documento"),
])
def test_fallo_del_portal_no_retira_nada(portal, tmp_path, fallo, mensaje):
    assert _ejecutar(tmp_path) == 0
    antes = _menores(tmp_path)
    SLEEP_REAL(1.1)
    fallo(portal)
    # Aunque además deje de enlazarse un documento, no se retira nada
    portal.quitar_enlace("registro-contratos-1t-2024", "Listado de contratos Menores 1T 2024.xlsx")
    assert _ejecutar(tmp_path) == 1
    assert mensaje in _log(tmp_path)
    despues = _menores(tmp_path)
    assert len(despues) == len(antes) and despues["_en_ultima_descarga"].all()
    assert all(e["publicado"] for e in _manifiesto(tmp_path).values())


def test_con_una_pagina_caida_un_documento_nuevo_no_ocupa_la_ruta_de_otro(portal, tmp_path):
    assert _ejecutar(tmp_path) == 0
    antes = (tmp_path / "raw" / "menores" / "menores_2023_2T-3T.xlsx").read_bytes()
    SLEEP_REAL(1.1)
    # La página de 2T-3T 2023 falla y otra enlaza un documento nuevo del mismo trimestre y tipo:
    # el de la página caída no se ha retirado, así que su fichero no se toca
    portal.paginas["registro-contratos-2t-2023"] = 503
    nuevo = portal.documento("CONTRATOS MENORES 2 y 3T 2023 (anexo).xlsx", _menores_c("CM0000099999/2023"))
    portal.paginas["registro-contratos-1t-2024"]["documentos"].append(nuevo)
    assert _ejecutar(tmp_path) == 1
    menores = tmp_path / "raw" / "menores"
    assert (menores / "menores_2023_2T-3T.xlsx").read_bytes() == antes
    assert len(M.versiones(menores / "menores_2023_2T-3T.xlsx")) == 1
    assert (menores / f"menores_2023_2T-3T_{nuevo['clave'][:8]}.xlsx").exists()
    df = _menores(tmp_path)
    assert df["_en_ultima_descarga"].all() and len(df) == 8


def test_documento_enlazado_que_falla_conserva_la_copia(portal, tmp_path):
    assert _ejecutar(tmp_path) == 0
    SLEEP_REAL(1.1)
    doc = portal.paginas["contratos-incidencias-inscritas-registro-contratos"]["documentos"][0]
    original = portal.ficheros[doc["ruta"]]
    portal.ficheros[doc["ruta"]] = b"<!DOCTYPE html><html><body>Mantenimiento</body></html>"
    assert _ejecutar(tmp_path, "--comprobar-todo") == 1
    raw = tmp_path / "raw" / "menores" / "menores_2026_2T.xlsx"
    assert len(M.versiones(raw)) == 1 and raw.read_bytes() == original
    assert "menores/menores_2026_2T.xlsx: la respuesta es HTML" in _log(tmp_path)
    assert _menores(tmp_path)["_en_ultima_descarga"].all()


def test_trimestre_confirmado_sin_listado_de_menores_es_error(portal, tmp_path):
    portal.quitar_enlace("registro-contratos-1t-2024", "Listado de contratos Menores 1T 2024.xlsx")
    assert _ejecutar(tmp_path) == 1
    assert "menores 1T 2024: trimestre publicado según las fuentes" in _log(tmp_path)


def test_informe_de_menores_por_trimestre(portal, tmp_path):
    assert _ejecutar(tmp_path) == 0
    log = _log(tmp_path)
    assert "CONTRATOS MENORES POR TRIMESTRE" in log
    assert "2023 2T-3T: 1 filas" in log and "2024 1T: 2 filas" in log


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def test_cli(portal, tmp_path, monkeypatch):
    # Como programa, el script usa su configuración real (sin los parches del
    # fixture): el portal simulado publica también los trimestres confirmados
    for anio, t in CONFIRMADOS_REALES:
        if (anio, t) not in CONFIRMADOS:
            portal.publicar(f"registro-contratos-{t}t-{anio}", f"{t}º Trimestre de {anio}", [
                portal.documento(f"LISTADO CONTRATOS MENORES {t}T {anio}.xlsx", _menores_c(f"CM{t}/{anio}"))])
    monkeypatch.setattr(sys, "argv", [str(SCRIPT), "--salida", str(tmp_path)])
    with pytest.raises(SystemExit) as salida:
        runpy.run_path(str(SCRIPT), run_name="__main__")
    assert salida.value.code == 0
    assert (tmp_path / "registro_contratos_menores.parquet").exists()


def test_solo_procesar_no_descarga(portal, tmp_path):
    assert _ejecutar(tmp_path, "--solo-descarga") == 0
    assert not list(tmp_path.glob("*.parquet"))
    llamadas = len(portal.llamadas)
    assert _ejecutar(tmp_path, "--solo-procesar") == 0
    assert len(portal.llamadas) == llamadas
    assert len(_menores(tmp_path)) == 7


def test_sin_cambios_el_parquet_no_cambia(portal, tmp_path):
    # Con pandas 3 las columnas leídas del Parquet anterior son 'str' y las de
    # un Excel 'object': el fichero no debe cambiar por eso
    assert _ejecutar(tmp_path) == 0
    antes = {p.name: p.read_bytes() for p in tmp_path.glob("*.parquet")}
    SLEEP_REAL(1.1)
    assert _ejecutar(tmp_path, "--solo-procesar") == 0
    assert {p.name: p.read_bytes() for p in tmp_path.glob("*.parquet")} == antes
    assert not (tmp_path / "_historico").exists()


def test_salida_por_defecto_en_el_repo():
    assert M.SALIDA == REPO_ROOT / "ccaa_extremadura"


def _en_mayusculas(resultado):
    """Una «lectura arreglada»: el mismo resultado con el texto de las columnas del origen en
    mayúsculas (las de control, _..., igual)."""
    df, *resto = resultado
    df = df.copy()
    for c in df.columns:
        if not str(c).startswith("_") and (df[c].dtype == object or pd.api.types.is_string_dtype(df[c])):
            df[c] = df[c].str.upper()
    return (df, *resto)


def _comprobar_arreglo(antes, despues):
    """Regla 3: mismas filas, historia intacta y el arreglo en TODAS las filas del origen."""
    assert len(despues) == len(antes) > 0
    assert despues["_primera_descarga"].tolist() == antes["_primera_descarga"].tolist()
    for c in despues.columns:
        if not str(c).startswith("_") and (despues[c].dtype == object or pd.api.types.is_string_dtype(despues[c])):
            valores = despues[c].dropna()
            assert (valores == valores.str.upper()).all(), c

def test_un_arreglo_de_lectura_llega_a_las_filas_ya_guardadas(portal, tmp_path, monkeypatch):
    """Regla 3: el Parquet se construye con el código actual desde todas las versiones del crudo.
    Antes solo se aplicaban las versiones posteriores al Parquet anterior y un arreglo de lectura no
    llegaba a las filas ya guardadas."""
    assert _ejecutar(tmp_path) == 0
    antes = {p.name: pd.read_parquet(p) for p in sorted(tmp_path.glob("*.parquet"))}
    leer = M.leer_tabla
    monkeypatch.setattr(M, "leer_tabla", lambda ruta: _en_mayusculas(leer(ruta)))
    assert _ejecutar(tmp_path, "--solo-parquet") == 0
    assert antes
    for nombre, df in antes.items():
        _comprobar_arreglo(df, pd.read_parquet(tmp_path / nombre))
