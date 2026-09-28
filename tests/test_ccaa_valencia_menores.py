"""Tests offline de scripts/ccaa_valencia_menores.py.

Las páginas índice, los ficheros, el CKAN de la UPV y el buscador de contratos
menores del Ajuntament de València (portlet Liferay: cookie de sesión, p_auth,
filtro por fecha y hora con 'hasta' a las 00:00 y máximo de 500 filas) se simulan con un
``requests.get`` y un ``requests.post`` falsos; los XLSX (también uno con la
caché de una tabla dinámica, como los de gastos menores de la UV), XLS y CSV se
generan en el propio test.
"""

import html
import importlib.util
import io
import json
import runpy
import sys
import time
import zipfile
from datetime import date, datetime, time as hora, timedelta
from pathlib import Path
from urllib.parse import quote, unquote

import openpyxl
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
import pytest
import requests

REPO_ROOT = Path(__file__).resolve().parents[1]
SCRIPT = REPO_ROOT / "scripts" / "ccaa_valencia_menores.py"
SLEEP_REAL = time.sleep


def _cargar():
    spec = importlib.util.spec_from_file_location("ccaa_valencia_menores", SCRIPT)
    modulo = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(modulo)
    return modulo


M = _cargar()
ANIO = datetime.now().year               # el script decide qué refrescar con el año real
CERRADO = ANIO - 3                       # un año que ya no se vuelve a pedir sin --comprobar-todo

UV = "https://www.uv.es/contratacion/PORTALTRANSPARENCIA/menores/"
UV_HTTP = "http://www.uv.es/contratacion/PORTALTRANSPARENCIA/menores/"
DIPA = "https://abierta.diputacionalicante.es/wp-content/uploads/"
UMH = "https://seguimientocontratacion.umh.es/files/"
PLACSP_DOC = ("https://contrataciondelestado.es/wps/wcm/connect/PLACE_es/Site/area/docAccCmpnt?srv=cmpnt"
              "&cmpntname=GetDocumentsById&source=library&DocumentIdParam=77f3c628-ec33")


# ---------------------------------------------------------------------------
# Ficheros
# ---------------------------------------------------------------------------

def _xlsx(hojas):
    libro = openpyxl.Workbook()
    libro.remove(libro.active)
    for nombre, filas in hojas.items():
        hoja = libro.create_sheet(nombre)
        for fila in filas:
            hoja.append(fila)
    datos = io.BytesIO()
    libro.save(datos)
    return datos.getvalue()


def _xls(filas):
    """XLS binario (con xlwt) o, si no está instalado, un XLSX: el script lee por
    el contenido, no por la extensión."""
    try:
        import xlwt
    except ImportError:
        return _xlsx({"Hoja1": filas})
    libro = xlwt.Workbook()
    hoja = libro.add_sheet("Sheet0")
    for i, fila in enumerate(filas):
        for j, valor in enumerate(fila):
            if valor is not None:
                hoja.write(i, j, valor)
    datos = io.BytesIO()
    libro.save(datos)
    return datos.getvalue()


NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
NS_OFF = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _xlsx_tabla_dinamica():
    """Como 'GM 2026_1 tabla dinamica cas.xlsx' de la UV: la hoja es el resumen
    de una tabla dinámica y los registros solo están en su caché (la hoja de
    origen es un libro externo). Un campo de agrupación (databaseField="0") no
    está en los registros."""
    base = _xlsx({"Hoja1": [["UNIDAD FUNCIONAL", "Importe Adjudicación\n(IVA excluido)", "Nº"],
                            ["Rectorat", 118.56, 2], ["B46669693", 118.56, 2], ["Biblioteca", 5991.22, 1],
                            ["Total general", 6109.78, 3]]})
    definicion = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<pivotCacheDefinition xmlns="{NS_MAIN}" xmlns:r="{NS_OFF}" r:id="rId1" recordCount="3">'
        f'<cacheSource type="worksheet"><worksheetSource ref="A1:G4" sheet="Hoja1" r:id="rId2"/></cacheSource>'
        f'<cacheFields count="7">'
        f'<cacheField name="UNIDAD FUNCIONAL" numFmtId="0"><sharedItems count="2"><s v="Rectorat"/>'
        f'<s v="Biblioteca"/></sharedItems></cacheField>'
        f'<cacheField name="NIF ADJUDICATARIO" numFmtId="0"><sharedItems count="2"><s v="B46669693"/>'
        f'<s v="A79486833"/></sharedItems></cacheField>'
        f'<cacheField name="OBJETO" numFmtId="0"><sharedItems/></cacheField>'
        f'<cacheField name="IMPORTE ADJUDICACIÓN (IVA EXCL.)" numFmtId="0"><sharedItems containsNumber="1"/>'
        f'</cacheField>'
        f'<cacheField name="PLAZO EJECUCIÓN_x000a_(meses)" numFmtId="0"><sharedItems containsBlank="1"/>'
        f'</cacheField>'
        f'<cacheField name="FECHA ADJUDICACIÓN" numFmtId="14"><sharedItems containsDate="1"/></cacheField>'
        f'<cacheField name="Meses" numFmtId="0" databaseField="0"><fieldGroup base="5"/></cacheField>'
        f'</cacheFields></pivotCacheDefinition>')
    registros = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><pivotCacheRecords xmlns="{NS_MAIN}" count="3">'
        f'<r><x v="0"/><x v="0"/><s v="Suministro de bridas para urnas"/><n v="18.559999999999999"/><m/>'
        f'<d v="2026-03-03T00:00:00"/></r>'
        f'<r><x v="1"/><x v="1"/><s v="Estores &quot;opacos&quot;"/><n v="5991.22"/><n v="0.42"/>'
        f'<d v="2026-02-11T00:00:00"/></r>'
        f'<r><x v="0"/><x v="0"/><s v="Lámpara"/><n v="100"/><n v="12"/><d v="2026-01-30T10:30:00"/></r>'
        f'</pivotCacheRecords>')
    rels_definicion = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="{NS_PKG}">'
        f'<Relationship Id="rId2" Type="{NS_OFF}/externalLinkPath" Target="/contratacion/disco/GM_1T_2026.xlsx" '
        f'TargetMode="External"/>'
        f'<Relationship Id="rId1" Type="{NS_OFF}/pivotCacheRecords" Target="pivotCacheRecords1.xml"/>'
        f'</Relationships>')
    tabla = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><pivotTableDefinition xmlns="{NS_MAIN}" '
             f'name="TablaDinámica4" cacheId="0" dataCaption="Valores"/>')
    rels_tabla = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="{NS_PKG}">'
                  f'<Relationship Id="rId1" Type="{NS_OFF}/pivotCacheDefinition" '
                  f'Target="../pivotCache/pivotCacheDefinition1.xml"/></Relationships>')
    rels_hoja = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="{NS_PKG}">'
                 f'<Relationship Id="rId1" Type="{NS_OFF}/pivotTable" Target="../pivotTables/pivotTable1.xml"/>'
                 f'</Relationships>')
    entrada, salida = zipfile.ZipFile(io.BytesIO(base)), io.BytesIO()
    with zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as nuevo:
        for nombre in entrada.namelist():
            nuevo.writestr(nombre, entrada.read(nombre))
        nuevo.writestr("xl/pivotCache/pivotCacheDefinition1.xml", definicion)
        nuevo.writestr("xl/pivotCache/pivotCacheRecords1.xml", registros)
        nuevo.writestr("xl/pivotCache/_rels/pivotCacheDefinition1.xml.rels", rels_definicion)
        nuevo.writestr("xl/pivotTables/pivotTable1.xml", tabla)
        nuevo.writestr("xl/pivotTables/_rels/pivotTable1.xml.rels", rels_tabla)
        nuevo.writestr("xl/worksheets/_rels/sheet1.xml.rels", rels_hoja)
    return salida.getvalue()


def _cm_uv(identificador, nif="B46669693"):
    """Como CM_2026_1_cas.xlsx: dos filas de título y la cabecera."""
    return _xlsx({"Hoja1": [
        [None, "Informe Menores Adjudicados"],
        ["FECHA ADJUDICACIÓN DESDE: 01/01/2026   FECHA ADJUDICACIÓN HASTA: 31/03/2026"],
        ["UNIDAD FUNCIONAL", "IDENTIFICADOR CONTRATO", "OBJETO", "NIF ADJUDICATARIO", "IMPORTE ADJUDICACIÓN (IVA EXCL.)",
         "FECHA ADJUDICACIÓN"],
        ["Rectorat", identificador, "Estores opacos", nif, 5991.22, datetime(2026, 3, 25)],
    ]})


# ---------------------------------------------------------------------------
# Páginas índice
# ---------------------------------------------------------------------------

def _a(href, texto=""):
    return f'<a href="{html.escape(href)}" target="_blank">{texto}</a>'


def pagina_uv(idioma="cas", sin=()):
    """Como la página de contratos menores de la UV: tabla con grupos (rowspan)
    en 2026, párrafos en 2023 y antes, enlaces http:// y trimestres comentados.
    `sin`: rutas que no se enlazan."""
    def enlace(ruta, texto="", http=False):
        return "" if ruta in sin else _a((UV_HTTP if http else UV) + ruta, texto)
    img = '<img alt="" src="//www.uv.es/contratacion/Imagenes/imagen excel.png"/> '
    return f"""<!DOCTYPE html><html><head><title>Contratos menores</title></head><body>
<div class="entry-content"><hr/>
<div class="desplegar12"><p><strong>EJERCICIO 2026</strong></p></div>
<div class="desplegar-int12" style="display:none"><table border="1" class="sortable"><tbody>
<tr class="fila1_a"><td><strong>Fichero</strong></td><td><strong>Per&iacute;odo</strong></td>
<td><strong>Descargar</strong></td></tr>
<!-- Contratos Menores-->
<tr class="fila1_b1"><td rowspan="2"><strong>1. Contratos menores</strong></td><td>1er Trimestre</td>
<td>{enlace(f"2026/1er_trimestre_26/CM_2026_1_{idioma}.xlsx", img)}</td></tr>
<tr class="fila1_b2"><td>2n Trimestre</td><td>{enlace(f"2026/2do_trimestre_26/CM_2T_{idioma}.xlsx", img)}</td></tr>
<!--
<tr class="fila1_b3"><td>3er Trimestre</td><td>{_a(UV + "2026/3er_trimestre_26/CM_2026_3_cas.xlsx", img)}</td></tr>
-->
<!--Gastos Menores-->
<tr class="fila2_b1"><td rowspan="1"><strong>2. Gastos menores</strong></td><td>1er Trimestre</td>
<td>{enlace(f"2026/1er_trimestre_26/GM 2026_1 tabla dinamica {idioma}.xlsx", img)}</td></tr>
</tbody></table></div>
<div class="desplegar9"><p><strong>EJERCICIO 2023</strong></p></div>
<div class="desplegar-int9">
<p>{enlace("2023/4_trimestre_2023/OTROS%20GASTOS_TRIMESTRE4_2023.pdf",
           "<u><strong>4r trimestre Otros gastos 2023 (pdf)</strong></u>", http=True)}</p>
<p>{enlace("2023/4_trimestre_2023/OTROS%20GASTOS_TRIMESTRE4_2023.xlsx",
           "<u><strong>4r trimestre Otros gastos 2023 (xls)</strong></u>")}</p>
<p>{_a("http://Z:" + chr(92) + "contratacion" + chr(92) + "OTROS GASTOS_TRIMESTRE3_2023.xlsx", ")")
    if idioma == "val" else ""}</p>
</div>
<div class="desplegar8"><p><strong>EJERCICIO 2022</strong></p></div>
<div class="desplegar-int8"><p>{enlace("2022/1T/2022_1T_CM_PORTALTRANSPARENCIA.xls", "- 1r&nbsp;trimestre 2022&nbsp;(xls)",
                                      http=True)}</p></div>
<div class="desplegar1"><p><strong>EJERCICIO 2015</strong></p></div>
<div class="desplegar-int1"><p>{enlace("LEY_TRANSPARENCIA_menores_4trimestre.xls",
                                      "- 4&ordm; trimestre&nbsp;2015 (format xls)", http=True)}</p></div>
</div>
<a href="https://www.uv.es/uvweb/transparencia-uv/es/presentacion.html">Contratación</a>
</body></html>""".encode("utf-8")


def publicar_uv(portal):
    portal.urls[M.UV_PAGINA] = pagina_uv("cas")
    portal.urls[M.UV_PAGINA_VA] = pagina_uv("val")
    for idioma in ("cas", "val"):
        portal.fichero(UV + f"2026/1er_trimestre_26/CM_2026_1_{idioma}.xlsx", _cm_uv(f"2026SU05097CM-{idioma}"))
        portal.fichero(UV + f"2026/2do_trimestre_26/CM_2T_{idioma}.xlsx", _cm_uv(f"2026SU06001CM-{idioma}"))
        portal.fichero(UV + f"2026/1er_trimestre_26/GM 2026_1 tabla dinamica {idioma}.xlsx", _xlsx_tabla_dinamica())
    portal.fichero(UV + "2023/4_trimestre_2023/OTROS GASTOS_TRIMESTRE4_2023.xlsx", _xlsx({
        "otros gastos trimestres 4- 2023": [["AM del que deriva el contrato", "Identificador", "NIF adjudicatario"],
                                           ["AM 5/20CC", "2023 054428 SU-ot", "A79206223"]]}))
    portal.fichero(UV + "2023/4_trimestre_2023/OTROS GASTOS_TRIMESTRE4_2023.pdf", b"%PDF-1.4 no se pide")
    portal.fichero(UV + "2022/1T/2022_1T_CM_PORTALTRANSPARENCIA.xls", _xls([
        ["Unidad funcional", "Identificador", "NIF adjudicatario", "Importe adjudicación (IVA excluido)"],
        ["Campus d'Ontinyent", "2022 011354 SU-cm", "A80837941", 74.42],
        ["Información publicada en cumplimiento de la Ley de Contratos del Sector Público", None, None, None]]))
    portal.fichero(UV + "LEY_TRANSPARENCIA_menores_4trimestre.xls", _xls([
        ["FECHAOP", "DESCRIPCION", "NIF TERCERO", "NOMTERCERO", "TOTALCONTRATO"],
        ["2015-10-01", "ACTIVITATS CULTURALS", "G54760699", "ASOC TEADA", 484]]))


def pagina_diputacion():
    """Como la de contratación de la Diputación de Alicante: los menores en la
    sección 'CONTRATOS MENORES' (h2); antes y después, otras secciones."""
    return f"""<html><body><h1>Contratación</h1>
<h2 class="PopUp_TituloPrincipal"><span>Relación de contratos celebrados por la Diputación de Alicante (no incluye los
menores)</span></h2>
<table><tbody><tr><td>{_a(DIPA + "RP-1447-Contratos-No-Menores-gestionados-por-Diputacion-2025-CAS.xlsx", "Año 2025")}
</td></tr></tbody></table>
<div class="PopUp_ContenedorPadre"><h2 class="PopUp_TituloPrincipal"><span class="gb-text">CONTRATOS MENORES</span>
<span class="gb-shape"><svg viewBox="0 0 16 16"><path d="M8 16"></path></svg></span></h2>
<p class="gb-text PopUp_Titulo">Contratos menores</p></div>
<div class="kt-accordion"><button><span class="kt-blocks-accordion-title">Contratos Menores 2026</span></button>
<figure class="wp-block-table"><table class="has-fixed-layout"><tbody>
<tr><td><strong>{_a(DIPA + "RP-1628-Adjudic-a-30-06-2026-cas.xlsx",
                    "Contratos menores adjudicados a 30 de junio de 2026 (LIGATE)")}</strong></td></tr>
</tbody></table></figure>
<figure class="wp-block-table"><table class="has-fixed-layout"><tbody>
<tr><td>{_a(DIPA + "1CORPORACION20232027/Eco-Presup/Contratos/RP-922-Adjudic-a-31-12-2024-cas.xlsx")}<strong>
{_a(DIPA + "Contratos-adjudicados-LIGATE-a-31-de-diciembre-de-2021-2.xlsx",
    "Contratos menores adjudicados a 31 de diciembre de 2021 (LIGATE)")}</strong></td></tr>
</tbody></table></figure>
<ul><li>{_a(DIPA + "Instruccion-contratos-menores.pdf", "Ver contenido completo de la Instrucción")}</li></ul>
<figure><table><tbody><tr><td>Relación de contratos menores. Primer trimestre 2016:</td>
<td>{_a(DIPA + "2016-Contratos-Menores-Primer-Trimestre-Vers26112018.xlsx", "<strong>XSLX</strong>")}</td>
<td>{_a(DIPA + "2016-Contratos-Menores-Primer-Trimestre-Vers26112018.pdf", "<strong>PDF</strong>")}</td></tr>
</tbody></table></figure></div>
<h2 class="PopUp_TituloPrincipal"><span>Relación de contratos de arrendamiento de inmuebles</span></h2>
<p>{_a(DIPA + "CAS-Contratos-arrendam-INMUEBLES-anualidad-2022.xlsx", "Contratos de arrendamiento año 2022")}</p>
</body></html>""".encode("utf-8")


def publicar_diputacion(portal):
    portal.urls[M.DIPA_PAGINA] = pagina_diputacion()
    portal.fichero(DIPA + "RP-1628-Adjudic-a-30-06-2026-cas.xlsx", _xlsx({"30-06-2026": [
        [None, "Versión núm. 1: 11 de agosto de 2026", None, None, None, "Documento reelaborado por Transparencia"],
        [None, None, None, None, None, "Fuente: Registro de Contratos Menores (LIGATE)"],
        [None, "RELACIÓN DE CONTRATOS MENORES QUE CONSTAN ADJUDICADOS EN EL REGISTRO ELECTRÓNICO LIGATE"],
        [None, "Contratos registrados del 01/04/2026 a 30/06/2026 "],
        [None, "TERCERO", "TIPO CONTRATO", "OBJETO", "ADJUDICADO (sin IVA)", "REFERENCIA"],
        [None, "ARQUES MORANTE, MARIA JESUS", "Servicios", "Formación de conductores", 1500, "DIPU/2026/93/0001"],
        [None, "LYRECO ESPAÑA, S.A.", "Suministros", "Destructora de papel", 357.19, "DIPU/2026/93/0005"]]}))
    portal.fichero(DIPA + "Contratos-adjudicados-LIGATE-a-31-de-diciembre-de-2021-2.xlsx", _xlsx({"2021": [
        ["TERCERO", "TIPO CONTRATO", "OBJETO", "ADJUDICADO (sin IVA)", "REFERENCIA"],
        ["ZUMALABE LOZANO, RAFAEL", "Servicios", "Topografía", 3463.02, "DIPU/2021/35/0052"]]}))
    portal.fichero(DIPA + "2016-Contratos-Menores-Primer-Trimestre-Vers26112018.xlsx", _xlsx({
        "Contratos Menores 1er Trim 2016": [
            ["Versión núm. 3: 26 de noviembre de 2018", None, None, None, None, "Documento reelaborado"],
            ["C.Gestor", "Tipo Contrato", "Fecha", "Importe", "Nombre Ter.", "Texto Libre"],
            ["01-PRESIDENCIA", "SERVICIOS", datetime(2016, 2, 19), 1999.99, "INFOEXPRESS, S.A.", "Publicidad"]]}))
    portal.fichero(DIPA + "RP-1447-Contratos-No-Menores-gestionados-por-Diputacion-2025-CAS.xlsx",
                   _xlsx({"No": [["no"], ["se pide"]]}))
    portal.fichero(DIPA + "CAS-Contratos-arrendam-INMUEBLES-anualidad-2022.xlsx", _xlsx({"No": [["no"], ["x"]]}))


def pagina_umh():
    return f"""<html><body><h2>Transparencia</h2>
<h3>Plan Anual de Contratación</h3><ul><li>{_a(UMH + "2025/02/RR_00414_2025.xlsx", "Plan 2025")}</li></ul>
<h3><strong>Contratos Menores</strong></h3>
<ul><li>{_a("https://sicgef.umh.es/contratos-menores-transparencia-sicgef/", "Contratos menores 2014-18")}</li>
<li>Contratos menores 2019<ul><li>{_a(PLACSP_DOC, "Primer trimestre 2019")}</li></ul></li>
<li>Contratos menores 2026<ul>
<li>{_a(UMH + "2026/09/DOC20260729100619CONTRATOS-MENORES-2T-2026.xlsx", "Segundo trimestre 2026")}</li></ul></li></ul>
<h3>Contratos Basados Acuerdos Marco</h3>
<ul><li>{_a(UMH + "2026/09/DOC20260717091255BASADOS-2T-2026.xlsx", "Segundo trimestre 2026")}</li></ul>
</body></html>""".encode("utf-8")


def publicar_umh(portal):
    portal.urls[M.UMH_PAGINA] = pagina_umh()
    portal.fichero(PLACSP_DOC, _xlsx({"CONTRATOS MENORES": [
        ["Código", "N.I.F. del proveedor", "Nombre del proveedor", "Importe total"],
        ["2019/0001", "B81152217", "ADLER INSTRUMENTOS, S.L.", 11132]]}))
    portal.fichero(UMH + "2026/09/DOC20260729100619CONTRATOS-MENORES-2T-2026.xlsx", _xlsx({"INFORME": [
        ["LISTADO DE CONTRATOS MENORES"], ["Segundo Trimestre 2026"],
        ["Código del expediente", "N.I.F. del proveedor", "Nombre del proveedor", "Importe total"],
        ["2026/0000000365", "B81152217", "ADLER INSTRUMENTOS, S.L.", 118]]}))
    portal.fichero(UMH + "2026/09/DOC20260717091255BASADOS-2T-2026.xlsx", _xlsx({"No": [["no"], ["x"]]}))


# ---------------------------------------------------------------------------
# UPV: CKAN
# ---------------------------------------------------------------------------

UPV_BASE = "https://upvtransparent.upv.es/dataset/"


def paquetes_upv(acumulado="ley_2026-07-20.csv"):
    return {
        "contratos-menores": {"name": "contratos-menores", "resources": [
            {"format": "HTML", "name": "Diccionario de datos", "url": "https://intranet.upv.es/pls/soalu/ayuda"},
            {"format": "CSV", "year": "2023", "name": "(Año 2023) Contratos menores",
             "url": UPV_BASE + "a/resource/1/download/2024-143914_menores_2024-03-04_trimestral.csv"},
            {"format": "XLSX", "year": "2023", "name": "(Año 2023) Contratos menores",
             "url": UPV_BASE + "a/resource/2/download/2024-143916_menores_2024-02-15_trimestral.xlsx"}]},
        "contratos-menores-ley-1-2022": {"name": "contratos-menores-ley-1-2022", "resources": [
            {"format": "CSV", "year": "2026", "name": "(Acumulado hasta 2º Trimestre 2026) Contratos menores",
             "url": UPV_BASE + f"b/resource/3/download/{acumulado}"}]},
    }


CSV_UPV_2023 = ("EXPEDIENTE;TERCERO;CONCEPTO ECONÓMICO; IMPORTE ;CENTRO DIRECTIVO\r\n"
                "PR000284971/195713;ASOCIACION PERITOS;otros gastos diversos; 150,00   ;GESTIÓN DE LA INVESTIGACIÓN\r\n"
                "PA000286606/206567;KAURI SPORTWEAR SL;Suministros de vestuario; 28,28   ;DEPORTES\r\n").encode("cp1252")
CSV_UPV_2026 = ("﻿EXPEDIENTE;DURACION;TERCERO;OBJETO;IMP_ADJUDICA\r\n"
                "PB000470959/097608;18;K-TUIN SISTEMAS INFORMATICOS, S.A.;Equipos;1841,44\r\n").encode("utf-8")


def publicar_upv(portal, acumulado="ley_2026-07-20.csv"):
    portal.ckan = paquetes_upv(acumulado)
    portal.fichero(UPV_BASE + "a/resource/1/download/2024-143914_menores_2024-03-04_trimestral.csv", CSV_UPV_2023)
    portal.fichero(UPV_BASE + f"b/resource/3/download/{acumulado}", CSV_UPV_2026)


# ---------------------------------------------------------------------------
# Buscador de València
# ---------------------------------------------------------------------------

ESPACIO = "_contratos_menores_ContratosMenoresPortlet_INSTANCE_TEST_"
ACCION = (M.VLC_PAGINA + "?p_p_id=contratos_menores_ContratosMenoresPortlet_INSTANCE_TEST&p_p_lifecycle=1"
          f"&{ESPACIO}javax.portlet.action=buscarContrato&p_auth=abc123")
FICHA = (M.VLC_PAGINA + "?p_p_id=contratos_menores_ContratosMenoresPortlet_INSTANCE_TEST&p_p_lifecycle=0"
         f"&{ESPACIO}jspPage=%2Fcontrato.jsp&{ESPACIO}numContrato=")
ESTADO_FILA = {"ADJUDICADOS": "ADJUDICADO", "MODIFICADOS": "EJECUTADO", "RESUELTOS": "RESUELTO"}


def contrato(num, fecha, estado="ADJUDICADOS", nif="B98316326"):
    """fecha: un día (se graba a las 10:00) o un datetime (p.ej. a las 00:00)."""
    momento = fecha if isinstance(fecha, datetime) else datetime.combine(fecha, hora(10, 0))
    return {"num": num, "fecha": momento, "estado": estado, "nif": nif,
            "objeto": f'SUMINISTRO "{num}" PARA EL SERVICIO DE BIBLIOTECAS MUNICIPALES'}


class FakeBuscador:
    """Portlet del buscador: la acción del formulario lleva un p_auth y solo
    vale con la cookie de la sesión que lo dio; compara fecha y hora con
    desde <= momento <= hasta a las 00:00 (como el real: de 'hasta' solo entran
    los grabados a las 00:00) y devuelve como mucho maxResultados filas."""

    def __init__(self):
        self.contratos = []
        self.consultas = []          # (estado, desde, hasta, filas devueltas)
        self.sesiones = 0
        self.fallo = None            # None | 'aviso' | int
        self.formularios = 0

    def pagina(self, desde="27/06/2026", hasta="27/09/2026", situacion="", filas=(), error="ocultar"):
        self.formularios += 1
        cuerpo = "".join(
            f'<tr><td>{html.escape(c["objeto"])}</td><td><a class="a-titulo-cm" href="{html.escape(FICHA + c["num"])}" '
            f'title="Enlace a {c["objeto"]}"><span>{html.escape(c["objeto"][:45])}</span></a></td>'
            f'<td>SUMINISTROS</td><td>{c["fecha"]:%d/%m/%Y}</td><td class="dt-body-right">2.473,63 €</td>'
            f'<td>{c["nif"]}</td><td>{ESTADO_FILA[c["estado"]]}</td><td></td></tr>' for c in filas)
        return f"""<!DOCTYPE html><html lang="es-ES"><head><title>Buscador de contratos menores</title></head><body>
<form action="{html.escape(ACCION)}" class="form" data-fm-namespace="{ESPACIO}" id="{ESPACIO}buscarContrato"
 method="post" name="{ESPACIO}buscarContrato" enctype="multipart/form-data" id="formContrato">
<input class="field" id="{ESPACIO}formDate" name="{ESPACIO}formDate" type="hidden" value="{1790546567622 + self.formularios}">
<input id="{ESPACIO}fch-inicio" name="{ESPACIO}fechaInicio" type="text" value="{desde}">
<input id="{ESPACIO}fch-fin" name="{ESPACIO}fechaFin" type="text" value="{hasta}">
</form>
<table id="tablaContratos" class="display"><thead><tr><th>Objeto del contrato</th><th>Objeto del contrato</th>
<th>Tipo</th><th>Fecha</th><th>Importe (sin IVA)</th><th>NIF/CIF/NIE adjudicatario</th><th>Estado</th>
<th>Fecha de pago</th></tr></thead>
<tbody>{cuerpo}</tbody></table>
<div class="alert-notifications"><strong class="lead">Nodo: sweb{self.formularios % 7}:8080:</strong></div>
<script type="text/javascript">$(document).ready(function(){{ var tipo = ""; var situacion = "{situacion}";
var mostrarError = "{error}"; }});</script></body></html>""".encode("utf-8")

    def get(self):
        self.sesiones += 1
        return FakeResponse(body=self.pagina(), cookies={"JSESSIONID": f"S{self.sesiones}", "COOKIE_SUPPORT": "true"})

    def post(self, url, files, cookies):
        if url != ACCION or (cookies or {}).get("JSESSIONID") != f"S{self.sesiones}":
            return FakeResponse(body=self.pagina())          # sin sesión: el formulario por defecto
        if isinstance(self.fallo, int):
            return FakeResponse(status=self.fallo, body=b"<html>error</html>")
        campos = {nombre[len(ESPACIO):]: valor[1] for nombre, valor in files}
        estado = campos["selectEstado"]
        desde = datetime.strptime(campos["fechaInicio"], "%d/%m/%Y")
        hasta = datetime.strptime(campos["fechaFin"], "%d/%m/%Y")
        filas = sorted((c for c in self.contratos if c["estado"] == estado and desde <= c["fecha"] <= hasta),
                       key=lambda c: (c["fecha"], c["num"]), reverse=True)[:int(campos["maxResultados"])]
        self.consultas.append((estado, desde.date(), hasta.date(), len(filas)))
        if self.fallo == "aviso":
            return FakeResponse(body=self.pagina(campos["fechaInicio"], campos["fechaFin"], estado, error="mostrar"))
        return FakeResponse(body=self.pagina(campos["fechaInicio"], campos["fechaFin"], estado, filas))

    def pedidas(self, estado=None):
        return [(d, h) for e, d, h, _ in self.consultas if estado in (None, e)]


# ---------------------------------------------------------------------------
# Portal simulado
# ---------------------------------------------------------------------------

class FakeResponse:
    def __init__(self, status=200, json_data=None, body=b"", cookies=None):
        self.status_code = status
        self._json = json_data
        self.content = body
        self.headers = {}
        self.cookies = cookies or {}

    def json(self):
        if self._json is None:
            raise ValueError("no es JSON")
        return self._json

    @property
    def text(self):
        return self.content.decode("utf-8", "replace")

    def iter_content(self, chunk_size=8192):
        yield self.content

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False


class FakePortal:
    """urls: url -> bytes | código HTTP (las rutas se comparan sin escapar); las
    URL http:// dan 403, como el proxy; ckan: {paquete: result} o un código."""

    def __init__(self):
        self.urls = {}
        self.ckan = {}
        self.llamadas = []
        self.buscador = FakeBuscador()

    @staticmethod
    def _clave(url):
        return unquote(url)

    def fichero(self, url, contenido):
        self.urls[self._clave(url)] = contenido

    def get(self, url, params=None, headers=None, timeout=None, stream=False):
        self.llamadas.append(url)
        if url.startswith("http://"):
            return FakeResponse(status=403, body=b"<html>Forbidden (HTTP plano)</html>")
        if url == M.VLC_PAGINA:
            return self.buscador.get()
        if url == M.UPV_CKAN:
            if isinstance(self.ckan, int):
                return FakeResponse(status=self.ckan)
            paquete = self.ckan.get(params["id"])
            if paquete is None:
                return FakeResponse(status=404, json_data={"success": False})
            return FakeResponse(json_data={"success": True, "result": paquete})
        cuerpo = self.urls.get(self._clave(url), 404)
        if isinstance(cuerpo, int):
            return FakeResponse(status=cuerpo, body=b"<html><body>Not Found</body></html>")
        return FakeResponse(body=cuerpo)

    def post(self, url, files=None, data=None, cookies=None, headers=None, timeout=None):
        self.llamadas.append(url)
        return self.buscador.post(url, files, cookies)

    def pedidas(self, url):
        return sum(1 for u in self.llamadas if self._clave(u) == self._clave(url))


@pytest.fixture
def portal(monkeypatch):
    fake = FakePortal()
    monkeypatch.setattr(requests, "get", fake.get)
    monkeypatch.setattr(requests, "post", fake.post)
    monkeypatch.setattr(time, "sleep", lambda s: None)
    return fake


def _ejecutar(salida, fuentes, *args):
    return M.main(["--salida", str(salida), "--fuentes", fuentes, *args])


def _log(salida):
    return (salida / "raw" / "descarga_log.txt").read_text(encoding="utf-8")


def _manifiesto(salida):
    return json.loads((salida / "raw" / "_manifiesto.json").read_text(encoding="utf-8"))


def _v(serie):
    return [None if pd.isna(v) else v for v in serie]


# ---------------------------------------------------------------------------
# Universitat de València
# ---------------------------------------------------------------------------

def test_uv_descubre_los_ficheros_de_la_pagina_y_los_pide_por_https(portal, tmp_path):
    publicar_uv(portal)
    assert _ejecutar(tmp_path, "uv") == 0

    # Los enlaces http:// se piden por https:// (HTTP plano da 403)
    assert all(not u.startswith("http://") for u in portal.llamadas)
    assert portal.pedidas(UV + "2022/1T/2022_1T_CM_PORTALTRANSPARENCIA.xls") == 1
    # Un trimestre comentado en el HTML no se pide; los PDF con su hoja tampoco
    assert portal.pedidas(UV + "2026/3er_trimestre_26/CM_2026_3_cas.xlsx") == 0
    assert portal.pedidas(UV + "2023/4_trimestre_2023/OTROS GASTOS_TRIMESTRE4_2023.pdf") == 0

    df = pd.read_parquet(tmp_path / "universitat_valencia.parquet")
    por_fichero = df.groupby("_archivo_origen").first()
    assert sorted(por_fichero.index) == sorted(f"universitat_valencia/{r}" for r in [
        "2026/1er_trimestre_26/CM_2026_1_cas.xlsx", "2026/2do_trimestre_26/CM_2T_cas.xlsx",
        "2026/1er_trimestre_26/GM 2026_1 tabla dinamica cas.xlsx",
        "2023/4_trimestre_2023/OTROS GASTOS_TRIMESTRE4_2023.xlsx",
        "2022/1T/2022_1T_CM_PORTALTRANSPARENCIA.xls", "LEY_TRANSPARENCIA_menores_4trimestre.xls"])
    fila = por_fichero.loc["universitat_valencia/2026/2do_trimestre_26/CM_2T_cas.xlsx"]
    # Grupo de la fila (rowspan de la anterior) y año del título "EJERCICIO 2026"
    assert (fila["_periodo"], fila["_anio"], fila["_trimestre"], fila["_categoria"]) == (
        "1. Contratos menores 2n Trimestre", "2026", "2", "contratos menores")
    gm = por_fichero.loc["universitat_valencia/2026/1er_trimestre_26/GM 2026_1 tabla dinamica cas.xlsx"]
    assert (gm["_periodo"], gm["_trimestre"], gm["_categoria"]) == ("2. Gastos menores 1er Trimestre", "1",
                                                                   "gastos menores")
    otros = por_fichero.loc["universitat_valencia/2023/4_trimestre_2023/OTROS GASTOS_TRIMESTRE4_2023.xlsx"]
    assert (otros["_anio"], otros["_trimestre"], otros["_categoria"]) == ("2023", "4", "otros gastos")
    viejo = por_fichero.loc["universitat_valencia/LEY_TRANSPARENCIA_menores_4trimestre.xls"]
    assert (viejo["_anio"], viejo["_trimestre"]) == ("2015", "4")
    assert viejo["_fuente"] == UV + "LEY_TRANSPARENCIA_menores_4trimestre.xls"

    # Todas las celdas como texto, filas de título en _titulo_hoja y el pie como fila
    tabla = pq.read_table(tmp_path / "universitat_valencia.parquet")
    assert all(pa.types.is_string(f.type) for f in tabla.schema if f.name != "_en_ultima_descarga")
    cm = df[df["_archivo_origen"].str.endswith("CM_2026_1_cas.xlsx")]
    assert cm["IDENTIFICADOR CONTRATO"].tolist() == ["2026SU05097CM-cas"]
    assert cm["IMPORTE ADJUDICACIÓN (IVA EXCL.)"].tolist() == ["5991.22"]
    assert cm["FECHA ADJUDICACIÓN"].tolist() == ["2026-03-25"]
    assert "FECHA ADJUDICACIÓN DESDE: 01/01/2026" in cm["_titulo_hoja"].iloc[0]
    x22 = df[df["_archivo_origen"].str.endswith("2022_1T_CM_PORTALTRANSPARENCIA.xls")]
    assert _v(x22["Identificador"]) == ["2022 011354 SU-cm", None]
    assert x22["Unidad funcional"].iloc[1].startswith("Información publicada")
    # Originales tal cual, con su ruta publicada
    raw = tmp_path / "raw" / "universitat_valencia"
    assert (raw / "2023/4_trimestre_2023/OTROS GASTOS_TRIMESTRE4_2023.xlsx").exists()
    assert (raw / "_enlaces.json").exists() and (raw / "_pagina.html").read_bytes() == pagina_uv("cas")
    assert "2 PDF no descargados" not in _log(tmp_path) and "1 PDF no descargados" in _log(tmp_path)


def test_uv_tabla_dinamica_se_leen_los_registros_de_su_cache(portal, tmp_path):
    publicar_uv(portal)
    assert _ejecutar(tmp_path, "uv") == 0
    df = pd.read_parquet(tmp_path / "universitat_valencia.parquet")
    gm = df[df["_archivo_origen"].str.endswith("GM 2026_1 tabla dinamica cas.xlsx")]

    resumen = gm[gm["_tabla_dinamica"] == "resumen"]
    assert resumen["UNIDAD FUNCIONAL"].tolist() == ["Rectorat", "B46669693", "Biblioteca", "Total general"]
    registros = gm[gm["_tabla_dinamica"] == "registros"]
    assert len(registros) == 3
    assert registros["UNIDAD FUNCIONAL"].tolist() == ["Rectorat", "Biblioteca", "Rectorat"]   # valores compartidos
    assert registros["NIF ADJUDICATARIO"].tolist() == ["B46669693", "A79486833", "B46669693"]
    assert registros["OBJETO"].tolist() == ["Suministro de bridas para urnas", 'Estores "opacos"', "Lámpara"]
    assert registros["IMPORTE ADJUDICACIÓN (IVA EXCL.)"].tolist() == ["18.56", "5991.22", "100"]
    assert _v(registros["PLAZO EJECUCIÓN\n(meses)"]) == [None, "0.42", "12"]                 # _x000a_ y <m/>
    assert registros["FECHA ADJUDICACIÓN"].tolist() == ["2026-03-03", "2026-02-11", "2026-01-30 10:30:00"]
    assert "Meses" not in gm.columns                               # agrupación: no está en los registros
    assert set(registros["_hoja"]) == {"pivotCacheDefinition1 (origen: Hoja1!A1:G4)"}
    assert "3 registros en la caché de la tabla dinámica" in _log(tmp_path)


def test_uv_la_pagina_valenciana_solo_aporta_sus_propios_ficheros(portal, tmp_path):
    publicar_uv(portal)
    assert _ejecutar(tmp_path, "uv") == 0

    # Lo que enlazan las dos páginas se baja una vez (serie en castellano)
    assert portal.pedidas(UV + "2023/4_trimestre_2023/OTROS GASTOS_TRIMESTRE4_2023.xlsx") == 1
    assert portal.pedidas(UV + "2026/1er_trimestre_26/CM_2026_1_val.xlsx") == 1
    va = pd.read_parquet(tmp_path / "universitat_valencia_va.parquet")
    assert sorted(set(va["_archivo_origen"])) == [
        "universitat_valencia_va/2026/1er_trimestre_26/CM_2026_1_val.xlsx",
        "universitat_valencia_va/2026/1er_trimestre_26/GM 2026_1 tabla dinamica val.xlsx",
        "universitat_valencia_va/2026/2do_trimestre_26/CM_2T_val.xlsx"]
    es = pd.read_parquet(tmp_path / "universitat_valencia.parquet")
    assert not es["_archivo_origen"].str.contains(r"[_ ]val\.xlsx$").any()
    assert "enlace mal formado (no se puede pedir): http://Z:" in _log(tmp_path)


def test_uv_fichero_que_la_pagina_deja_de_enlazar_queda_retirado(portal, tmp_path):
    publicar_uv(portal)
    assert _ejecutar(tmp_path, "uv") == 0
    SLEEP_REAL(1.1)
    portal.urls[M.UV_PAGINA] = pagina_uv("cas", sin=("2026/2do_trimestre_26/CM_2T_cas.xlsx",))
    assert _ejecutar(tmp_path, "uv") == 0

    df = pd.read_parquet(tmp_path / "universitat_valencia.parquet")
    quitado = df["_archivo_origen"] == "universitat_valencia/2026/2do_trimestre_26/CM_2T_cas.xlsx"
    assert quitado.sum() == 1 and not df.loc[quitado, "_en_ultima_descarga"].any()
    assert df.loc[~quitado, "_en_ultima_descarga"].all()
    man = _manifiesto(tmp_path)
    assert man["universitat_valencia/2026/2do_trimestre_26/CM_2T_cas.xlsx"]["publicado"] is False
    assert (tmp_path / "raw/universitat_valencia/2026/2do_trimestre_26/CM_2T_cas.xlsx").exists()


@pytest.mark.parametrize("pagina", [503, b"<html><body>Sin enlaces</body></html>"])
def test_uv_pagina_caida_o_sin_enlaces_no_retira_nada(portal, tmp_path, pagina):
    publicar_uv(portal)
    assert _ejecutar(tmp_path, "uv") == 0
    antes = pd.read_parquet(tmp_path / "universitat_valencia.parquet")
    SLEEP_REAL(1.1)
    portal.urls[M.UV_PAGINA] = pagina
    assert _ejecutar(tmp_path, "uv", "--comprobar-todo") == 1
    despues = pd.read_parquet(tmp_path / "universitat_valencia.parquet")
    assert len(despues) == len(antes) and despues["_en_ultima_descarga"].all()
    assert pd.read_parquet(tmp_path / "universitat_valencia_va.parquet")["_en_ultima_descarga"].all()
    assert "sin la página en castellano no se puede saber" in _log(tmp_path)


# ---------------------------------------------------------------------------
# Diputación de Alicante y UMH (secciones de la página)
# ---------------------------------------------------------------------------

def test_diputacion_solo_la_seccion_de_menores(portal, tmp_path):
    publicar_diputacion(portal)
    portal.urls[M.DIPA_PAGINA_VA] = pagina_diputacion().replace(b"CONTRATOS MENORES", b"CONTRACTES MENORS")
    assert _ejecutar(tmp_path, "alicante") == 0      # el enlace roto (404) es un aviso, no un error

    assert portal.pedidas(DIPA + "RP-1447-Contratos-No-Menores-gestionados-por-Diputacion-2025-CAS.xlsx") == 0
    assert portal.pedidas(DIPA + "CAS-Contratos-arrendam-INMUEBLES-anualidad-2022.xlsx") == 0
    assert portal.pedidas(DIPA + "2016-Contratos-Menores-Primer-Trimestre-Vers26112018.pdf") == 0
    assert portal.pedidas(DIPA + "1CORPORACION20232027/Eco-Presup/Contratos/RP-922-Adjudic-a-31-12-2024-cas.xlsx") == 1
    log = _log(tmp_path)
    assert "enlace roto (HTTP 404)" in log and "RP-922-Adjudic-a-31-12-2024-cas" in log
    assert "PDF sin hoja de cálculo del mismo nombre" in log and "Instruccion-contratos-menores.pdf" in log

    df = pd.read_parquet(tmp_path / "diputacion_alicante.parquet")
    ligate = df[df["_archivo_origen"] == "diputacion_alicante/RP-1628-Adjudic-a-30-06-2026-cas.xlsx"]
    assert ligate["TERCERO"].tolist() == ["ARQUES MORANTE, MARIA JESUS", "LYRECO ESPAÑA, S.A."]
    assert ligate["ADJUDICADO (sin IVA)"].tolist() == ["1500", "357.19"]
    titulo = ligate["_titulo_hoja"].iloc[0]
    assert "Versión núm. 1: 11 de agosto de 2026" in titulo and "Contratos registrados del 01/04/2026 a 30/06/2026" in titulo
    assert _v([ligate["_anio"].iloc[0], ligate["_trimestre"].iloc[0]]) == ["2026", None]
    t16 = df[df["_archivo_origen"].str.contains("2016-Contratos-Menores")]
    assert (t16["_anio"].iloc[0], t16["_trimestre"].iloc[0]) == ("2016", "1")         # no 2018 ("Vers26112018")
    assert t16["_periodo"].iloc[0] == "Relación de contratos menores. Primer trimestre 2016:"
    assert t16["Fecha"].tolist() == ["2016-02-19"]
    # La página valenciana enlaza los mismos ficheros: su serie no tiene nada
    assert not (tmp_path / "diputacion_alicante_va.parquet").exists()


def test_umh_documento_de_la_placsp_sin_extension(portal, tmp_path):
    publicar_umh(portal)
    assert _ejecutar(tmp_path, "umh") == 0

    assert portal.pedidas(UMH + "2026/09/DOC20260717091255BASADOS-2T-2026.xlsx") == 0   # otra sección
    assert portal.pedidas(UMH + "2025/02/RR_00414_2025.xlsx") == 0
    rel = "universidad_miguel_hernandez/contrataciondelestado.es/77f3c628-ec33.xlsx"
    assert (tmp_path / "raw" / rel).exists()
    assert _manifiesto(tmp_path)[rel]["url"] == PLACSP_DOC
    df = pd.read_parquet(tmp_path / "universidad_miguel_hernandez.parquet")
    doc = df[df["_archivo_origen"] == rel]
    assert doc["N.I.F. del proveedor"].tolist() == ["B81152217"]
    assert (doc["_anio"].iloc[0], doc["_trimestre"].iloc[0]) == ("2019", "1")
    assert set(df["_anio"]) == {"2019", "2026"}


# ---------------------------------------------------------------------------
# UPV (CKAN)
# ---------------------------------------------------------------------------

def test_upv_descarga_los_csv_de_los_paquetes_y_retira_los_sustituidos(portal, tmp_path):
    publicar_upv(portal)
    assert _ejecutar(tmp_path, "upv") == 0
    assert portal.pedidas(UPV_BASE + "a/resource/2/download/2024-143916_menores_2024-02-15_trimestral.xlsx") == 0
    df = pd.read_parquet(tmp_path / "universitat_politecnica_valencia.parquet")
    d23 = df[df["_anio"] == "2023"]
    assert d23["CONCEPTO ECONÓMICO"].tolist() == ["otros gastos diversos", "Suministros de vestuario"]   # cp1252
    assert d23[" IMPORTE "].tolist() == [" 150,00   ", " 28,28   "]
    assert d23["_periodo"].iloc[0] == "(Año 2023) Contratos menores"
    assert df.loc[df["_anio"] == "2026", "TERCERO"].tolist() == ["K-TUIN SISTEMAS INFORMATICOS, S.A."]
    assert (tmp_path / "raw/universitat_politecnica_valencia/contratos-menores/_paquete.json").exists()

    # El acumulado de 2026 se sustituye por otro fichero: el anterior queda retirado
    SLEEP_REAL(1.1)
    publicar_upv(portal, acumulado="ley_2026-10-20.csv")
    assert _ejecutar(tmp_path, "upv") == 0
    df = pd.read_parquet(tmp_path / "universitat_politecnica_valencia.parquet")
    viejo = df["_archivo_origen"].str.endswith("ley_2026-07-20.csv")
    assert viejo.sum() == 1 and not df.loc[viejo, "_en_ultima_descarga"].any()
    assert df.loc[~viejo, "_en_ultima_descarga"].all() and len(df) == 4


def test_upv_ckan_caido_no_retira_nada(portal, tmp_path):
    publicar_upv(portal)
    assert _ejecutar(tmp_path, "upv") == 0
    SLEEP_REAL(1.1)
    portal.ckan = 503
    assert _ejecutar(tmp_path, "upv", "--comprobar-todo") == 1
    assert pd.read_parquet(tmp_path / "universitat_politecnica_valencia.parquet")["_en_ultima_descarga"].all()


# ---------------------------------------------------------------------------
# Buscador del Ajuntament de València
# ---------------------------------------------------------------------------

def _valencia(salida, anio, *args):
    return _ejecutar(salida, "valencia", "--desde", str(anio), "--hasta", str(anio), *args)


def test_valencia_consulta_por_meses_y_estados_con_la_fecha_final_excluida(portal, tmp_path):
    b = portal.buscador
    b.contratos = [contrato("019042016000001", date(2016, 12, 29)),                 # anterior a --desde
                   contrato("014012017000001", date(2017, 1, 31)),                  # último día del mes
                   contrato("014012017000002", date(2017, 2, 1)),
                   contrato("014012017000003", datetime(2017, 3, 1, 0, 0)),         # día 1 a las 00:00
                   contrato("023102017000204", date(2017, 5, 10), "MODIFICADOS", nif="*****607X"),
                   contrato("019052017000103", date(2017, 11, 3), "RESUELTOS")]
    assert _valencia(tmp_path, 2017) == 0

    # Una consulta de control con lo anterior, doce meses de adjudicados y un año
    # de modificados y de resueltos; cada mes del día 1 al día 1 del siguiente
    assert b.pedidas("ADJUDICADOS") == [(date(2000, 1, 1), date(2017, 1, 1))] + [
        (date(2017, m, 1), date(2017 + (m == 12), m % 12 + 1, 1)) for m in range(1, 13)]
    assert b.pedidas("MODIFICADOS") == [(date(2000, 1, 1), date(2017, 1, 1)), (date(2017, 1, 1), date(2018, 1, 1))]
    assert b.sesiones == 1                          # una sesión (cookie + p_auth) para todas las consultas

    df = pd.read_parquet(tmp_path / "ajuntament_valencia.parquet")
    assert sorted(set(df["_num_contrato"])) == sorted(c["num"] for c in b.contratos)
    # El del día 1 a las 00:00 lo devuelven las consultas de febrero y de marzo: dos filas
    assert sorted(df.loc[df["_num_contrato"] == "014012017000003", "_mes"]) == ["02", "03"]
    assert len(df) == len(b.contratos) + 1
    enero = df[df["_num_contrato"] == "014012017000001"].iloc[0]
    assert (enero["Fecha"], enero["_anio"], enero["_mes"], enero["_estado_consulta"]) == (
        "31/01/2017", "2017", "01", "ADJUDICADOS")
    assert (enero["_consulta_desde"], enero["_consulta_hasta"]) == ("2017-01-01", "2017-02-01")
    assert enero["_archivo_origen"] == "ajuntament_valencia/ADJUDICADOS/2017/2017-01.html"
    assert enero["Objeto del contrato"] == 'SUMINISTRO "014012017000001" PARA EL SERVICIO DE BIBLIOTECAS MUNICIPALES'
    assert enero["Objeto del contrato.1"] == 'SUMINISTRO "014012017000001" PARA EL SERVICIO'
    assert enero["Objeto del contrato.1 (enlace)"] == FICHA + "014012017000001"
    assert pd.isna(enero["Fecha de pago"])
    previo = df[df["_num_contrato"] == "019042016000001"].iloc[0]
    assert _v([previo["_periodo"], previo["_anio"], previo["_consulta_desde"]]) == ["anteriores-2017", None,
                                                                                   "2000-01-01"]
    mod = df[df["_estado_consulta"] == "MODIFICADOS"].iloc[0]
    assert _v([mod["Estado"], mod["NIF/CIF/NIE adjudicatario"], mod["_mes"]]) == ["EJECUTADO", "*****607X", None]
    # La respuesta de cada consulta se guarda tal cual
    raw = tmp_path / "raw" / "ajuntament_valencia"
    assert b"014012017000001" in (raw / "ADJUDICADOS/2017/2017-01.html").read_bytes()
    assert (raw / "RESUELTOS/2017.html").exists() and (raw / "ADJUDICADOS/anteriores-2017.html").exists()

    # Año cerrado: una segunda ejecución no vuelve a consultar
    consultas = len(b.consultas)
    assert _valencia(tmp_path, 2017) == 0
    assert len(b.consultas) == consultas


def test_valencia_un_mes_con_500_filas_se_pide_por_dias(portal, tmp_path):
    b = portal.buscador
    marzo = [contrato(f"0140{CERRADO}{i:06d}", date(CERRADO, 3, 1) + timedelta(days=i % 31)) for i in range(520)]
    b.contratos = marzo + [contrato("014019990000001", date(CERRADO, 4, 2))]
    assert _valencia(tmp_path, CERRADO) == 0

    dias = [(date(CERRADO, 3, d), date(CERRADO, 3, d) + timedelta(days=1)) for d in range(1, 32)]
    assert [p for p in b.pedidas("ADJUDICADOS") if p[0].month == 3] == [
        (date(CERRADO, 3, 1), date(CERRADO, 4, 1))] + dias
    df = pd.read_parquet(tmp_path / "ajuntament_valencia.parquet")
    en_marzo = df[df["_mes"] == "03"]
    assert len(en_marzo) == 520 and en_marzo["_num_contrato"].is_unique       # ni cortado en 500 ni duplicado
    assert set(en_marzo["_consulta_hasta"]) == {(d + timedelta(days=1)).isoformat() for d, _ in dias}
    assert len(df) == 521
    raw = tmp_path / "raw" / "ajuntament_valencia"
    assert (raw / f"_troceadas/ADJUDICADOS/{CERRADO}/{CERRADO}-03.html").exists()   # la respuesta cortada, tal cual
    assert not (raw / f"ADJUDICADOS/{CERRADO}/{CERRADO}-03.html").exists()
    assert (raw / f"ADJUDICADOS/{CERRADO}/{CERRADO}-03/{CERRADO}-03-31.html").exists()
    assert "1 troceadas por llegar a 500 filas" in _log(tmp_path)

    # Año cerrado: se vuelven a recorrer las partes sin consultar nada
    consultas = len(b.consultas)
    assert _valencia(tmp_path, CERRADO) == 0
    assert len(b.consultas) == consultas


def test_valencia_registro_retirado_se_conserva_y_una_misma_tabla_no_crea_version(portal, tmp_path):
    b = portal.buscador
    b.contratos = [contrato("A", date(CERRADO, 6, 3)), contrato("B", date(CERRADO, 6, 20)),
                   contrato("C", date(CERRADO, 7, 1))]
    assert _valencia(tmp_path, CERRADO) == 0
    junio = tmp_path / "raw" / "ajuntament_valencia" / "ADJUDICADOS" / f"{CERRADO}" / f"{CERRADO}-06.html"
    SLEEP_REAL(1.1)
    assert _valencia(tmp_path, CERRADO, "--comprobar-todo") == 0
    assert len(M.versiones(junio)) == 1        # otra sesión, otro p_auth y otro nodo, pero la misma tabla

    SLEEP_REAL(1.1)
    b.contratos = [c for c in b.contratos if c["num"] != "B"]
    assert _valencia(tmp_path, CERRADO, "--comprobar-todo") == 0
    df = pd.read_parquet(tmp_path / "ajuntament_valencia.parquet")
    assert df.set_index("_num_contrato")["_en_ultima_descarga"].to_dict() == {"A": True, "B": False, "C": True}
    assert len(M.versiones(junio)) == 2


@pytest.mark.parametrize("fallo", ["aviso", 503])
def test_valencia_respuesta_invalida_reintenta_y_no_retira_nada(portal, tmp_path, fallo):
    b = portal.buscador
    b.contratos = [contrato("A", date(CERRADO, 6, 3)), contrato("B", date(CERRADO, 6, 20))]
    assert _valencia(tmp_path, CERRADO) == 0
    SLEEP_REAL(1.1)
    b.fallo = fallo
    sesiones = b.sesiones
    assert _valencia(tmp_path, CERRADO, "--comprobar-todo") == 1
    if fallo == "aviso":
        assert b.sesiones > sesiones + 1           # cada respuesta no válida abre una sesión nueva
    df = pd.read_parquet(tmp_path / "ajuntament_valencia.parquet")
    assert len(df) == 2 and df["_en_ultima_descarga"].all()
    junio = tmp_path / "raw" / "ajuntament_valencia" / "ADJUDICADOS" / f"{CERRADO}" / f"{CERRADO}-06.html"
    assert len(M.versiones(junio)) == 1 and b"B" in junio.read_bytes()
    assert f"ajuntament_valencia ADJUDICADOS {CERRADO}-06" in _log(tmp_path)


def test_valencia_sin_la_cookie_de_sesion_no_hay_resultados(portal, tmp_path, monkeypatch):
    """El p_auth solo vale con la cookie de la sesión que lo dio: sin ella el
    portlet devuelve el formulario por defecto y la consulta no es válida."""
    b = portal.buscador
    b.contratos = [contrato("A", date(CERRADO, 6, 3))]
    real = M.Buscador.abrir

    def sin_cookies(self):
        real(self)
        self.cookies = {}
    monkeypatch.setattr(M.Buscador, "abrir", sin_cookies)
    assert _valencia(tmp_path, CERRADO) == 1
    assert "la respuesta no es la de la consulta" in _log(tmp_path)


# ---------------------------------------------------------------------------
# Utilidades y CLI
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("texto, trimestre", [
    ("1. Contratos menores 1er Trimestre", "1"), ("2. Gastos menores 2n Trimestre", "2"),
    ("4t Trimestre", "4"), ("1. Contratos menores 4º Trimestre", "4"),
    ("1r - 2n - 3er trimestre Otros gastos 2024", None), ("Otros gastos Año 2024", None),
    ("- Otros gastos 2º trimestre 2023 (xls)", "2"), ("Relación de contratos menores. Primer trimestre 2016:", "1"),
    ("Tercer y cuarto trimestre 2020", None), ("Contratos Menores 2º trimestre del ejercicio 2026", "2"),
    ("MENORES_TRIMESTRE4_2023.xlsx", "4"), ("LEY_TRANSPARENCIA_menores_3trimestre jul_sept_2015_ok.xls", "3"),
    ("Contratos menores adjudicados a 30 de junio de 2026 (LIGATE)", None),
])
def test_trimestre_de(texto, trimestre):
    assert M.trimestre_de(texto) == trimestre


def test_categoria_y_url():
    assert M.categoria_de("3. Altres Despeses 1er Trimestre") == "otros gastos"
    assert M.categoria_de("", "GM_2025_2T_tabla_dinamica_cas.xlsx") == "gastos menores"
    assert M.categoria_de("4r trimestre 2023 (xls)") is None
    assert M.normalizar_url("http://www.uv.es/a/4%c2%ba_trim/Año 2024.xlsx", ("www.uv.es",)) == \
        "https://www.uv.es/a/4%C2%BA_trim/A%C3%B1o%202024.xlsx"
    assert M.ruta_local(PLACSP_DOC, "/files/") == "contrataciondelestado.es/77f3c628-ec33"


def test_cli(portal, tmp_path, monkeypatch):
    publicar_umh(portal)
    monkeypatch.setattr(sys, "argv", [str(SCRIPT), "--salida", str(tmp_path), "--fuentes", "umh"])
    with pytest.raises(SystemExit) as salida:
        runpy.run_path(str(SCRIPT), run_name="__main__")
    assert salida.value.code == 0
    assert (tmp_path / "universidad_miguel_hernandez.parquet").exists()


def test_fuente_desconocida(tmp_path):
    with pytest.raises(SystemExit):
        M.main(["--salida", str(tmp_path), "--fuentes", "castello"])


def test_salida_por_defecto_en_el_repo():
    assert M.SALIDA == REPO_ROOT / "ccaa_valencia_menores"


def test_quote_de_espacios_en_rutas_publicadas():
    # Las rutas locales conservan el nombre publicado (con espacios y º)
    url = M.normalizar_url(UV + quote("2024/4º_trimestre_24/Año 2024CM cas.xlsx"))
    assert M.ruta_local(url, M.SERIES_PAGINA["universitat_valencia"]["base"]) == "2024/4º_trimestre_24/Año 2024CM cas.xlsx"
