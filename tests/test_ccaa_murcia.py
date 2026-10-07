"""Tests offline de scripts/ccaa_murcia.py.

Los ficheros anuales de datosabiertos.carm.es / transparencia.carm.es y el
catálogo CKAN regional se simulan con un ``requests.get`` falso; los CSV y
XLSX se generan en el propio test.
"""

import csv
import importlib.util
import io
import json
import runpy
import sys
import time
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
import pytest
import requests

REPO_ROOT = Path(__file__).resolve().parents[1]
SCRIPT = REPO_ROOT / "scripts" / "ccaa_murcia.py"
SLEEP_REAL = time.sleep


def _cargar():
    spec = importlib.util.spec_from_file_location("ccaa_murcia", SCRIPT)
    modulo = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(modulo)
    return modulo


M = _cargar()
ANIO = datetime.now().year          # el script decide qué refrescar con el año real


def url_carm(anio):
    return f"https://datosabiertos.carm.es/odata/transparencia/contratosOD{anio}.csv"


def url_menores(anio):
    return f"https://datosabiertos.carm.es/odata/Hacienda/CONTRA_ContratosMenores_{anio}.csv"


DIR_SMS = "https://transparencia.carm.es/wres/transparencia/doc/Sector_Publico/SMS/Contratos_menores/"
# Nombres publicados en la página de sector público (2026-09-27): cada año con otro
SMS_PUBLICADOS = {**{f"PT_SMS_{t}T2019.xlsx": 2019 for t in range(1, 5)},
                  "Contratos_menores_SMS_2020.xlsx": 2020, "SMS_Contratos_Menores_2021.xlsx": 2021,
                  **{f"Contratos_Menores_SMS_{a}.xlsx": a for a in (2022, 2023, 2024)},
                  "SMS_Contratos_menores_2025.xlsx": 2025}


def url_sms(anio, ext="xlsx"):
    return f"{DIR_SMS}Contratos_menores_SMS_{anio}.{ext}"


def pagina_sector_publico(nombres):
    """HTML como el de la página real: enlaces absolutos y relativos, y otros enlaces."""
    enlaces = [f'<a href="{DIR_SMS if i % 2 else "/wres/transparencia/doc/Sector_Publico/SMS/Contratos_menores/"}'
               f'{n}">Contratos menores</a>' for i, n in enumerate(nombres)]
    return ("<html><body><a href=\"/web/transparencia/convenios-sms\">Convenios</a>"
            '<a href="/wres/transparencia/doc/Sector_Publico/SMS/Convenios/Convenios_2024.xlsx">x</a>'
            + "".join(enlaces) + "</body></html>").encode("utf-8")


# ---------------------------------------------------------------------------
# Portal simulado
# ---------------------------------------------------------------------------

class FakeResponse:
    def __init__(self, status=200, json_data=None, body=b""):
        self.status_code = status
        self._json = json_data
        self._body = body
        self.headers = {}

    def json(self):
        return self._json

    @property
    def text(self):
        return self._body.decode("utf-8", "replace")

    def iter_content(self, chunk_size=8192):
        yield self._body

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False


class FakePortal:
    """urls: url -> bytes | código HTTP; ckan: lista de paquetes (o código)."""

    def __init__(self):
        self.urls = {}
        self.ckan = []
        self.llamadas = []

    def get(self, url, params=None, headers=None, timeout=None, stream=False):
        self.llamadas.append((url, dict(params or {})))
        if url == M.URL_CKAN:
            if isinstance(self.ckan, int):
                return FakeResponse(status=self.ckan)
            ini, n = int(params["start"]), int(params["rows"])
            return FakeResponse(json_data={"success": True, "result": {
                "count": len(self.ckan), "results": self.ckan[ini:ini + n]}})
        cuerpo = self.urls.get(url, 404)
        if isinstance(cuerpo, int):
            return FakeResponse(status=cuerpo, body=b"<html><body>Not Found</body></html>")
        return FakeResponse(body=cuerpo)

    def pedidas(self, url):
        return sum(1 for u, _ in self.llamadas if u == url)


def _xlsx(hojas):
    libro = openpyxl.Workbook()
    libro.remove(libro.active)
    for nombre, filas in hojas.items():
        hoja = libro.create_sheet(nombre)
        for fila in filas:
            hoja.append(fila)
    datos = io.BytesIO()
    libro.save(datos)
    return datos.getvalue()


def _csv_carm(anio, objeto="Señalización – 5 €"):
    return f"NUMERO;OBJETO;IMPORTE;CIF\n{anio}-0001;{objeto};1.000,00;B0012345\n".encode("cp1252")


def _publicar_confirmados(portal):
    for anio in range(2019, 2024):
        portal.urls[url_carm(anio)] = _csv_carm(anio)
    for anio in range(2022, 2026):
        portal.urls[url_menores(anio)] = f"Expediente,Importe\nM{anio}-1,100\n".encode("utf-8")
    for nombre, anio in SMS_PUBLICADOS.items():
        portal.urls[DIR_SMS + nombre] = _xlsx({"Hoja1": [["Expediente", "Importe"], [f"S{nombre[:-5]}-1", 50]]})
    portal.urls[M.URL_SECTOR_PUBLICO] = pagina_sector_publico(SMS_PUBLICADOS)


@pytest.fixture
def portal(monkeypatch):
    fake = FakePortal()
    monkeypatch.setattr(requests, "get", fake.get)
    monkeypatch.setattr(time, "sleep", lambda s: None)
    _publicar_confirmados(fake)
    return fake


def _ejecutar(salida, *args, desde=2019):
    return M.main(["--salida", str(salida), "--desde", str(desde), "--hasta", str(ANIO), *args])


def _log(salida):
    return (salida / "raw" / "descarga_log.txt").read_text(encoding="utf-8")


def _v(serie):
    return [None if pd.isna(v) else v for v in serie]


# ---------------------------------------------------------------------------
# Años: sondeo, 404 y confirmados
# ---------------------------------------------------------------------------

def test_sondea_todos_los_anios_y_anota_los_no_publicados(portal, tmp_path):
    assert _ejecutar(tmp_path, desde=2017) == 0

    for anio in range(2017, ANIO + 1):
        assert portal.pedidas(url_carm(anio)) == 1
        assert portal.pedidas(url_menores(anio)) == 1
    # El SMS no se sondea por año: se baja lo que enlaza la página, con su nombre
    assert portal.pedidas(M.URL_SECTOR_PUBLICO) == 1
    assert all(portal.pedidas(DIR_SMS + n) == 1 for n in SMS_PUBLICADOS)
    assert portal.pedidas(url_sms(2021)) == 0
    log = _log(tmp_path)
    rangos = M.Resumen._rangos
    assert f"contratos_carm: {rangos([2017, 2018] + list(range(2024, ANIO + 1)))}" in log
    assert f"contratos_menores_carm: {rangos(list(range(2017, 2022)) + list(range(2026, ANIO + 1)))}" in log


def test_parquet_por_serie_con_todos_los_anios_como_texto(portal, tmp_path):
    assert _ejecutar(tmp_path) == 0

    assert sorted(p.name for p in tmp_path.glob("*.parquet")) == [
        "contratos_carm.parquet", "contratos_menores_carm.parquet", "contratos_menores_sms.parquet"]
    tabla = pq.read_table(tmp_path / "contratos_carm.parquet")
    assert all(pa.types.is_string(f.type) for f in tabla.schema if f.name != "_en_ultima_descarga")
    df = tabla.to_pandas()
    assert df["_anio_fichero"].tolist() == [str(a) for a in range(2019, 2024)]
    assert df["NUMERO"].tolist() == [f"{a}-0001" for a in range(2019, 2024)]
    assert set(df["OBJETO"]) == {"Señalización – 5 €"}                 # cp1252
    assert set(df["IMPORTE"]) == {"1.000,00"} and set(df["CIF"]) == {"B0012345"}
    assert df["_fuente"].tolist() == [url_carm(a) for a in range(2019, 2024)]
    assert df["_archivo_origen"].tolist() == [f"contratos_carm/contratosOD{a}.csv" for a in range(2019, 2024)]
    assert list(df.columns[:4]) == ["NUMERO", "OBJETO", "IMPORTE", "CIF"]
    # Originales tal cual
    assert (tmp_path / "raw" / "contratos_carm" / "contratosOD2019.csv").read_bytes() == _csv_carm(2019)
    manifiesto = json.loads((tmp_path / "raw" / "_manifiesto.json").read_text(encoding="utf-8"))
    assert manifiesto["contratos_carm/contratosOD2019.csv"]["url"] == url_carm(2019)


def test_anio_confirmado_que_no_esta_es_un_error(portal, tmp_path):
    del portal.urls[url_carm(2021)]
    assert _ejecutar(tmp_path) == 1
    assert "contratos_carm 2021: año publicado según las fuentes" in _log(tmp_path)


def test_pagina_html_en_un_anio_inexistente_es_no_publicado(portal, tmp_path):
    portal.urls[url_carm(2018)] = b"<!DOCTYPE html><html><body>No encontrado</body></html>"
    assert _ejecutar(tmp_path, desde=2018) == 0
    assert not (tmp_path / "raw" / "contratos_carm" / "contratosOD2018.csv").exists()
    assert "contratos_carm: 2018" in _log(tmp_path)


def test_solo_se_vuelven_a_pedir_el_anio_actual_y_el_anterior(portal, tmp_path):
    portal.urls[url_menores(ANIO)] = b"Expediente,Importe\nX,1\n"
    _ejecutar(tmp_path)
    _ejecutar(tmp_path)
    assert portal.pedidas(url_carm(2019)) == 1
    assert portal.pedidas(url_menores(ANIO)) == 2
    assert portal.pedidas(url_menores(ANIO - 1)) == 2
    _ejecutar(tmp_path, "--comprobar-todo")
    assert portal.pedidas(url_carm(2019)) == 2
    assert not (tmp_path / "raw" / "contratos_carm" / "_historico").exists()   # nada cambió


def test_catalogo_ckan_aporta_anios_y_urls_y_lista_el_resto(portal, tmp_path, monkeypatch):
    monkeypatch.setattr(M, "FILAS_CKAN", 1)
    portal.ckan = [
        {"id": "1", "name": "contratos-carm-historico", "resources": [
            {"name": "Contratos 2015", "url": url_carm(2015)},
            {"name": "Contratos 2016", "url": "https://otro.carm.es/descargas/contratosOD2016.csv"}]},
        {"id": "2", "name": "contratos-trabajo", "resources": [
            {"name": "Contratos de trabajo registrados", "url": "https://x.carm.es/paro.csv"}]},
    ]
    portal.urls[url_carm(2015)] = _csv_carm(2015)
    portal.urls["https://otro.carm.es/descargas/contratosOD2016.csv"] = _csv_carm(2016)

    assert _ejecutar(tmp_path) == 0            # --desde 2019: 2015 y 2016 solo están en el catálogo

    assert portal.pedidas(M.URL_CKAN) == 2
    df = pd.read_parquet(tmp_path / "contratos_carm.parquet")
    assert df["_anio_fichero"].tolist() == [str(a) for a in [2015, 2016, 2019, 2020, 2021, 2022, 2023]]
    assert df["_fuente"].iloc[1] == "https://otro.carm.es/descargas/contratosOD2016.csv"
    assert (tmp_path / "raw" / "contratos_carm" / "contratosOD2016.csv").exists()
    assert len(json.loads((tmp_path / "raw" / "catalogo_ckan.json").read_text(encoding="utf-8"))) == 2
    assert "contratos-trabajo: Contratos de trabajo registrados" in _log(tmp_path)


def test_catalogo_ckan_caido_no_impide_la_descarga(portal, tmp_path):
    portal.ckan = 500
    assert _ejecutar(tmp_path) == 0
    assert "catálogo CKAN regional" in _log(tmp_path)


# ---------------------------------------------------------------------------
# Excel del SMS
# ---------------------------------------------------------------------------

def test_excel_con_titulo_varias_hojas_y_tipos_se_guarda_como_texto(portal, tmp_path):
    portal.urls[url_sms(2020)] = _xlsx({
        "2020": [["Contratos menores SMS 2020"], [], ["Expediente", "Importe", "Fecha", "Hora", "CP", None],
                 ["00123", 1234.5, datetime(2020, 1, 15), datetime(2020, 2, 1, 10, 30), "03001"],
                 ["S-2", 10, None, None, None], [None, None, None, None, None, None], ["S-3", 7.0, None, None, "NA"]],
        "Anexo": [["Expediente", "Importe"], ["S-4", 0.1]],
    })
    assert _ejecutar(tmp_path) == 0

    df = pd.read_parquet(tmp_path / "contratos_menores_sms.parquet")
    df = df[df["_anio_fichero"] == "2020"]
    assert list(df.columns[:5]) == ["Expediente", "Importe", "Fecha", "Hora", "CP"]
    assert df["Expediente"].tolist() == ["00123", "S-2", "S-3", "S-4"]
    assert df["Importe"].tolist() == ["1234.5", "10", "7", "0.1"]
    assert _v(df["Fecha"]) == ["2020-01-15", None, None, None]
    assert _v(df["Hora"]) == ["2020-02-01 10:30:00", None, None, None]
    assert _v(df["CP"]) == ["03001", None, "NA", None]
    assert df["_hoja"].tolist() == ["2020", "2020", "2020", "Anexo"]
    log = _log(tmp_path)
    assert "2 filas antes de la cabecera" in log and "Contratos menores SMS 2020" in log
    assert "2 hojas con datos" in log


def test_sms_se_toman_los_ficheros_que_enlaza_la_pagina(portal, tmp_path):
    # Antes solo se encontraba 2020 (plantilla Contratos_menores_SMS_{año})
    assert _ejecutar(tmp_path) == 0
    raw = tmp_path / "raw" / "contratos_menores_sms"
    assert sorted(p.name for p in raw.iterdir() if p.is_file()) == sorted(SMS_PUBLICADOS)
    df = pd.read_parquet(tmp_path / "contratos_menores_sms.parquet")
    assert sorted(df["_anio_fichero"]) == sorted(str(a) for a in SMS_PUBLICADOS.values())
    assert (df["_anio_fichero"] == "2019").sum() == 4          # cuatro trimestres
    assert set(df["_fuente"]) == {DIR_SMS + n for n in SMS_PUBLICADOS}
    man = json.loads((tmp_path / "raw" / "_manifiesto.json").read_text(encoding="utf-8"))
    assert man["contratos_menores_sms/PT_SMS_3T2019.xlsx"]["anio"] == 2019

    # La página deja de enlazar un trimestre: queda retirado y sus filas se conservan
    SLEEP_REAL(1.1)
    portal.urls[M.URL_SECTOR_PUBLICO] = pagina_sector_publico(
        [n for n in SMS_PUBLICADOS if n != "PT_SMS_4T2019.xlsx"])
    assert _ejecutar(tmp_path, "--comprobar-todo") == 0
    df = pd.read_parquet(tmp_path / "contratos_menores_sms.parquet")
    cuarto = df["_archivo_origen"] == "contratos_menores_sms/PT_SMS_4T2019.xlsx"
    assert cuarto.sum() == 1 and not df.loc[cuarto, "_en_ultima_descarga"].any()
    assert df.loc[~cuarto, "_en_ultima_descarga"].all()
    man = json.loads((tmp_path / "raw" / "_manifiesto.json").read_text(encoding="utf-8"))
    assert man["contratos_menores_sms/PT_SMS_4T2019.xlsx"]["publicado"] is False


@pytest.mark.parametrize("pagina", [503, b"<html><body>Sin enlaces</body></html>"])
def test_sms_pagina_caida_o_sin_enlaces_no_retira_nada(portal, tmp_path, pagina):
    assert _ejecutar(tmp_path) == 0
    antes = pd.read_parquet(tmp_path / "contratos_menores_sms.parquet")
    SLEEP_REAL(1.1)
    portal.urls[M.URL_SECTOR_PUBLICO] = pagina
    assert _ejecutar(tmp_path, "--comprobar-todo") == 1
    despues = pd.read_parquet(tmp_path / "contratos_menores_sms.parquet")
    assert len(despues) == len(antes) and despues["_en_ultima_descarga"].all()
    assert "contratos_menores_sms" in _log(tmp_path)


def test_hoja_sin_cabecera_conserva_su_primera_fila(portal, tmp_path):
    # Como la Hoja1 del SMS de 2021: códigos de acreedor y NIF, sin cabecera
    portal.urls[DIR_SMS + "SMS_Contratos_Menores_2021.xlsx"] = _xlsx({
        "Contratos menores 2021": [["Doc.compr.", " Cif", "Adjudicatario Nombre"],
                                   [4431033264, "B87867446", "BIOMARIN"]],
        "Hoja1": [[1000002621, "34789348P"], [1000003683, "A08015646"]],
    })
    assert _ejecutar(tmp_path) == 0
    df = pd.read_parquet(tmp_path / "contratos_menores_sms.parquet")
    df = df[df["_anio_fichero"] == "2021"]
    aux = df[df["_hoja"] == "Hoja1"]
    assert aux["columna_1"].tolist() == ["1000002621", "1000003683"]
    assert aux["columna_2"].tolist() == ["34789348P", "A08015646"]
    assert "1000002621" not in df.columns
    assert df.loc[df["_hoja"] != "Hoja1", " Cif"].tolist() == ["B87867446"]
    assert "Hoja1]: sin fila de cabecera" in _log(tmp_path)


def test_xls_binario(tmp_path):
    xlwt = pytest.importorskip("xlwt")
    pytest.importorskip("xlrd")
    libro = xlwt.Workbook()
    hoja = libro.add_sheet("Contratos")
    fecha = xlwt.easyxf(num_format_str="DD/MM/YYYY")
    for c, v in enumerate(["Expediente", "Importe", "Fecha"]):
        hoja.write(0, c, v)
    hoja.write(1, 0, "00123")
    hoja.write(1, 1, 1234.5)
    hoja.write(1, 2, datetime(2021, 3, 4), fecha)
    hoja.write(2, 0, "X")
    hoja.write(2, 1, 10)
    ruta = tmp_path / "c.xls"
    libro.save(str(ruta))
    df, _ = M.leer_tabla(ruta)
    assert df["Expediente"].tolist() == ["00123", "X"]
    assert df["Importe"].tolist() == ["1234.5", "10"]
    assert _v(df["Fecha"]) == ["2021-03-04", None]


# ---------------------------------------------------------------------------
# Sesgo del superviviente
# ---------------------------------------------------------------------------

def test_registro_retirado_y_modificado_se_conservan(portal, tmp_path):
    portal.urls[url_menores(ANIO)] = b"Expediente,Importe\nA,10\nB,20\nC,30\n"
    assert _ejecutar(tmp_path) == 0
    SLEEP_REAL(1.1)          # la versión nueva tiene que llevar otra fecha (resolución: 1 s)
    portal.urls[url_menores(ANIO)] = b"Expediente,Importe\nA,10\nC,35\n"
    assert _ejecutar(tmp_path) == 0

    df = pd.read_parquet(tmp_path / "contratos_menores_carm.parquet")
    actual = df[df["_anio_fichero"] == str(ANIO)]
    assert actual[["Expediente", "Importe", "_en_ultima_descarga"]].values.tolist() == [
        ["A", "10", True], ["B", "20", False], ["C", "30", False], ["C", "35", True]]
    otros = df[df["_anio_fichero"] != str(ANIO)]
    assert len(otros) == 4 and otros["_en_ultima_descarga"].all()
    raw = tmp_path / "raw" / "contratos_menores_carm" / f"CONTRA_ContratosMenores_{ANIO}.csv"
    assert len(M.versiones(raw)) == 2


def test_anio_retirado_por_el_portal_conserva_sus_filas(portal, tmp_path):
    portal.urls[url_menores(ANIO)] = b"Expediente,Importe\nA,10\n"
    assert _ejecutar(tmp_path) == 0
    SLEEP_REAL(1.1)
    del portal.urls[url_menores(ANIO)]
    assert _ejecutar(tmp_path) == 0

    df = pd.read_parquet(tmp_path / "contratos_menores_carm.parquet")
    assert df.loc[df["_anio_fichero"] == str(ANIO), "_en_ultima_descarga"].tolist() == [False]
    assert (tmp_path / "raw" / "contratos_menores_carm" / f"CONTRA_ContratosMenores_{ANIO}.csv").exists()
    assert f"contratos_menores_carm {ANIO}: el portal ya no lo sirve" in _log(tmp_path)


@pytest.mark.parametrize("fallo", [503, b"<html><body>Mantenimiento</body></html>"])
def test_descarga_fallida_no_pierde_nada(portal, tmp_path, fallo):
    portal.urls[url_menores(ANIO)] = b"Expediente,Importe\nA,10\n"
    assert _ejecutar(tmp_path) == 0
    antes = pd.read_parquet(tmp_path / "contratos_menores_carm.parquet")
    SLEEP_REAL(1.1)
    portal.urls[url_menores(ANIO)] = fallo
    assert _ejecutar(tmp_path) == 1            # no se sabe si se ha retirado: error, nada cambia
    despues = pd.read_parquet(tmp_path / "contratos_menores_carm.parquet")
    assert despues["_en_ultima_descarga"].all()
    assert despues["Expediente"].tolist() == antes["Expediente"].tolist()
    raw = tmp_path / "raw" / "contratos_menores_carm" / f"CONTRA_ContratosMenores_{ANIO}.csv"
    assert raw.read_bytes() == b"Expediente,Importe\nA,10\n" and len(M.versiones(raw)) == 1


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def test_cli(portal, tmp_path, monkeypatch):
    monkeypatch.setattr(sys, "argv", [str(SCRIPT), "--salida", str(tmp_path), "--desde", "2019"])
    with pytest.raises(SystemExit) as salida:
        runpy.run_path(str(SCRIPT), run_name="__main__")
    assert salida.value.code == 0
    assert (tmp_path / "contratos_carm.parquet").exists()


def test_salida_por_defecto_en_el_repo():
    assert M.SALIDA == REPO_ROOT / "ccaa_murcia"


def test_comilla_literal_no_se_traga_registros(tmp_path):
    """Una comilla al principio de un campo que no se cierra hasta varias
    líneas después (registros enteros) es literal: se conservan todas las
    filas y la comilla queda en el texto (comun.lectura_csv)."""
    cabecera = ";".join(f"c{i}" for i in range(5))
    ruta = tmp_path / "contratosOD2024.csv"
    ruta.write_bytes((f'{cabecera}\n1;"Obra A;org;10;x\n2;Obra B;org;20;y\n3;Obra "C";org;30;z\n').encode("utf-8"))
    df, avisos = M.leer_tabla(ruta)
    assert df["c0"].tolist() == ["1", "2", "3"]
    assert df["c1"].tolist() == ['"Obra A', "Obra B", 'Obra "C"']
    assert any("comillas literales" in a for a in avisos)


@pytest.mark.parametrize("texto, codificacion, esperada", [
    # contratosOD 2014-2018 de la CARM: cp850 ('Ó' = 0xE0); leído como cp1252 daba 'NEGOCIACIàN'
    ("COD;PROCEDIMIENTO;ORGANO\n1;NEGOCIACIÓN SIN PUBLICIDAD;ÁREA DE SALUD IX\n2;ADJUDICACIÓN;CONSEJERÍA\n",
     "cp850", "cp850"),
    ("COD;OBJETO\n1;Señalización – 5 € de la Consejería\n", "cp1252", "cp1252"),
    ("COD;OBJETO\n1;Póliza año\x81\n", "latin-1", "latin-1"),     # 0x81 no existe en cp1252
    ("COD;OBJETO\n1;Señalización – 5 €\n", "utf-8", "utf-8"),
    ("COD;OBJETO\n1;SIN ACENTOS\n", "cp1252", "utf-8"),              # ASCII: vale UTF-8
])
def test_detectar_codificacion(tmp_path, texto, codificacion, esperada):
    ruta = tmp_path / "f.csv"
    ruta.write_bytes(texto.encode(codificacion))
    assert M._detectar_codificacion(ruta) == esperada
    df, _ = M.leer_tabla(ruta)
    assert df.iloc[0, 1] == texto.splitlines()[1].split(";")[1]      # el texto publicado, bien leído


def test_carm_en_cp850_se_lee_bien(portal, tmp_path):
    portal.urls[url_carm(2019)] = ("NUMERO;OBJETO;IMPORTE;CIF\n2019-0001;NEGOCIACIÓN SIN PUBLICIDAD - ÁREA IX;"
                                   "1.000,00;B0012345\n").encode("cp850")
    assert _ejecutar(tmp_path) == 0
    df = pq.read_table(tmp_path / "contratos_carm.parquet").to_pandas()
    assert df.loc[df["_anio_fichero"] == "2019", "OBJETO"].tolist() == ["NEGOCIACIÓN SIN PUBLICIDAD - ÁREA IX"]
    assert set(df.loc[df["_anio_fichero"] != "2019", "OBJETO"]) == {"Señalización – 5 €"}   # cp1252, como antes


def test_un_arreglo_de_lectura_llega_a_las_filas_ya_guardadas(portal, tmp_path, monkeypatch):
    """Regla 3: el Parquet se construye con el código actual desde todas las versiones del crudo.
    Antes solo se aplicaban las versiones posteriores al Parquet anterior y un arreglo de lectura
    (p.ej. la codificación cp850 de los contratosOD 2014-2018) no llegaba a las filas guardadas."""
    assert _ejecutar(tmp_path) == 0
    antes = pq.read_table(tmp_path / "contratos_carm.parquet").to_pandas()
    leer = M.leer_tabla

    def lectura_arreglada(ruta, *a, **k):
        df, avisos = leer(ruta, *a, **k)
        if "OBJETO" in df.columns:
            df = df.assign(OBJETO=df["OBJETO"].str.upper())
        return df, avisos

    monkeypatch.setattr(M, "leer_tabla", lectura_arreglada)
    assert _ejecutar(tmp_path, "--solo-parquet") == 0
    despues = pq.read_table(tmp_path / "contratos_carm.parquet").to_pandas()
    assert len(despues) == len(antes)                                   # sin duplicar filas
    assert set(despues["OBJETO"]) == {"SEÑALIZACIÓN – 5 €"}             # el arreglo llega a todas
    assert despues["_en_ultima_descarga"].all()
    assert despues["_primera_descarga"].tolist() == antes["_primera_descarga"].tolist()   # historia intacta


# ---------------------------------------------------------------------------
# CSV del exportador JSON de datosabiertos.carm.es/odata (líneas reales de los
# crudos de producción, 2026-09-28): comillas escapadas con barra, cortes de línea
# escritos como \n cada 80 caracteres y restos de la lista JSON al final
# ---------------------------------------------------------------------------

CAB_OD = ["ejercicio", "codinscripcion", "tipocontrato", "procedimiento", "objeto", "importlicitacion",
          "importadjudicacion", "adjudicatariodescripcion", "adjudicatariocodigo", "fechaformalizacion",
          "duracion", "organo", "cpvcodigo", "cpvdescripcion", "nummodificaciones", "codigoOrgano", "fechaInicio"]
# El texto de los campos va como lo publica el portal (r'...': \" y \n son barra + carácter)
OD_1262 = [r'2019', r'1262/2019', r'OBRAS', r'OTROS',
           r'EMERGENCIA DE OBRAS EN LOS CENTROS IES \"MENARGUEZ COSTA\", CEIP \"BIENVENIDO CONEJ\nERO\" Y CEIP '
           r'\"PETRA SÁNCHEZ ROLLÁN\" Y SERVICIOS DE LIMPIEZA EN 29 INSTITUTOS DE L\nA REGIÓN DE MURCIA.',
           r'273928.75', r'273928.75', r'CONSTRUCCIONES URDECON SA', r'A30032205', r'20/09/2019', r'3 meses y 8 días',
           r'C. DE EDUCACION Y CULTURA', r'45200000', r'TRABAJOS GENERALES DE CONSTRUCCIÓN', r'0', r'15', r'20/09/2019']
OD_649 = [r'2019', r'649/2019', r'PRIVADOS', r'NEGOCIACIÓN SIN PUBLICIDAD',
          r'CONTRATACIÓN DEL SERVICIO DE MANTENIMIENTO, PREVENTIVO Y CORRECTIVO, A REALIZAR\nEN LOS EQUIPOS INSTAL'
          r'ADOS EN EL ANIMALARIO DEL INSTITUTO MURCIANO DE INVESTIGACI\nÓN BIOSANITARIA ¿VIRGEN DE LA ARRIXACA\"',
          r'26136', r'26133.97', r'ANTONIO MATACHANA S.A.', r'A08238578', r'02/04/2019', r'12 meses',
          r'FUNDACIÓN PARA LA FORMACIÓN E INVESTIGACIÓN SANITARIAS', r'50000000', r'SERVICIOS DE REPARACIÓN', r'0',
          r'91', r'02/04/2019']
OD_ESCAPES = [r'2019', r'1068/2019', r'PRIVADOS', r'NEGOCIACIÓN SIN PUBLICIDAD',
              r'OBRAS EN EL CENTRO DE SALUD DE\tLA MANGA PARA LA EJECUCIÓ\u0093N', r'100', r'90', r'EMPRESA, S.L.',
              r'B30000000', r'01/01/2019', r'1 mes', r'C. DE SALUD', r'45000000', r'OBRAS', r'0', r'14', r'01/01/2019']
# contratosOD2023: 1636/2023 es uno de los 33 contratos que la simulación del revisor duplicaba; 464/2023 tiene
# saltos de línea del texto dentro de la ventana de 80 ('REFORMA PARCIA\nL DE')
OD_1636_2023 = [r'2023', r'1636/2023', r'OBRAS', r'NEGOCIACIÓN SIN PUBLICIDAD',
                r'OBRAS DE EMERGENCIA EN EL CEIP \"HERMANOS ISIDORO Y SANTA FLORENTINA\" (CASA EL NI\nÑO) DE '
                r'CARTAGENA', r'33500', r'33500', r'RESTAURALIA CARTAGO, S.L.', r'B30802565', r'05/12/2023',
                r'1 meses', r'C. DE EDUCACIÓN, FORMACIÓN PROFESIONAL Y EMPLEO', r'45200000',
                r'TRABAJOS GENERALES DE CONSTRUCCIÓN DE INMUEBLES Y OBRAS DE INGENIERÍA CIVIL', r'0', r'15',
                r'06/12/2023']
OD_464_2023 = [r'2023', r'464/2023', r'SERVICIOS', r'OTROS',
               r'BASADO 30 -\nCOORDINACIÓN DE SEGURIDAD Y SALUD DURANTE PROYECTO DE REFORMA PARCIA\nL DE\n'
               r'ACONDICIONAMIENTO DE SALA DE EXPLORACIÓN PARA SU RENOVACIÓN Y POSTERIOR\nINS\nTALACIÓN DE UNA '
               r'GAMMACÁMARA EN EL HOSPITAL GENERAL UNIVERSITARIO SANTA\nLUCÍA EN\nCARTAGENA.', r'968', r'597.74',
               r'JEVEAL INGENIERIA Y SEGURIDAD SL', r'B73416067', r'13/04/2023', r'37 días',
               r'SERVICIOS CENTRALES (SMS)', r'71317200', r'SERVICIOS DE SALUD Y SEGURIDAD', r'0', r'00', r'13/04/2023']
# La fila que la simulación del revisor añadía a contratosOD2023: una barra que no es escape de JSON, como
# las de los menores de 2024-2025 (aquí con un corte de 80 delante, para ver que el texto sin cortes sigue)
OD_BARRA_SUELTA = [r'2023', r'9999/2023', r'SERVICIOS', r'ABIERTO',
                   r'ADAPTACIÓN DE LAS INSTALACIONES DE PROTECCIÓN CONTRA INCENDIOS DEL EDIFICIO SEDE\n DE LA '
                   r'CONSEJERÍA AL RD 390\2021', r'1000', r'900', r'EMPRESA PRUEBA SL', r'B00000000', r'01/12/2023',
                   r'1 mes', r'C. DE PRUEBA', r'45000000', r'OBRAS', r'0', r'99', r'01/12/2023']


def _csv_json(cabecera, filas):
    """CSV como los del exportador: todo entre comillas, separado por comas y con LF."""
    return "\n".join(",".join(f'"{v}"' for v in fila) for fila in [cabecera] + filas).encode("utf-8")


def test_comillas_escapadas_con_barra_no_corren_columnas(tmp_path):
    """contratosOD 2019-2023: el portal escapa las comillas del texto con barra ('\\"'). Antes el
    campo se cerraba en esa comilla: 1262/2019 se corría a la derecha (adjudicatario '273928.75' y
    _columna_extra_1) y 649/2019 a la izquierda sin ninguna marca (el adjudicatario en
    importadjudicacion)."""
    ruta = tmp_path / "contratosOD2019.csv"
    ruta.write_bytes(_csv_json(CAB_OD, [OD_1262, OD_649, OD_ESCAPES]))
    df, avisos = M.leer_tabla(ruta)
    assert not [c for c in df.columns if c.startswith("_columna_extra")]
    assert df["codinscripcion"].tolist() == ["1262/2019", "649/2019", "1068/2019"]
    assert df["importlicitacion"].tolist() == ["273928.75", "26136", "100"]
    assert df["importadjudicacion"].tolist() == ["273928.75", "26133.97", "90"]
    assert df["adjudicatariodescripcion"].tolist() == ["CONSTRUCCIONES URDECON SA", "ANTONIO MATACHANA S.A.",
                                                      "EMPRESA, S.L."]
    assert df["adjudicatariocodigo"].tolist() == ["A30032205", "A08238578", "B30000000"]
    assert df["fechaInicio"].tolist() == ["20/09/2019", "02/04/2019", "01/01/2019"]
    # \" es la comilla del texto (como '""'); el resto de secuencias se sirven tal cual
    assert df["objeto"].iloc[0] == (
        r'EMERGENCIA DE OBRAS EN LOS CENTROS IES "MENARGUEZ COSTA", CEIP "BIENVENIDO CONEJ\nERO" Y CEIP '
        r'"PETRA SÁNCHEZ ROLLÁN" Y SERVICIOS DE LIMPIEZA EN 29 INSTITUTOS DE L\nA REGIÓN DE MURCIA.')
    assert df["objeto"].iloc[1].endswith(r'INVESTIGACI\nÓN BIOSANITARIA ¿VIRGEN DE LA ARRIXACA"')
    assert df["objeto"].iloc[2] == r'OBRAS EN EL CENTRO DE SALUD DE\tLA MANGA PARA LA EJECUCIÓ\u0093N'
    assert any("comillas escapadas con barra" in a for a in avisos)


def test_cortes_de_linea_cada_80_caracteres_van_aparte(tmp_path):
    """El exportador parte el texto en líneas de 80 caracteres (\\n escrito) y recorta el espacio del
    final: un trozo de 80 se une tal cual y uno de 79 con el espacio recortado. La columna original
    no cambia; el texto sin cortes va en _<columna>_sin_cortes (nulo en las filas sin cortes)."""
    ruta = tmp_path / "contratosOD2019.csv"
    ruta.write_bytes(_csv_json(CAB_OD, [OD_1262, OD_649, OD_ESCAPES]))
    df, avisos = M.leer_tabla(ruta)
    assert _v(df["_objeto_sin_cortes"]) == [
        'EMERGENCIA DE OBRAS EN LOS CENTROS IES "MENARGUEZ COSTA", CEIP "BIENVENIDO CONEJERO" Y CEIP '
        '"PETRA SÁNCHEZ ROLLÁN" Y SERVICIOS DE LIMPIEZA EN 29 INSTITUTOS DE LA REGIÓN DE MURCIA.',
        'CONTRATACIÓN DEL SERVICIO DE MANTENIMIENTO, PREVENTIVO Y CORRECTIVO, A REALIZAR EN LOS EQUIPOS '
        'INSTALADOS EN EL ANIMALARIO DEL INSTITUTO MURCIANO DE INVESTIGACIÓN BIOSANITARIA ¿VIRGEN DE LA '
        'ARRIXACA"',
        None]
    assert "\\n" in df["objeto"].iloc[0]                      # el original, tal cual
    assert "_adjudicatariodescripcion_sin_cortes" not in df.columns
    assert any("_objeto_sin_cortes" in a for a in avisos)


@pytest.mark.parametrize("valor, esperado", [
    ("A" * 80 + r"\nB", "A" * 80 + "B"),                              # 80: el corte cae dentro de la palabra
    ("A" * 79 + r"\nB", "A" * 79 + " B"),                             # 79: el carácter 80 era un espacio
    ("A" * 80 + r"\n B", "A" * 80 + " B"),                            # 80 y el espacio en la línea siguiente
    ("Guantes latex" + r"\n- 4", "Guantes latex - 4"),                 # salto de línea del texto original
    ("A" * 78 + "\u00d3" + r"\u0093" + r"\nN", "A" * 78 + "\u00d3" + r"\u0093" + "N"),   # \u0093 cuenta uno
    ("A" * 78 + r"\t" + "B" + r"\nC", "A" * 78 + r"\t" + "BC"),                       # \t también
    (r"\nA", "A"), ("A" * 80 + r"\n", "A" * 80), ("A" + r"\n\n" + "B", "A B"),
    ("SIN CORTES", None), (None, None),
])
def test_texto_sin_cortes(valor, esperado):
    assert M.texto_sin_cortes(valor) == esperado


# Valores reales (crudos de producción, 2026-09-28) con saltos de línea del texto dentro de la ventana de 80
# caracteres: el contador del exportador no se reinicia en ellos (cuentan uno). Contando cada línea por
# separado salía lo que dice cada comentario.
@pytest.mark.parametrize("valor, esperado", [
    # menores 2023: 70 + 1 (salto del texto) + 9 = 80 ('Dichos ma teriales')
    (r"materiales de sensibilización, tales como carteles, pegatinas, pulseras, cubreva\nsos, etc., así como "
     r"la maquetación y edición del protocolo mencionado.\nDichos ma\nteriales se utilizarán para la difusión del "
     r"Protocolo de actuación ante la viole\nncia sexual en locales turísticos y hosteleros",
     "materiales de sensibilización, tales como carteles, pegatinas, pulseras, cubrevasos, etc., así como la "
     "maquetación y edición del protocolo mencionado. Dichos materiales se utilizarán para la difusión del "
     "Protocolo de actuación ante la violencia sexual en locales turísticos y hosteleros"),
    # contratosOD2023, 464/2023 ('REFORMA PARCIA L DE' e 'INS TALACIÓN')
    (r"BASADO 30 -\nCOORDINACIÓN DE SEGURIDAD Y SALUD DURANTE PROYECTO DE REFORMA PARCIA\nL DE\nACONDICIONAMIENTO "
     r"DE SALA DE EXPLORACIÓN PARA SU RENOVACIÓN Y POSTERIOR\nINS\nTALACIÓN DE UNA GAMMACÁMARA EN EL HOSPITAL "
     r"GENERAL UNIVERSITARIO SANTA\nLUCÍA EN\nCARTAGENA.",
     "BASADO 30 - COORDINACIÓN DE SEGURIDAD Y SALUD DURANTE PROYECTO DE REFORMA PARCIAL DE ACONDICIONAMIENTO DE "
     "SALA DE EXPLORACIÓN PARA SU RENOVACIÓN Y POSTERIOR INSTALACIÓN DE UNA GAMMACÁMARA EN EL HOSPITAL GENERAL "
     "UNIVERSITARIO SANTA LUCÍA EN CARTAGENA."),
    # contratosOD2022 ('R UBÉOLA')
    (r"BASADO EN EL ACUERDO MARCO SUMINISTRO DE VACUNAS DE CALENDARIO Y OTRAS PARA LOS\nAÑOS\n2022 Y 2023 (LOTES "
     r"4, 5, 13 Y 14 AM). LOTE 13: VACUNA FRENTE A SARAMPIÓN, R\nUBÉOLA, PAROTIDITIS Y VARICELA",
     "BASADO EN EL ACUERDO MARCO SUMINISTRO DE VACUNAS DE CALENDARIO Y OTRAS PARA LOS AÑOS 2022 Y 2023 (LOTES 4, "
     "5, 13 Y 14 AM). LOTE 13: VACUNA FRENTE A SARAMPIÓN, RUBÉOLA, PAROTIDITIS Y VARICELA"),
    # menores 2023 ('968,0 0 euros')
    (r'Contratación de "Coordinador de seguridad y salud de las obras de cerramiento de\n porche en el IES REY '
     r'CARLOS III de ÁGUILAS".\n- Presupuesto de licitación: 968,0\n0 euros\n- Plazo de ejecución: 2 Meses',
     'Contratación de "Coordinador de seguridad y salud de las obras de cerramiento de porche en el IES REY CARLOS '
     'III de ÁGUILAS". - Presupuesto de licitación: 968,00 euros - Plazo de ejecución: 2 Meses'),
    # menores 2023: el salto del texto tras 'Nitrito sodio 3' es el carácter 80 y el corte, el \n siguiente;
    # tomar el primero por el corte descuadra la cuenta ('Hidr óxido')
    (r"Eter de petróleo 4\nAcido acético glaciar 6\nAcido perclórico 2\nCloruro de amonio\n6\nAcido EDTA 4\n"
     r"Nitrato calcio 4\nCloruro sodio 4\nAcetato sodio 3\nNitrito sodio 3\n\nAcido bórico 5\nMetanol 12\n"
     r"Acetonitrilo 6\nEter dietílico 5\nIoduro potásico 4\nHidr\nóxido de calcio 6",
     "Eter de petróleo 4 Acido acético glaciar 6 Acido perclórico 2 Cloruro de amonio 6 Acido EDTA 4 Nitrato "
     "calcio 4 Cloruro sodio 4 Acetato sodio 3 Nitrito sodio 3 Acido bórico 5 Metanol 12 Acetonitrilo 6 Eter "
     "dietílico 5 Ioduro potásico 4 Hidróxido de calcio 6"),
    # menores 2023: la primera línea se cortó a 70 (espacios recortados): 'General de ...' (66) no cabría tras
    # un salto del texto. Tomado por un salto, la cuenta llega a 137 ('yEducación Per manente')
    (r"Adquisición de 3 monitores para ordenador sobremesa, para la Dirección\nGeneral de Formación Profesional, "
     r"Enseñanzas de Régimen Especial y\nEducación Per\nmanente.",
     "Adquisición de 3 monitores para ordenador sobremesa, para la Dirección General de Formación Profesional, "
     "Enseñanzas de Régimen Especial y Educación Permanente."),
    # Lo mismo con trozos sintéticos
    ("A" * 79 + r"\n\n" + "B" * 79 + r"\nC", "A" * 79 + " " + "B" * 79 + " C"),   # salto en el 80 y un 79 detrás
    ("A" * 70 + r"\n" + "B" * 20 + r"\nC", "A" * 70 + " " + "B" * 20 + " C"),     # 70 + 1 + 20 no cabe: corte
    ("A" * 60 + r"\n" + "B" * 19 + r"\nC", "A" * 60 + " " + "B" * 19 + "C"),      # 60 + 1 + 19 = 80
], ids=["dichos-materiales", "reforma-parcial", "rubeola", "968-euros", "salto-en-el-80", "corte-a-70",
        "sintetico-salto-en-el-80", "sintetico-no-cabe", "sintetico-80-con-salto"])
def test_texto_sin_cortes_con_saltos_del_texto(valor, esperado):
    assert M.texto_sin_cortes(valor) == esperado


def test_barras_que_no_son_escapes_de_json_se_leen_como_siempre(tmp_path):
    """Los menores de 2024-2025 traen barras del texto ('RD 390\\2021'): no es el exportador JSON
    (no traen '\\"' y la mitad o más de sus barras no son secuencias válidas), así que '\\n' no se
    interpreta."""
    ruta = tmp_path / "CONTRA_ContratosMenores_2024.csv"
    ruta.write_bytes('CODEXPEDIENTE,OBJETO_CONTRATO_MENOR,IMPORTE\n2024/1,"Según el RD 390\\2021 y C:\\nuevo",10\n'
                     '2024/2,Normal,20\n'.encode("utf-8"))
    df, avisos = M.leer_tabla(ruta)
    assert df["OBJETO_CONTRATO_MENOR"].tolist() == ["Según el RD 390\\2021 y C:\\nuevo", "Normal"]
    assert [c for c in df.columns if c.startswith("_")] == []
    assert avisos == []                         # ni rastro del exportador en el log ('comillas escapadas...')


@pytest.mark.parametrize("texto, es_json", [
    (r'"OBRAS EN EL CEIP \"X\" Y EN EL IES\nY","A\tB \u0093"', True),     # exportador: todas válidas
    (r'"x\ny","CEIP \"Z\"","ADAPTACIÓN AL RD 390\2021"', True),          # y una barra suelta: la mayoría
    (r'"Según el RD 390\2021 y C:\nuevo"', False),                       # tantas sueltas como válidas
    (r'"Decreto 532\2017","USB-C\A","Conversor\es"', False),             # menores 2024-2025: todas sueltas
    ('"SIN BARRAS"', False),
])
def test_es_texto_json_por_mayoria_de_las_barras(texto, es_json):
    assert M._es_texto_json(texto) is es_json


def test_comillas_con_barra_se_leen_aunque_haya_una_barra_suelta(tmp_path):
    """La barra se decide solo por la presencia de '\\"': antes una barra suelta en otro campo ('RD 390\\2021',
    como en los menores de 2024-2025) devolvía todo el fichero al lector normal y las comillas con barra
    volvían a correr columnas. El texto sin cortes también sigue (la mayoría de las barras son de JSON)."""
    ruta = tmp_path / "contratosOD2023.csv"
    ruta.write_bytes(_csv_json(CAB_OD, [OD_1262, OD_1636_2023, OD_BARRA_SUELTA]))
    df, avisos = M.leer_tabla(ruta)
    assert not [c for c in df.columns if c.startswith("_columna_extra")]
    assert df["codinscripcion"].tolist() == ["1262/2019", "1636/2023", "9999/2023"]
    assert df["importlicitacion"].tolist() == ["273928.75", "33500", "1000"]
    assert df["adjudicatariodescripcion"].tolist() == ["CONSTRUCCIONES URDECON SA", "RESTAURALIA CARTAGO, S.L.",
                                                      "EMPRESA PRUEBA SL"]
    assert df["objeto"].iloc[1] == (r'OBRAS DE EMERGENCIA EN EL CEIP "HERMANOS ISIDORO Y SANTA FLORENTINA" (CASA EL '
                                    r'NI\nÑO) DE CARTAGENA')
    assert df["objeto"].iloc[2] == OD_BARRA_SUELTA[4]                    # la barra suelta, tal cual
    assert _v(df["_objeto_sin_cortes"])[1:] == [
        'OBRAS DE EMERGENCIA EN EL CEIP "HERMANOS ISIDORO Y SANTA FLORENTINA" (CASA EL NIÑO) DE CARTAGENA',
        "ADAPTACIÓN DE LAS INSTALACIONES DE PROTECCIÓN CONTRA INCENDIOS DEL EDIFICIO SEDE DE LA CONSEJERÍA "
        "AL RD 390\\2021"]
    assert any("comillas escapadas con barra" in a for a in avisos)


@pytest.mark.parametrize("barra", [False, True])
@pytest.mark.parametrize("texto", [
    'A,B,C\n"1","x"y,"3"\n',                     # texto pegado tras la comilla de cierre: es del campo
    'A,B,C\n"1","2",\n',                         # separador al final de la línea: un último campo vacío
    'A,B,C\n"1","2","3"\n"4","5",',              # separador al final del texto
    'A,B\n"con ""comillas""","l1\nl2"\n"3",4\r\n',
])
def test_lector_del_exportador_lee_como_el_modulo_csv(texto, barra):
    """Sin barras ni relleno, _registros_json lee como el módulo csv (su docstring): así se leen los menores
    de 2018 (solo restos de la lista al final) y, con la barra decidida por '\\"', cualquier fichero que lo
    traiga. No se pierde lo pegado tras una comilla de cierre ni el último campo vacío. Ningún crudo de producción
    tiene estos casos (comprobado): son la garantía de no perder texto con un fichero mal formado."""
    assert M._registros_json(texto, ",", barra, False)[0] == list(csv.reader(io.StringIO(texto)))


def test_menores_con_comillas_dobladas_y_cortes(tmp_path):
    """Los menores de 2023 traen los cortes \\n del exportador pero doblan las comillas ('""'):
    se leen como siempre (pandas) y el texto sin cortes va aparte."""
    ruta = tmp_path / "CONTRA_ContratosMenores_2023.csv"
    obj = "Contratación de la difusión de varias campañas publicitarias en el ámbito de la" + r"\n" + "salud"
    ruta.write_bytes(('"CODEXPEDIENTE","OBJETO_CONTRATO_MENOR","ADJUDICATARIODESCRIPCION","EJERCICIONUM"\n'
                      f'"2023/003698","{obj}","EDICIONES ""VITALIDAD"" SL","2023"\n'
                      '"2023/007031","Gastos alojamiento","VIAJES SL","2023"\n').encode("utf-8"))
    df, avisos = M.leer_tabla(ruta)
    assert df["ADJUDICATARIODESCRIPCION"].tolist() == ['EDICIONES "VITALIDAD" SL', "VIAJES SL"]
    assert df["OBJETO_CONTRATO_MENOR"].iloc[0] == obj
    assert _v(df["_OBJETO_CONTRATO_MENOR_sin_cortes"]) == [
        "Contratación de la difusión de varias campañas publicitarias en el ámbito de la salud", None]


CAB_MENORES = ["CONSEJERIA-OA-COD", "CONSEJERIA-OA-DESCRIPCION", "ADJUDICATARIO_CIF", "IMPORTE(IVAINCLUIDO)",
               "EJERCICIO", "TRIMESTRE_NUM"]


def test_menores_2018_restos_de_la_lista_json_al_final(tmp_path):
    """CONTRA_ContratosMenores_2018 y 2022 acaban en '"4"]' y una línea ']'. Antes: TRIMESTRE_NUM '4]'
    y una fila con ']' como código de consejería. Ahora el ']' va a _resto_json y la línea ']' sigue
    siendo una fila (solo con _resto_json): ninguna fila ya descargada desaparece."""
    ruta = tmp_path / "CONTRA_ContratosMenores_2018.csv"
    ruta.write_bytes(('"' + '","'.join(CAB_MENORES) + '"\n'
                      '"19","Consejería de Turismo","B73656985","1289.99","2018","4"\n'
                      '"16","Consejería de Empleo","B73335069","208","2018","4"]\n]').encode("utf-8"))
    df, avisos = M.leer_tabla(ruta)
    assert len(df) == 3
    assert _v(df["CONSEJERIA-OA-COD"]) == ["19", "16", None]
    assert _v(df["TRIMESTRE_NUM"]) == ["4", "4", None]
    assert _v(df["_resto_json"]) == [None, "]", "]"]
    assert df.drop(columns="_resto_json").iloc[2].isna().all()
    assert any("restos de la lista JSON" in a for a in avisos)


def test_menores_2017_y_2019_lineas_con_espacios_y_final(tmp_path):
    """Los menores de 2017 y 2019 llevan cada línea rodeada de espacios (' "14",...,"4" '): antes el
    código de consejería era ' "14"' (con las comillas) y el trimestre '4 '. El de 2019 acaba además en
    '"4"","' y una línea ' ' (que no era ni es una fila)."""
    cabecera = '"' + '","'.join(CAB_MENORES) + '"\n'
    for anio, final, resto in (("2017", '"4" \n ', None), ("2019", '"4"","\n ', '","')):
        ruta = tmp_path / f"CONTRA_ContratosMenores_{anio}.csv"
        ruta.write_bytes((cabecera + f' "14","Consejería de Fomento","B73802571","35952.29","{anio}","4" \n'
                          f' "51","IMAS (""ASPRODES"")","B73083461","72,77","{anio}",' + final).encode("utf-8"))
        df, avisos = M.leer_tabla(ruta)
        assert df["CONSEJERIA-OA-COD"].tolist() == ["14", "51"]
        assert df["TRIMESTRE_NUM"].tolist() == ["4", "4"]
        assert df["CONSEJERIA-OA-DESCRIPCION"].tolist() == ["Consejería de Fomento", 'IMAS ("ASPRODES")']
        assert df["IMPORTE(IVAINCLUIDO)"].tolist() == ["35952.29", "72,77"]
        if resto:
            assert _v(df["_resto_json"]) == [None, resto]
        else:
            assert "_resto_json" not in df.columns
        assert any("espacios alrededor" in a for a in avisos)


def test_menores_2021_comillas_con_barra(tmp_path):
    """CONTRA_ContratosMenores_2021 y 2022 escapan las comillas con barra: antes el adjudicatario
    quedaba 'ASOC DE PROMOCION AL DEFECIENTE \\ASPRODES\\""'; los ficheros de 2020 y 2023 publican el
    mismo nombre como 'ASOC DE PROMOCION AL DEFECIENTE "ASPRODES"'. Las filas de 2021 no traen el
    último campo (TRIMESTRENUM), como en el portal."""
    cab = ["UNIDAD", "CPVDESCRIPCION", "ADJUDICATARIOCODIGO", "ADJUDICATARIODESCRIPCION", "VALORCONTABPAGO",
           "EJERCICIONUM", "TRIMESTRENUM"]
    ruta = tmp_path / "CONTRA_ContratosMenores_2021.csv"
    ruta.write_bytes(_csv_json(cab, [
        [r"51", r"HUEVOS", r"G30033146", r'ASOC DE PROMOCION AL DEFECIENTE \"ASPRODES\"', r"131,04", r"2021"],
        [r"14", r"SERVICIOS DE REPARACIÓN Y MANTENIMIENTO DE MAQUINARIA ELÉCTRICA, APARATOS Y EQUI\nPO ASOCIADO",
         r"B30811814", r"GISPERT SL", r"10", r"2021"]]))
    df, _ = M.leer_tabla(ruta)
    assert df["ADJUDICATARIODESCRIPCION"].tolist() == ['ASOC DE PROMOCION AL DEFECIENTE "ASPRODES"', "GISPERT SL"]
    assert df["VALORCONTABPAGO"].tolist() == ["131,04", "10"]
    assert _v(df["TRIMESTRENUM"]) == [None, None]
    assert _v(df["_CPVDESCRIPCION_sin_cortes"]) == [
        None, "SERVICIOS DE REPARACIÓN Y MANTENIMIENTO DE MAQUINARIA ELÉCTRICA, APARATOS Y EQUIPO ASOCIADO"]


def test_exportador_json_de_punta_a_punta(portal, tmp_path):
    """Por el portal simulado: contratosOD2019 y los menores de 2018 llegan al Parquet leídos bien
    (sin _columna_extra_N, con _objeto_sin_cortes y _resto_json) y los originales quedan tal cual en raw/."""
    portal.urls[url_carm(2019)] = _csv_json(CAB_OD, [OD_1262, OD_649])
    menores_2018 = ('"' + '","'.join(CAB_MENORES) + '"\n'
                    '"16","Consejería de Empleo","B73335069","208","2018","4"]\n]').encode("utf-8")
    portal.urls[url_menores(2018)] = menores_2018
    assert _ejecutar(tmp_path, desde=2018) == 0
    df = pq.read_table(tmp_path / "contratos_carm.parquet").to_pandas()
    od = df[df["_anio_fichero"] == "2019"]
    assert od["adjudicatariodescripcion"].tolist() == ["CONSTRUCCIONES URDECON SA", "ANTONIO MATACHANA S.A."]
    assert not [c for c in df.columns if c.startswith("_columna_extra")]
    assert od["_objeto_sin_cortes"].str.contains("INVESTIGACIÓN BIOSANITARIA").tolist() == [False, True]
    assert df.loc[df["_anio_fichero"] != "2019", "_objeto_sin_cortes"].isna().all()
    men = pq.read_table(tmp_path / "contratos_menores_carm.parquet").to_pandas()
    m18 = men[men["_anio_fichero"] == "2018"]
    assert _v(m18["TRIMESTRE_NUM"]) == ["4", None] and _v(m18["_resto_json"]) == ["]", "]"]
    assert m18["_en_ultima_descarga"].all()
    raw = tmp_path / "raw"
    assert (raw / "contratos_carm" / "contratosOD2019.csv").read_bytes() == _csv_json(CAB_OD, [OD_1262, OD_649])
    assert (raw / "contratos_menores_carm" / "CONTRA_ContratosMenores_2018.csv").read_bytes() == menores_2018


def test_una_barra_suelta_en_la_descarga_nueva_no_duplica_contratos(portal, tmp_path):
    """Simulación del revisor con los crudos de producción: contratosOD2023 vuelve a descargarse con una fila más que
    trae una barra suelta ('RD 390\\2021'). Antes esa barra devolvía la versión nueva entera al lector normal:
    los contratos con '\\"' cambiaban, la versión buena quedaba retirada y la corrupta vigente (33 contratos
    duplicados, 15.433 filas en vez de 15.400). Ahora las dos versiones se leen igual: nada se retira ni se
    duplica, solo entra la fila nueva, con su texto sin cortes."""
    primera = _csv_json(CAB_OD, [OD_1636_2023, OD_464_2023])
    portal.urls[url_carm(2023)] = primera
    assert _ejecutar(tmp_path) == 0
    SLEEP_REAL(1.1)                  # la versión nueva tiene que llevar otra fecha (resolución: 1 s)
    portal.urls[url_carm(2023)] = _csv_json(CAB_OD, [OD_1636_2023, OD_464_2023, OD_BARRA_SUELTA])
    assert _ejecutar(tmp_path, "--comprobar-todo") == 0

    crudo = tmp_path / "raw" / "contratos_carm" / "contratosOD2023.csv"
    assert len(M.versiones(crudo)) == 2 and M.versiones(crudo)[0].read_bytes() == primera
    df = pq.read_table(tmp_path / "contratos_carm.parquet").to_pandas()
    assert not [c for c in df.columns if c.startswith("_columna_extra")]
    od = df[df["_anio_fichero"] == "2023"]
    assert od["codinscripcion"].tolist() == ["1636/2023", "464/2023", "9999/2023"]
    assert od["_en_ultima_descarga"].all()
    assert od["adjudicatariodescripcion"].tolist() == ["RESTAURALIA CARTAGO, S.L.", "JEVEAL INGENIERIA Y SEGURIDAD SL",
                                                      "EMPRESA PRUEBA SL"]
    assert od["objeto"].iloc[0].startswith('OBRAS DE EMERGENCIA EN EL CEIP "HERMANOS ISIDORO')
    assert _v(od["_objeto_sin_cortes"]) == [
        'OBRAS DE EMERGENCIA EN EL CEIP "HERMANOS ISIDORO Y SANTA FLORENTINA" (CASA EL NIÑO) DE CARTAGENA',
        "BASADO 30 - COORDINACIÓN DE SEGURIDAD Y SALUD DURANTE PROYECTO DE REFORMA PARCIAL DE ACONDICIONAMIENTO DE "
        "SALA DE EXPLORACIÓN PARA SU RENOVACIÓN Y POSTERIOR INSTALACIÓN DE UNA GAMMACÁMARA EN EL HOSPITAL GENERAL "
        "UNIVERSITARIO SANTA LUCÍA EN CARTAGENA.",
        "ADAPTACIÓN DE LAS INSTALACIONES DE PROTECCIÓN CONTRA INCENDIOS DEL EDIFICIO SEDE DE LA CONSEJERÍA "
        "AL RD 390\\2021"]
    assert od["_primera_descarga"].iloc[0] == od["_primera_descarga"].iloc[1] != od["_primera_descarga"].iloc[2]
