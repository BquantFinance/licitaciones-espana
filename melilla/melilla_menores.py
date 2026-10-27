#!/usr/bin/env python3
"""
=============================================================================
MELILLA - CONTRATOS MENORES DE LA CIUDAD AUTÓNOMA (relaciones trimestrales)
=============================================================================
La Ciudad Autónoma de Melilla publica cada trimestre la relación de los
contratos menores de toda la Ciudad (art. 63.4 LCSP), unidad gestora a
unidad gestora, en la pestaña «Documentos» del perfil del contratante de su
Consejería de Hacienda en la Plataforma de Contratación del Sector Público
(PLACSP), epígrafe «Publicidad de la contratación menor»: un XLS por
trimestre desde 2019 y, de 2019 a 2021, también el mismo listado en PDF y
tres relaciones anuales. Esos menores no están en los datos abiertos de la
PLACSP (la Ciudad no los carga contrato a contrato). Este script lista los
documentos del epígrafe, baja los que todavía no tiene, los guarda tal cual
y genera un Parquet con una fila por contrato y relación.

Ejecutar:  python melilla/melilla_menores.py [--salida DIR] [--solo-parquet]
           [--comprobar-todo] [--max-descargas N] [--sin-cotejo]

Salida (por defecto <repo>/melilla_menores/):
    raw/documentos/<id>_<nombre>   cada documento del epígrafe, tal cual (XLS, PDF), con el nombre
                                   con el que lo sirve la PLACSP; <id>: 12 cifras hexadecimales del
                                   sha256 de su DocumentIdParam (el identificador del documento)
    raw/documentos/_historico/     versiones anteriores (solo con --comprobar-todo; nunca se borran)
    raw/_manifiesto.json           un registro por documento: URL, título, fecha de publicación,
                                   posición en el epígrafe, nombre en origen, formato, clase
                                   (trimestral, anual u otro), año y trimestre, sha256, fecha de
                                   descarga, última comprobación y si se sigue publicando
    raw/_cotejo.json               cotejo XLS-PDF de cada periodo publicado en los dos formatos
    raw/descarga_log.txt           resumen de cada ejecución (se añade al final)
    melilla_menores.parquet        registros acumulados de todas las relaciones (texto)
    _historico/                    versiones anteriores del Parquet

Qué documento es cada uno: el periodo se lee del contenido cuando lo trae (la
fila de título de los XLS de 2019-T1 a T3, «de 20190101 a 20190331», y la de
los de 2025-T4 en adelante, «Del 01/10/2025 al 31/12/2025») y, si no, del
título que le da la PLACSP y del nombre del fichero («4º trimestre 2020»,
«CME_2025_2T», «1-10-2025 a 31-12-2025», «2019 completo»). Un título sin año
(«Contratos menores CAM 2º trimestre Excel») toma el del último trimestre
con ese número acabado antes de su fecha de publicación. Lo que no es una
relación (el oficio de 2019) se guarda y no se lee (clase «otro»). Si el
contenido y el título no dicen el mismo periodo, manda el contenido y se
avisa.

Lectura: de cada periodo, el XLS (xlrd); el PDF (pdfplumber: las tablas de
cada página, con su cuadrícula) solo si el periodo no tiene XLS. Cuando un
periodo trae los dos, se cotejan fila a fila por expediente (importe, fecha
de adjudicación, NIF, persona física, unidad y nº de filas); una diferencia
es un aviso, y el Parquet sale del XLS. El resultado se guarda en
raw/_cotejo.json por la pareja de sha256 y no se repite. Las relaciones
anuales (2019, 2020 y 2021) también se leen: rellenan importes que el
trimestre daba a 0 (columna _relacion).

Columnas (todo texto; nulo lo que viene vacío): los dos formatos de la Ciudad
se leen a las mismas columnas.
- expediente ('2654/2019/CME'), objeto, tipo_contrato (tal cual: 'SERVICIOS',
  'Servicios'...), duracion (formato B, texto: '1 MES'; formato A, el plazo
  de ejecución: '12'), presupuesto (B), importe_adjudicacion, impuestos (A:
  la cuota publicada), n_empresas_invitadas (B, «Publicidad Nº Empresas
  Invitadas»), n_licitadores (B, «Número de Licitadores participantes»),
  publicado_tablon (B), fecha_adjudicacion y fecha_informe_necesidad
  (AAAA-MM-DD, sin la hora; una fecha imposible en origen, '26/08/0119', se
  da como viene, '0119-08-26'), forma_tramitacion, legislacion y
  adjudicatario_extranjero (A), codigo_unidad ('07.7') y unidad (la unidad
  gestora: 'Dirección General de Cultura y Festejos'), adjudicatario y
  adjudicatario_nif (en B vienen juntos, «NOMBRE CIF: B12345678»),
  adjudicatario_persona_fisica ('True'/'False') y formato ('A': 2019-T1 a T3,
  con la cuota de impuestos; 'B': desde 2019-T4, con presupuesto, invitadas
  y licitadores). Importes con punto decimal y sin ceros de más ('5000',
  '14196.44'): el XLS los trae como número y el PDF con coma.
- Añadidas (de origen: no cuentan al comparar versiones): _fuente (URL del
  documento), _archivo_origen (ruta en raw/), _documento_id, _titulo (el de la
  PLACSP), _publicado (fecha y hora de subida del fichero, del sello
  'DOCAAAAMMDDhhmmss' de su nombre, o la fecha de publicación), _relacion
  ('trimestral' o 'anual'), _anio, _trimestre, _formato_origen ('xls',
  'xlsx' o 'pdf'), _fila (fila de la hoja, desde 1, o nº de fila del PDF) y
  _fecha_descarga; de comun/historico.py, _primera_descarga,
  _ultima_descarga y _en_ultima_descarga.

Personas físicas, sin nombre ni NIF en el Parquet: un adjudicatario cuyo
identificador no es el NIF de una persona jurídica (DNI o NIE, con o sin un
0 delante; pasaportes y otros, a menudo con 'null' pegado: 'NYRD27253null'),
o cuyo nombre lleva la forma «APELLIDO*APELLIDO,NOMBRE», un DNI o, sin NIF,
la forma «APELLIDOS,NOMBRE» sin ninguna forma jurídica, va con adjudicatario
y adjudicatario_nif nulos y adjudicatario_persona_fisica = 'True'. Ante la
duda se oculta. Los documentos originales de raw/ sí los traen: no se
publican.

Sesgo del superviviente (docs/PRINCIPIOS.md, regla 3; comun/historico.py):
- Capa cruda: cada documento pasa por guardar_version. La PLACSP da a cada
  documento un DocumentIdParam fijo: uno ya descargado no se vuelve a pedir
  (salvo con --comprobar-todo); una relación corregida es un documento nuevo.
- Parquet: se construye con el código actual desde todos los documentos de
  raw/ (comun.salida.construir_parquet). Un documento que el epígrafe deja de
  listar conserva sus filas con _en_ultima_descarga=False; si vuelve a
  listarse, vuelven a True. Solo se retira algo si la lista se ha leído
  entera y trae algún documento.
- El mismo expediente puede salir en dos relaciones (dos trimestres seguidos,
  o el trimestre y la anual): salen las dos filas; quien lea el Parquet elige.

Qué se pide en cada ejecución (con PAUSA segundos entre una petición y la
siguiente; un 403 o un 429 para todo sin reintentar; los fallos de red y los
5xx se reintentan INTENTOS veces):
- el perfil del contratante (1 GET, con dos redirecciones del servidor) y su
  pestaña «Documentos» (1 POST del formulario JSF de la página);
- cada documento que no se tiene (1 GET), como mucho --max-descargas (por
  defecto TOPE_DESCARGAS) por ejecución; el resto, en la siguiente.
La primera descarga son 49 peticiones (51 con las redirecciones); la
semanal, 2 (4) y una más por relación nueva, unas cuatro al año.
Sale con 1 si la lista no se pudo leer o algún documento nuevo no se pudo
bajar o leer (sin tocar lo anterior).

FUENTES
-------
Verificado en vivo el 2026-10-27 (confianza A):
  - Perfil: https://contrataciondelestado.es/wps/poc?uri=deeplink:perfilContratante&idBp=QOY179oppWwQK2TEfXGy%2BA%3D%3D
    («Consejería de Hacienda de la Ciudad Autónoma de Melilla», NIF S7900010E; el
    idBp es el del BuyerProfileURIID de sus anuncios en la PLACSP). Pestaña
    «Documentos»: el botón submit de valor «Documentos» del formulario JSF del
    perfil (campos ocultos javax.faces.ViewState, javax.faces.encodedURL y
    <form>_SUBMIT; la sesión va en cookies).
  - Epígrafe «Publicidad de la contratación menor» (`tituloEpigrafe<N>` y sus
    filas `epigrafe_<N>_<i>`: título, enlace GetDocumentByIdServlet con
    DocumentIdParam, fecha de publicación «18 Jun 2019» y descripción): 47
    documentos de 2019-T1 a 2026-T2: 30 relaciones trimestrales en XLS, 3
    anuales en XLS (2019, 2020 y 2021), 13 relaciones en PDF (2019-T1 a
    2021-T3 y las anuales de 2019 y 2020: el mismo listado que su XLS) y el
    oficio de 2019. Desde 2021-T4, solo XLS. Ningún trimestre sin relación.
  - XLS (BIFF, una hoja 'Report'): una o dos filas de título y la cabecera.
    Formato A (2019-T1 a T3, 14 columnas): EXPEDIENTE, FECHA DE ADJUDICACIÓN,
    OBJETO, IMPORTE DE ADJUDICACIÓN, PLAZO DE EJECUCIÓN, IMPUESTOS, TIPO, FORMA
    DE TRAMITACIÓN, LEGISLACIÓN APLICABLE, CIF ADJUDICATARIO, NOMBRE
    ADJUDICATARIO, ¿ADJUDICATARIO EXTRANJERO? (SI/NO), CODIGO UNIDAD,
    DESCRIPCION UNIDAD. Formato B (desde 2019-T4, 14 columnas): Expediente,
    Objeto, Tipo, Duración, Presupuesto, Importe de Adjudicación, Publicidad Nº
    Empresas Invitadas, Número de Licitadores participantes, Publicado en Tablón
    Electrónico, Adjudicatario («NOMBRE CIF: X»), Fecha de Adjudicación (con la
    hora), Fecha Informe de Necesidad, Codigo Unidad, Nombre de la unidad
    responsable. Importes como número; fechas como texto o fecha de Excel.
  - PDF: las mismas 14 columnas en una tabla con cuadrícula (páginas de
    8.504 puntos de alto; la cabecera solo en la primera); importes con coma
    decimal. 18.471 filas en las 30 relaciones trimestrales (2.241 a 2.876 al
    año), 188 códigos de unidad.
  - La cabecera no dice si el importe lleva impuesto (Melilla no tiene IVA,
    tiene IPSI): el escalón en el umbral del menor dice que sí lo lleva (de los
    servicios y suministros de más de 15.000 €, 635 de 662 están por debajo de
    16.500 €).
  - robots.txt de contrataciondelestado.es: no prohíbe /wps/ ni /FileSystem/.
=============================================================================
"""

import argparse
import hashlib
import html as html_mod
import json
import os
import re
import sys
import time
import unicodedata
from datetime import date, datetime, timezone
from pathlib import Path
from urllib.parse import parse_qs, urljoin, urlparse

import pandas as pd
import requests

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from comun.historico import guardar_version  # noqa: E402
from comun.lectura import formato_contenido  # noqa: E402
from comun.red import CABECERAS, ERRORES_RED  # noqa: E402
from comun.salida import fecha_version, iso  # noqa: E402
import comun.salida  # noqa: E402

# ============================================================================
# CONFIGURACIÓN
# ============================================================================

URL_PERFIL = ("https://contrataciondelestado.es/wps/poc?uri=deeplink:perfilContratante"
              "&idBp=QOY179oppWwQK2TEfXGy%2BA%3D%3D")
EPIGRAFE = "publicidad de la contratacion menor"   # título del epígrafe, normalizado (normalizar)
BOTON_DOCUMENTOS = "Documentos"                    # valor del submit de la pestaña (en español)
BOTON_DOCUMENTOS_NOMBRE = ":linkPrepDocs"          # final del name de ese submit (el mismo en cualquier idioma)
# El portal WPS localiza los rótulos por la cabecera Accept-Language: se piden en español
CABECERAS_PORTAL = dict(CABECERAS, **{"Accept-Language": "es-ES,es;q=0.9"})
PAUSA = 5.0               # segundos entre peticiones (como mínimo)
INTENTOS = 3              # red y 5xx; un 403 o un 429 no se reintenta
ESPERAS = (15.0, 60.0)    # antes del 2.º y del 3.er intento
TIMEOUT = 120
TOPE_DESCARGAS = 100      # documentos nuevos por ejecución, como mucho
CODIGOS_BLOQUEO = (403, 429)

SALIDA = Path(__file__).resolve().parent.parent / "melilla_menores"
TITULO = "MELILLA - CONTRATOS MENORES DE LA CIUDAD AUTÓNOMA (PLACSP, DOCUMENTOS)"
PARQUET = "melilla_menores.parquet"

COLUMNAS = ("expediente", "objeto", "tipo_contrato", "duracion", "presupuesto", "importe_adjudicacion", "impuestos",
            "n_empresas_invitadas", "n_licitadores", "publicado_tablon", "fecha_adjudicacion",
            "fecha_informe_necesidad", "forma_tramitacion", "legislacion", "adjudicatario_extranjero",
            "codigo_unidad", "unidad", "adjudicatario", "adjudicatario_nif", "adjudicatario_persona_fisica",
            "formato")
METADATOS_ORIGEN = ("_fuente", "_archivo_origen", "_documento_id", "_titulo", "_publicado", "_relacion", "_anio",
                    "_trimestre", "_formato_origen", "_fila", "_fecha_descarga")
ORDEN_METADATOS = METADATOS_ORIGEN + ("_primera_descarga", "_ultima_descarga", "_en_ultima_descarga")

# Cabecera normalizada (normalizar) -> columna. Las que acaban en ':' casan por prefijo
CABECERA = {
    "expediente": "expediente",
    "objeto": "objeto",
    "tipo": "tipo_contrato",
    "duracion": "duracion",
    "plazo de ejecucion": "duracion",
    "presupuesto": "presupuesto",
    "importe de adjudicacion": "importe_adjudicacion",
    "impuestos": "impuestos",
    "publicidad no empresas invitadas": "n_empresas_invitadas",
    "numero de licitadores participantes": "n_licitadores",
    "publicado en tablon electronico": "publicado_tablon",
    "fecha de adjudicacion": "fecha_adjudicacion",
    "fecha informe de necesidad": "fecha_informe_necesidad",
    "forma de tramitacion": "forma_tramitacion",
    "legislacion aplicable": "legislacion",
    "adjudicatario extranjero si/no": "adjudicatario_extranjero",
    "cif adjudicatario": "_nif_publicado",
    "nombre adjudicatario": "_nombre_publicado",
    "adjudicatario": "_adjudicatario_publicado",
    "codigo unidad": "codigo_unidad",
    "descripcion unidad": "unidad",
    "nombre de la unidad responsable": "unidad",
}
OBLIGATORIAS = ("expediente", "importe_adjudicacion", "fecha_adjudicacion", "codigo_unidad")
NUMERICAS = ("presupuesto", "importe_adjudicacion", "impuestos")
ENTERAS = ("n_empresas_invitadas", "n_licitadores")
FECHAS = ("fecha_adjudicacion", "fecha_informe_necesidad")

# NIF de persona jurídica (letra de entidad, 7 cifras y control). Cualquier otro identificador es de una persona
NIF_PJ = re.compile(r"^[ABCDEFGHJNPQRSUVW]\d{7}[0-9A-J]$")
DNI_EN_TEXTO = re.compile(r"\b\d{8}[A-Z]\b")
NOMBRE_ASTERISCO = re.compile(r"^[^,*]+\*[^,*]+,\s*\S")          # APELLIDO*APELLIDO,NOMBRE
NOMBRE_COMA = re.compile(r"^[^,\d]+,\s*[^,\d]+$")                # APELLIDOS,NOMBRE
FORMA_JURIDICA = re.compile(
    r"\b(S\s?L|S\s?A|S\s?L\s?U|S\s?A\s?U|S\s?L\s?L|S\s?C|C\s?B|S\s?COOP|COOP|COOPERATIVA|SOCIEDAD|ASOCIACION|"
    r"FUNDACION|FEDERACION|CLUB|LTD|LIMITED|GMBH|SARL|SAS|INC|AIE|UTE|A/S|AYUNTAMIENTO|UNIVERSIDAD|CONSORCIO|"
    r"INSTITUTO|COLEGIO|COMUNIDAD|HERMANOS|HNOS|EDITORIAL|GRUPO|EMPRESA)\b")
MESES = {"jan": 1, "ene": 1, "feb": 2, "mar": 3, "apr": 4, "abr": 4, "may": 5, "jun": 6, "jul": 7, "aug": 8,
         "ago": 8, "sep": 9, "set": 9, "oct": 10, "nov": 11, "dec": 12, "dic": 12}


def ahora():
    return datetime.now(timezone.utc)


def normalizar(texto):
    """Minúsculas, sin tildes, sin '¿?()' y con los espacios (y saltos de línea) simples."""
    s = unicodedata.normalize("NFKD", str(texto or ""))
    s = "".join(c for c in s if not unicodedata.combining(c)).lower()
    s = s.replace("º", "o").replace("ª", "a").replace("_", " ")
    s = re.sub(r"[¿?()]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


# ============================================================================
# RED: una petición cada vez, con pausa; un 403 o un 429 lo para todo
# ============================================================================

class Bloqueo(Exception):
    """El servidor respondió 403 o 429: no se pide nada más en esta ejecución."""


class ErrorPeticion(Exception):
    """Fallo de red, 5xx tras los reintentos u otro 4xx."""


class Cliente:
    """Sesión HTTP (cookies del portal) con PAUSA entre peticiones y el registro de cada una."""

    def __init__(self, sesion=None, pausa=PAUSA, dormir=time.sleep, reloj=time.monotonic):
        self.sesion = sesion if sesion is not None else requests.Session()
        self.sesion.headers.update(CABECERAS_PORTAL)
        self.pausa = pausa
        self.dormir = dormir
        self.reloj = reloj
        self.ultima = None
        self.peticiones = 0       # peticiones HTTP, contando cada salto de redirección
        self.registro = []

    def _esperar(self, segundos):
        if self.ultima is not None:
            falta = self.ultima + segundos - self.reloj()
            if falta > 0:
                self.dormir(falta)

    def pedir(self, metodo, url, datos=None):
        """Respuesta de un GET o POST (datos: el formulario). Bloqueo con 403/429; ErrorPeticion si falla."""
        detalle = ""
        for intento in range(1, INTENTOS + 1):
            self._esperar(self.pausa if intento == 1 else max(self.pausa, ESPERAS[intento - 2]))
            t0 = self.reloj()
            try:
                if metodo == "POST":
                    r = self.sesion.post(url, data=datos, timeout=TIMEOUT)
                else:
                    r = self.sesion.get(url, timeout=TIMEOUT)
            except ERRORES_RED as e:
                self.ultima = self.reloj()
                self.peticiones += 1
                detalle = f"{type(e).__name__}: {str(e)[:150]}"
                self.registro.append({"metodo": metodo, "url": url, "codigo": None, "error": detalle})
                continue
            except requests.exceptions.RequestException as e:
                self.ultima = self.reloj()
                self.peticiones += 1
                raise ErrorPeticion(f"{type(e).__name__}: {str(e)[:150]} ({url})") from e
            self.ultima = self.reloj()
            saltos = len(getattr(r, "history", None) or [])
            self.peticiones += 1 + saltos
            self.registro.append({"metodo": metodo, "url": url, "codigo": r.status_code, "saltos": saltos,
                                  "bytes": len(r.content or b""), "segundos": round(self.ultima - t0, 1)})
            if r.status_code in CODIGOS_BLOQUEO:
                raise Bloqueo(f"HTTP {r.status_code} en {url}")
            if r.status_code >= 500:
                detalle = f"HTTP {r.status_code}"
                continue
            if r.status_code >= 400:
                raise ErrorPeticion(f"HTTP {r.status_code} en {url}")
            return r
        raise ErrorPeticion(f"{detalle} (tras {INTENTOS} intentos) en {url}")


# ============================================================================
# PÁGINAS DE LA PLACSP: formulario del perfil y lista del epígrafe
# ============================================================================

def formulario_documentos(pagina, base):
    """(acción, campos) del POST que abre la pestaña «Documentos»: los campos ocultos del formulario que
    contiene el submit de valor «Documentos» (o cuyo name acaba en ':linkPrepDocs'), más ese submit. ValueError
    si la página no lo trae."""
    from bs4 import BeautifulSoup

    sopa = BeautifulSoup(pagina, "html.parser")
    botones = [b for b in sopa.find_all("input", attrs={"type": "submit"}) if b.get("name")]
    boton = next((b for b in botones if (b.get("value") or "").strip() == BOTON_DOCUMENTOS),
                 next((b for b in botones if b["name"].endswith(BOTON_DOCUMENTOS_NOMBRE)), None))
    if boton is None:
        raise ValueError("el perfil no trae el botón «Documentos» (¿otra página o un bloqueo?)")
    form = boton.find_parent("form")
    if form is None or not form.get("action"):
        raise ValueError("el botón «Documentos» no está en un formulario con acción")
    campos = [(i["name"], i.get("value") or "") for i in form.find_all("input", attrs={"type": "hidden"})
              if i.get("name")]
    campos.append((boton["name"], boton.get("value") or ""))
    return urljoin(base, html_mod.unescape(form["action"])), campos


def fecha_lista(texto):
    """'18 Jun 2019' (o 'Ene', 'Abr', 'Ago', 'Dic') -> '2019-06-18'; None si no se entiende."""
    m = re.match(r"^\s*(\d{1,2})\s+([A-Za-z]{3})[a-z]*\.?\s+(\d{4})\s*$", texto or "")
    if not m or m.group(2).lower() not in MESES:
        return None
    try:
        return date(int(m.group(3)), MESES[m.group(2).lower()], int(m.group(1))).isoformat()
    except ValueError:
        return None


def id_documento(url):
    """El DocumentIdParam del enlace (o el enlace entero si no lo trae)."""
    valores = parse_qs(urlparse(url).query).get("DocumentIdParam")
    return valores[0] if valores else url


def id_corto(doc_id):
    return hashlib.sha256(doc_id.encode("utf-8")).hexdigest()[:12]


def leer_lista(pagina):
    """Documentos del epígrafe «Publicidad de la contratación menor» de la pestaña Documentos:
    [{posicion, titulo, fecha_publicacion, descripcion, url, doc_id}]. ValueError si no está el epígrafe."""
    from bs4 import BeautifulSoup

    sopa = BeautifulSoup(pagina, "html.parser")
    titulo = next((a for a in sopa.find_all("a", id=re.compile(r"^tituloEpigrafe\d+$"))
                   if normalizar(a.get_text()) == EPIGRAFE), None)
    if titulo is None:
        raise ValueError("la pestaña Documentos no trae el epígrafe «Publicidad de la contratación menor»")
    n = titulo["id"][len("tituloEpigrafe"):]
    bloque = sopa.find("div", id=f"epigrafe_{n}")
    if bloque is None:
        raise ValueError(f"el epígrafe {n} no trae su bloque de documentos")
    documentos = []
    for fila in bloque.find_all("div", id=re.compile(rf"^epigrafe_{n}_\d+$")):
        enlace = fila.find("a", href=re.compile(r"GetDocumentByIdServlet"))
        if enlace is None:
            continue
        celdas = [c for c in fila.find_all("div", recursive=False)]
        textos = [" ".join(c.get_text(" ").split()) for c in celdas]
        url = html_mod.unescape(enlace["href"])
        descripcion = textos[3] if len(textos) > 3 else ""
        documentos.append({
            "posicion": int(fila["id"].rsplit("_", 1)[1]),
            "titulo": textos[0] if textos else "",
            "fecha_publicacion": fecha_lista(textos[2]) if len(textos) > 2 else None,
            "descripcion": None if descripcion in ("", "vacio") else descripcion,
            "url": url,
            "doc_id": id_documento(url),
        })
    return documentos


def nombre_en_origen(respuesta):
    """Nombre del fichero que da el servidor (Content-Disposition), o None."""
    cd = respuesta.headers.get("Content-Disposition") or ""
    m = re.search(r"filename\*=UTF-8''([^;]+)", cd, re.I) or re.search(r'filename="?([^";]+)"?', cd, re.I)
    if not m:
        return None
    nombre = m.group(1).strip()
    try:   # el servidor manda el nombre en Latin-1 sin declararlo
        nombre = nombre.encode("latin-1").decode("utf-8")
    except (UnicodeEncodeError, UnicodeDecodeError):
        pass
    return nombre


def nombre_seguro(nombre, maximo=90):
    """Nombre de fichero sin rutas ni caracteres raros (las tildes se quitan)."""
    s = unicodedata.normalize("NFKD", Path(nombre or "documento").name)
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r"[^A-Za-z0-9._-]+", "_", s).strip("._") or "documento"
    raiz, ext = os.path.splitext(s)
    return raiz[:maximo] + ext[:8]


def formato_fichero(cabeza):
    """'pdf', 'xls', 'xlsx', 'html' u otro (comun.lectura.formato_contenido)."""
    if cabeza.lstrip()[:5] == b"%PDF-":
        return "pdf"
    return formato_contenido(cabeza)


# ============================================================================
# QUÉ PERIODO ES CADA DOCUMENTO
# ============================================================================

def periodo_de_fechas(d1, d2):
    """('trimestral', año, trimestre), ('anual', año, None) o None para un intervalo de fechas."""
    if d1.year != d2.year or d1.day != 1 or d2 < d1:
        return None
    if (d1.month, d2.month, d2.day) == (1, 12, 31):
        return ("anual", d1.year, None)
    trimestre = (d1.month - 1) // 3 + 1
    fin_mes = trimestre * 3
    fin_dia = 31 if fin_mes in (3, 12) else 30
    if d1.month == fin_mes - 2 and (d2.month, d2.day) == (fin_mes, fin_dia):
        return ("trimestral", d1.year, trimestre)
    return None


def periodo_en_texto(texto):
    """Periodo de un intervalo escrito en el texto: '20190101 a 20190331', '01/10/2025 al 31/12/2025',
    '1-10-2025 a 31-12-2025'. None si no hay ninguno."""
    s = normalizar(texto)
    m = re.search(r"\b(\d{4})(\d{2})(\d{2}) a (\d{4})(\d{2})(\d{2})\b", s)
    try:
        if m:
            g = [int(x) for x in m.groups()]
            return periodo_de_fechas(date(g[0], g[1], g[2]), date(g[3], g[4], g[5]))
        m = re.search(r"\b(\d{1,2})[/-](\d{1,2})[/-](\d{4}) (?:a|al|-) (\d{1,2})[/-](\d{1,2})[/-](\d{4})\b", s)
        if m:
            g = [int(x) for x in m.groups()]
            return periodo_de_fechas(date(g[2], g[1], g[0]), date(g[5], g[4], g[3]))
    except ValueError:
        return None
    return None


def anio_por_publicacion(trimestre, publicado):
    """El año del último trimestre `trimestre` acabado antes de la fecha de publicación (AAAA-MM-DD)."""
    if not publicado:
        return None
    p = date.fromisoformat(publicado[:10])
    fin_mes = trimestre * 3
    for anio in (p.year, p.year - 1):
        if date(anio, fin_mes, 28) < p:
            return anio
    return p.year - 2


def periodo_por_titulo(titulo, nombre, publicado):
    """Periodo por el título de la PLACSP y el nombre del fichero: (clase, año, trimestre) o None."""
    nombre = re.sub(r"^DOC\d{14}", "", nombre or "")
    for texto in (titulo, nombre):
        p = periodo_en_texto(texto)
        if p:
            return p
    s = normalizar(f"{titulo} | {nombre}")
    anios = [int(a) for a in re.findall(r"\b(20[1-3]\d)\b", s)]
    m = (re.search(r"\b([1-4]) ?o? ?trimestre\b", s) or re.search(r"\b([1-4])t\b", s))
    if m:
        trimestre = int(m.group(1))
        anio = anios[0] if anios else anio_por_publicacion(trimestre, publicado)
        return ("trimestral", anio, trimestre) if anio else None
    if anios and re.search(r"\b(completo|ano|anual|listado)\b", s):
        return ("anual", anios[0], None)
    return None


def texto_de_cabeza(ruta, formato):
    """Texto de lo que va encima de la cabecera (XLS) o de la primera página (PDF), para buscar el periodo."""
    if formato == "xls":
        import xlrd
        libro = xlrd.open_workbook(str(ruta), on_demand=True)
        try:
            hoja = libro.sheet_by_index(0)
            filas = []
            for i in range(min(hoja.nrows, 10)):
                valores = [str(v) for v in hoja.row_values(i) if str(v).strip()]
                if valores and normalizar(valores[0]) == "expediente":
                    break
                filas.append(" ".join(valores))
            return " | ".join(filas)
        finally:
            libro.release_resources()
    if formato == "xlsx":
        import openpyxl
        with open(ruta, "rb") as fichero:
            libro = openpyxl.load_workbook(fichero, read_only=True, data_only=True)
            try:
                filas = []
                for fila in libro.worksheets[0].iter_rows(max_row=10, values_only=True):
                    valores = [str(v) for v in fila if v is not None and str(v).strip()]
                    if valores and normalizar(valores[0]) == "expediente":
                        break
                    filas.append(" ".join(valores))
                return " | ".join(filas)
            finally:
                libro.close()
    if formato == "pdf":
        import pdfplumber
        with pdfplumber.open(str(ruta)) as pdf:
            if not pdf.pages:
                return ""
            pagina = pdf.pages[0]
            arriba = pagina.crop((0, 0, pagina.width, min(pagina.height, 120)))
            return arriba.extract_text() or ""
    return ""


def clasificar(entrada, ruta, avisos):
    """(clase, año, trimestre) de un documento: el periodo de su contenido y, si no lo trae, el de su título
    y su nombre. Clase 'otro' si no es una relación (no se lee)."""
    formato = entrada.get("formato")
    por_titulo = periodo_por_titulo(entrada.get("titulo"), entrada.get("nombre_en_origen"),
                                    entrada.get("fecha_publicacion"))
    por_contenido = None
    if formato in ("xls", "xlsx", "pdf"):
        try:
            por_contenido = periodo_en_texto(texto_de_cabeza(ruta, formato))
        except Exception as e:   # un fichero que no se puede abrir lo dirá la lectura
            avisos.append(f"{entrada.get('fichero')}: no se pudo leer su título ({type(e).__name__}: {e})")
    if por_contenido and por_titulo and por_contenido != por_titulo:
        avisos.append(f"{entrada.get('fichero')}: el contenido dice {por_contenido} y el título "
                      f"«{entrada.get('titulo')}» dice {por_titulo}; manda el contenido")
    periodo = por_contenido or por_titulo
    if periodo is None or formato not in ("xls", "xlsx", "pdf"):
        return ("otro", None, None)
    return periodo


# ============================================================================
# LECTURA DE UNA RELACIÓN (XLS o PDF) A LAS COLUMNAS COMUNES
# ============================================================================

def numero(valor):
    """Número de una celda: float del XLS o texto '14196,44', '5.000,50', '5000'. None si no lo es."""
    if valor is None:
        return None
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        return float(valor)
    s = re.sub(r"[\s€]", "", str(valor))
    if not s:
        return None
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):    # '5.000': punto de miles (el PDF escribe en español)
        s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        return None


def numero_texto(x):
    """'5000', '14196.44', '0.5' (dos decimales como mucho, sin ceros de más)."""
    s = f"{round(x, 2):.2f}".rstrip("0").rstrip(".")
    return "0" if s in ("-0", "") else s


def fecha_texto(valor, modo_fecha=None):
    """AAAA-MM-DD de una fecha de Excel (número de serie con modo_fecha), de un texto ISO con hora o de
    dd/mm/aaaa (sin validar el año: '26/08/0119' da '0119-08-26'). En el PDF, el final de un nombre largo de
    la columna de al lado se mete a veces en la celda de la fecha ('RO 2021-04-26 14:20:26.0'): la fecha se
    busca dentro del texto. El texto tal cual si no trae ninguna."""
    if valor is None:
        return None
    if isinstance(valor, float) and modo_fecha is not None:
        import xlrd
        return xlrd.xldate_as_datetime(valor, modo_fecha).date().isoformat()
    s = " ".join(str(valor).split())
    if not s:
        return None
    m = re.search(r"(?<!\d)(\d{4})-(\d{2})-(\d{2})(?!\d)", s)
    if m:
        return "-".join(m.groups())
    m = re.search(r"(?<!\d)(\d{1,2})/(\d{1,2})/(\d{4})(?!\d)", s)
    if m:
        return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
    return s


def texto(valor):
    """Texto de una celda (int si el número es entero), con los espacios y saltos de línea simples."""
    if valor is None:
        return None
    if isinstance(valor, float):
        valor = str(int(valor)) if valor.is_integer() else repr(valor)
    s = " ".join(str(valor).split())
    return s or None


def nif_normalizado(valor):
    """NIF sin espacios, puntos ni guiones, en mayúsculas; sin el 'null' que a veces lleva pegado. None si
    no hay ('', '0', 'null')."""
    s = re.sub(r"[\s.\-]", "", str(valor or "")).upper()
    if s in ("", "0", "NULL", "NULLNULL"):
        return None
    if s.endswith("NULL") and len(s) > 4:
        s = s[:-4]
    return s or None


def es_persona_fisica(nombre, nif):
    """¿El adjudicatario es una persona física? (ante la duda, sí: ver el docstring del módulo)"""
    if nif and not NIF_PJ.match(nif):
        return True
    nom = normalizar(nombre).upper() if nombre else ""
    if not nom:
        return False
    if NOMBRE_ASTERISCO.match(nom) or DNI_EN_TEXTO.search(nom):
        return True
    sin_puntos = re.sub(r"[.,]", " ", nom)
    return not nif and bool(NOMBRE_COMA.match(nom)) and not FORMA_JURIDICA.search(sin_puntos)


def adjudicatario_b(valor):
    """Formato B: 'NOMBRE CIF: B12345678' -> (nombre, nif). 'null CIF: nullnull' -> (None, None)."""
    s = texto(valor)
    if not s:
        return None, None
    m = re.match(r"^(.*?)\s*CIF:\s*(.*)$", s, re.S)
    nombre, nif = (m.group(1), m.group(2)) if m else (s, "")
    nombre = nombre.strip() or None
    if nombre and nombre.lower() == "null":
        nombre = None
    return nombre, nif_normalizado(nif)


def columnas_de(cabecera, origen):
    """Columna de cada posición de la cabecera. ValueError con una cabecera que no se conoce."""
    salida, desconocidas = [], []
    for celda in cabecera:
        clave = normalizar(celda)
        if not clave:
            salida.append(None)
            continue
        col = CABECERA.get(clave)
        if col is None:
            desconocidas.append(str(celda))
        salida.append(col)
    if desconocidas:
        raise ValueError(f"{origen}: cabecera no reconocida: {desconocidas}")
    faltan = [c for c in OBLIGATORIAS if c not in salida]
    if faltan:
        raise ValueError(f"{origen}: faltan columnas {faltan} en la cabecera")
    return salida


def registro(columnas, celdas, modo_fecha=None):
    """Una fila de la relación (celdas del XLS o del PDF) a las columnas comunes (todo texto)."""
    v = {}
    for col, celda in zip(columnas, celdas):
        if col is not None:
            v[col] = celda
    r = {c: None for c in COLUMNAS}
    for c in ("expediente", "objeto", "tipo_contrato", "duracion", "publicado_tablon", "forma_tramitacion",
              "legislacion", "adjudicatario_extranjero", "codigo_unidad", "unidad"):
        r[c] = texto(v.get(c))
    for c in NUMERICAS:
        x = numero(v.get(c))
        r[c] = numero_texto(x) if x is not None else texto(v.get(c))
    for c in ENTERAS:
        x = numero(v.get(c))
        r[c] = str(int(x)) if x is not None and float(x).is_integer() else texto(v.get(c))
    for c in FECHAS:
        r[c] = fecha_texto(v.get(c), modo_fecha if isinstance(v.get(c), float) else None)
    if "_adjudicatario_publicado" in v:
        nombre, nif = adjudicatario_b(v["_adjudicatario_publicado"])
        r["formato"] = "B"
    else:
        nombre, nif = texto(v.get("_nombre_publicado")), nif_normalizado(texto(v.get("_nif_publicado")))
        r["formato"] = "A"
    pf = es_persona_fisica(nombre, nif)
    r["adjudicatario"] = None if pf else nombre
    r["adjudicatario_nif"] = None if pf else nif
    r["adjudicatario_persona_fisica"] = "True" if pf else "False"
    return r


def filas_xls(ruta):
    """(cabecera, [(nº de fila desde 1, celdas)], modo_fecha) de la primera hoja de un XLS: la cabecera es la
    fila que empieza por 'Expediente'. Las fechas de Excel quedan como número (fecha_texto las convierte)."""
    import xlrd

    libro = xlrd.open_workbook(str(ruta), on_demand=True)
    try:
        hoja = libro.sheet_by_index(0)
        inicio = next((i for i in range(min(hoja.nrows, 15)) if normalizar(hoja.cell_value(i, 0)) == "expediente"),
                      None)
        if inicio is None:
            raise ValueError("no hay fila de cabecera (Expediente) en las 15 primeras")
        cabecera = [str(v) for v in hoja.row_values(inicio)]
        filas = []
        for i in range(inicio + 1, hoja.nrows):
            celdas = []
            for c in hoja.row(i):
                if c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    celdas.append(None)
                elif c.ctype == xlrd.XL_CELL_DATE:
                    celdas.append(float(c.value))
                elif c.ctype == xlrd.XL_CELL_NUMBER:
                    celdas.append(float(c.value))
                else:
                    celdas.append(str(c.value))
            filas.append((i + 1, celdas))
        return cabecera, filas, libro.datemode
    finally:
        libro.release_resources()


def filas_xlsx(ruta):
    """Como filas_xls para un XLSX (openpyxl; las fechas como texto ISO). Ninguna relación de 2019-2026 lo
    es: está por si la Ciudad cambia de formato."""
    import openpyxl

    with open(ruta, "rb") as fichero:
        libro = openpyxl.load_workbook(fichero, read_only=True, data_only=True)
        try:
            hoja = libro.worksheets[0]
            hoja.reset_dimensions()
            todas = [list(f) for f in hoja.iter_rows(values_only=True)]
        finally:
            libro.close()
    inicio = next((i for i, f in enumerate(todas[:15]) if f and normalizar(f[0]) == "expediente"), None)
    if inicio is None:
        raise ValueError("no hay fila de cabecera (Expediente) en las 15 primeras")

    def celda(v):
        if v is None or (isinstance(v, str) and not v.strip()):
            return None
        if isinstance(v, datetime):
            return v.isoformat(sep=" ")
        if isinstance(v, date):
            return v.isoformat()
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            return float(v)
        return str(v)

    cabecera = ["" if v is None else str(v) for v in todas[inicio]]
    return cabecera, [(i + 1, [celda(v) for v in f]) for i, f in enumerate(todas) if i > inicio], None


def filas_pdf(ruta):
    """(cabecera, [(nº de fila, celdas)]) de un PDF de la relación: las tablas con cuadrícula de cada página
    (pdfplumber). La cabecera es la fila que empieza por 'Expediente' (solo en la primera página); lo de
    encima (el título) no es una fila."""
    import pdfplumber

    cabecera, filas, n = None, [], 0
    with pdfplumber.open(str(ruta)) as pdf:
        for pagina in pdf.pages:
            for tabla in pagina.extract_tables():
                for fila in tabla:
                    # '(cid:9)': un tabulador del texto original que la fuente del PDF no tiene
                    celdas = [None if c is None or not str(c).replace("(cid:9)", " ").strip()
                              else str(c).replace("(cid:9)", " ") for c in fila]
                    if cabecera is None:
                        if celdas and normalizar(celdas[0]) == "expediente":
                            cabecera = [c or "" for c in celdas]
                        continue
                    if normalizar(celdas[0]) == "expediente":   # cabecera repetida en otra página
                        continue
                    n += 1
                    filas.append((n, celdas))
    if cabecera is None:
        raise ValueError("no hay fila de cabecera (Expediente) en las tablas del PDF")
    return cabecera, filas


def leer_relacion(ruta, formato):
    """(lista de registros con _fila, avisos) de una relación. Se saltan las filas sin ningún valor; una
    fila con datos y sin expediente se conserva (y se avisa)."""
    avisos = []
    if formato == "xls":
        cabecera, filas, modo = filas_xls(ruta)
    elif formato == "xlsx":
        cabecera, filas, modo = filas_xlsx(ruta)
    elif formato == "pdf":
        cabecera, filas = filas_pdf(ruta)
        modo = None
    else:
        raise ValueError(f"formato {formato} no se lee")
    columnas = columnas_de(cabecera, Path(ruta).name)
    registros = []
    for n, celdas in filas:
        if not any(c is not None and str(c).strip() for c in celdas):
            continue
        r = registro(columnas, celdas, modo)
        if not r["expediente"]:
            avisos.append(f"{Path(ruta).name}: la fila {n} trae datos y no expediente; se conserva")
        r["_fila"] = str(n)
        registros.append(r)
    return registros, avisos


def leer_tabla(ruta):
    """(DataFrame de texto, avisos) de una versión de un documento (comun.salida.acumular_fichero)."""
    with open(ruta, "rb") as f:
        formato = formato_fichero(f.read(4096))
    registros, avisos = leer_relacion(ruta, formato)
    columnas = list(COLUMNAS) + ["_fila"]
    return pd.DataFrame([[r.get(c) for c in columnas] for r in registros], columns=columnas, dtype=object), avisos


# ============================================================================
# COTEJO XLS-PDF (los periodos publicados en los dos formatos)
# ============================================================================

def cotejar(xls, pdf):
    """Compara dos lecturas de la misma relación (listas de registros) por expediente (multiconjunto, en el
    orden de las filas). Devuelve los recuentos; ok si coinciden las filas y, en cada fila, importe, fecha de
    adjudicación, NIF, persona física y unidad. El objeto y el nombre del adjudicatario se cuentan aparte, sin
    espacios: el PDF parte las líneas largas y a veces mete el final de un nombre en la celda de al lado. Ese
    desborde puede caer encima de los dígitos de la fecha ('AM2E0D2 1-05-18'): una fecha del PDF que no se
    puede leer se cuenta en fecha_ilegible_pdf, no como distinta."""
    def por_expediente(registros):
        d = {}
        for r in registros:
            d.setdefault(r["expediente"], []).append(r)
        return d

    a, b = por_expediente(xls), por_expediente(pdf)
    campos = ("importe_adjudicacion", "fecha_adjudicacion", "adjudicatario_nif", "adjudicatario_persona_fisica",
              "codigo_unidad")
    distintos = {c: 0 for c in campos}
    distintos["objeto"] = distintos["adjudicatario"] = 0
    comunes = solo_xls = solo_pdf = ilegibles = 0
    ejemplos = []
    for exp in sorted(set(a) | set(b), key=str):
        la, lb = a.get(exp, []), b.get(exp, [])
        k = min(len(la), len(lb))
        comunes += k
        solo_xls += len(la) - k
        solo_pdf += len(lb) - k
        for ra, rb in zip(la, lb):
            for c in campos:
                if (c == "fecha_adjudicacion" and ra[c] and ra[c] != rb[c]
                        and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", rb[c] or "")):
                    ilegibles += 1
                    continue
                if (ra[c] or None) != (rb[c] or None):
                    distintos[c] += 1
                    if len(ejemplos) < 5:
                        ejemplos.append(f"{exp} {c}")
            for c in ("objeto", "adjudicatario"):
                if re.sub(r"\s+", "", ra[c] or "") != re.sub(r"\s+", "", rb[c] or ""):
                    distintos[c] += 1
    ok = solo_xls == 0 and solo_pdf == 0 and all(distintos[c] == 0 for c in campos)
    return {"filas_xls": len(xls), "filas_pdf": len(pdf), "filas_comunes": comunes, "solo_xls": solo_xls,
            "solo_pdf": solo_pdf, "distintos": distintos, "fecha_ilegible_pdf": ilegibles, "ejemplos": ejemplos,
            "ok": ok}


# ============================================================================
# MANIFIESTO Y RESUMEN
# ============================================================================

class Manifiesto:
    """raw/_manifiesto.json: un registro por documento descargado, con la ruta en raw/ como clave."""

    def __init__(self, raw):
        self.raw = Path(raw)
        self.ruta = self.raw / "_manifiesto.json"
        try:
            self.datos = json.loads(self.ruta.read_text(encoding="utf-8"))
        except (FileNotFoundError, ValueError):
            self.datos = {}

    def rel(self, ruta):
        return Path(ruta).relative_to(self.raw).as_posix()

    def get(self, rel):
        return self.datos.get(rel, {})

    def por_doc_id(self):
        return {e.get("doc_id"): rel for rel, e in self.datos.items() if e.get("doc_id")}

    def guardar(self):
        self.raw.mkdir(parents=True, exist_ok=True)
        tmp = self.ruta.with_name(f".{self.ruta.name}.tmp")
        tmp.write_text(json.dumps(self.datos, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8")
        os.replace(tmp, self.ruta)


class Resumen:
    """Lo descargado, lo retirado, la clasificación, el cotejo, los avisos y lo que falló."""

    def __init__(self, titulo):
        self.titulo = titulo
        self.inicio = ahora()
        self.descargados = []
        self.ya_descargados = 0
        self.retirados = []
        self.periodos = []
        self.cotejos = []
        self.avisos = []
        self.fallidos = []
        self.parquets = []
        self.peticiones = None

    def texto(self):
        lineas = ["=" * 70]
        if self.fallidos:
            lineas.append(f"⚠️ {self.titulo}: COMPLETADO CON ERRORES ({len(self.fallidos)})")
        else:
            lineas.append(f"✅ {self.titulo}: COMPLETADO")
        lineas += ["=" * 70, f"Inicio: {iso(self.inicio)}  Fin: {iso(ahora())}"]
        if self.peticiones is not None:
            lineas.append(f"Peticiones HTTP: {self.peticiones}")
        if self.ya_descargados:
            lineas.append(f"Documentos ya descargados (no se piden): {self.ya_descargados}")

        def bloque(titulo, elementos):
            if elementos:
                lineas.append(f"\n{titulo} ({len(elementos)}):")
                lineas.extend(f"   - {e}" for e in elementos)

        bloque("DESCARGADOS (versión nueva)", self.descargados)
        bloque("RETIRADOS (se conservan con _en_ultima_descarga=False)", self.retirados)
        bloque("RELACIONES", self.periodos)
        bloque("COTEJO XLS-PDF", self.cotejos)
        bloque("PARQUET", [f"{n}: {f:,} filas x {c} columnas ({r:,} ya no publicadas)"
                           for n, f, c, r in self.parquets])
        bloque("AVISOS", self.avisos)
        bloque("ERRORES - vuelve a ejecutar el script para reintentarlos", self.fallidos)
        return "\n".join(lineas)

    def cerrar(self, raw):
        texto_ = self.texto()
        print("\n" + texto_)
        raw = Path(raw)
        raw.mkdir(parents=True, exist_ok=True)
        with open(raw / "descarga_log.txt", "a", encoding="utf-8") as f:
            f.write(texto_ + "\n\n")
        return 1 if self.fallidos else 0


# ============================================================================
# DESCARGA: lista del epígrafe y documentos nuevos
# ============================================================================

def sello_subida(nombre):
    """Fecha y hora de subida del sello 'DOCAAAAMMDDhhmmss' con que la PLACSP nombra el fichero, o None."""
    m = re.match(r"^DOC(\d{4})(\d{2})(\d{2})(\d{2})(\d{2})(\d{2})", nombre or "")
    if not m:
        return None
    try:
        return datetime(*map(int, m.groups())).isoformat()
    except ValueError:
        return None


def descargar(salida, manifiesto, resumen, cliente, comprobar_todo=False, max_descargas=TOPE_DESCARGAS):
    """Lee la lista del epígrafe, baja los documentos nuevos (o todos, con comprobar_todo) y anota en el
    manifiesto qué sigue publicado. Sin lista no se retira nada."""
    raw = Path(salida) / "raw"
    carpeta = raw / "documentos"
    ultima = None
    try:
        ultima = perfil = cliente.pedir("GET", URL_PERFIL)
        accion, campos = formulario_documentos(perfil.text, perfil.url or URL_PERFIL)
        ultima = pagina = cliente.pedir("POST", accion, campos)
        lista = leer_lista(pagina.text)
    except Bloqueo as e:
        resumen.fallidos.append(f"contrataciondelestado.es: {e}; se para sin pedir nada más")
        return
    except ErrorPeticion as e:
        resumen.fallidos.append(f"contrataciondelestado.es (lista de documentos): {e}")
        return
    except ValueError as e:
        # La página que no se entiende se guarda para verla (raw/_pagina_sin_lista.html, con sus versiones)
        detalle = ""
        if ultima is not None:
            guardar_version(raw / "_pagina_sin_lista.html", ultima.content or b"")
            detalle = f" (HTTP {ultima.status_code}, {ultima.url}; guardada en raw/_pagina_sin_lista.html)"
        resumen.fallidos.append(f"contrataciondelestado.es (lista de documentos): {e}{detalle}")
        return
    if not lista:
        resumen.fallidos.append("contrataciondelestado.es: el epígrafe no trae ningún documento; no se retira nada")
        return
    print(f"  📄 {len(lista)} documentos en el epígrafe")
    momento = iso(ahora())
    conocidos = manifiesto.por_doc_id()
    listados = {doc["doc_id"] for doc in lista}   # la lista entera, aunque un bloqueo pare las descargas
    descargas = 0
    for doc in sorted(lista, key=lambda d: d["posicion"]):
        rel = conocidos.get(doc["doc_id"])
        entrada = manifiesto.datos.get(rel) if rel else None
        datos_lista = {k: doc[k] for k in ("posicion", "titulo", "fecha_publicacion", "descripcion", "url")}
        if entrada is not None and (raw / rel).exists() and not comprobar_todo:
            entrada.update(datos_lista, publicado=True, comprobado=momento)
            entrada.pop("retirado_desde", None)
            resumen.ya_descargados += 1
            continue
        if descargas >= max_descargas:
            resumen.avisos.append(f"«{doc['titulo']}»: pasa a la siguiente ejecución (tope de {max_descargas} "
                                  "descargas)")
            continue
        descargas += 1
        try:
            r = cliente.pedir("GET", doc["url"])
        except Bloqueo as e:
            resumen.fallidos.append(f"contrataciondelestado.es: {e} («{doc['titulo']}»); se para sin pedir nada más")
            break
        except ErrorPeticion as e:
            resumen.fallidos.append(f"contrataciondelestado.es («{doc['titulo']}»): {e}")
            continue
        contenido = r.content or b""
        formato = formato_fichero(contenido[:4096])
        if formato not in ("xls", "xlsx", "pdf"):
            resumen.fallidos.append(f"contrataciondelestado.es («{doc['titulo']}»): la respuesta no es un documento "
                                    f"({formato}, {len(contenido)} bytes); se conserva lo anterior")
            continue
        origen = nombre_en_origen(r)
        if entrada is not None:
            destino = raw / rel
        else:
            destino = carpeta / f"{id_corto(doc['doc_id'])}_{nombre_seguro(origen or doc['titulo'] + '.' + formato)}"
        estado = guardar_version(destino, contenido)
        rel = manifiesto.rel(destino)
        entrada = manifiesto.datos.setdefault(rel, {})
        if estado != "sin_cambios" or not entrada.get("fecha_descarga"):
            entrada["fecha_descarga"] = fecha_version(destino)
            resumen.descargados.append(f"«{doc['titulo']}» -> {rel} ({estado})")
        entrada.update(datos_lista, doc_id=doc["doc_id"], id=id_corto(doc["doc_id"]), fichero=rel,
                       nombre_en_origen=origen, formato=formato, bytes=len(contenido),
                       sha256=hashlib.sha256(contenido).hexdigest(), publicado=True, comprobado=momento,
                       subido=sello_subida(origen))
        entrada.pop("retirado_desde", None)
        entrada.pop("clase", None)   # se vuelve a clasificar
    for rel, entrada in manifiesto.datos.items():
        if entrada.get("doc_id") not in listados and entrada.get("publicado", True):
            entrada["publicado"] = False
            entrada.setdefault("retirado_desde", momento)
            entrada["comprobado"] = momento
            resumen.retirados.append(f"«{entrada.get('titulo')}» ({rel})")
    manifiesto.guardar()


# ============================================================================
# CLASIFICACIÓN, COTEJO Y PARQUET
# ============================================================================

def clasificar_todo(salida, manifiesto, resumen):
    """Clase, año y trimestre de cada documento del manifiesto (solo los que no lo tienen o han cambiado)."""
    raw = Path(salida) / "raw"
    for rel, entrada in sorted(manifiesto.datos.items()):
        ruta = raw / rel
        if not ruta.exists():
            continue
        if entrada.get("clase") and entrada.get("clase_sha256") == entrada.get("sha256"):
            continue
        clase, anio, trimestre = clasificar(entrada, ruta, resumen.avisos)
        entrada.update(clase=clase, anio=anio, trimestre=trimestre, clase_sha256=entrada.get("sha256"))
    manifiesto.guardar()
    trimestres = sorted({(e["anio"], e["trimestre"]) for e in manifiesto.datos.values()
                         if e.get("clase") == "trimestral" and e.get("anio")})
    anuales = sorted({e["anio"] for e in manifiesto.datos.values() if e.get("clase") == "anual" and e.get("anio")})
    otros = [e.get("titulo") for e in manifiesto.datos.values() if e.get("clase") == "otro"]
    if trimestres:
        (a0, t0), (a1, t1) = trimestres[0], trimestres[-1]
        resumen.periodos.append(f"trimestrales: {len(trimestres)} ({a0}-T{t0} a {a1}-T{t1})")
        todos = [(a, t) for a in range(a0, a1 + 1) for t in range(1, 5) if (a0, t0) <= (a, t) <= (a1, t1)]
        huecos = [f"{a}-T{t}" for a, t in todos if (a, t) not in trimestres]
        if huecos:
            resumen.avisos.append(f"trimestres sin relación entre el primero y el último: {', '.join(huecos)}")
    if anuales:
        resumen.periodos.append(f"anuales: {', '.join(map(str, anuales))}")
    if otros:
        resumen.periodos.append(f"no son relaciones (no se leen): {len(otros)} ({'; '.join(map(str, otros))[:200]})")


def _clave_periodo(e):
    return (e.get("clase"), e.get("anio"), e.get("trimestre"))


def para_el_parquet(manifiesto):
    """Rutas de raw/ que se leen: de cada periodo, sus XLS; sus PDF solo si no tiene ningún XLS."""
    por_periodo = {}
    for rel, e in manifiesto.datos.items():
        if e.get("clase") in ("trimestral", "anual"):
            por_periodo.setdefault(_clave_periodo(e), []).append(rel)
    elegidos = []
    for rels in por_periodo.values():
        xls = [r for r in rels if manifiesto.get(r).get("formato") in ("xls", "xlsx")]
        elegidos += xls or [r for r in rels if manifiesto.get(r).get("formato") == "pdf"]
    return sorted(elegidos)


def cotejar_todo(salida, manifiesto, resumen):
    """Cotejo XLS-PDF de los periodos que traen los dos (el XLS de fecha de subida más cercana a la del
    PDF). Se guarda en raw/_cotejo.json por la pareja de sha256 y no se repite."""
    raw = Path(salida) / "raw"
    ruta_cotejo = raw / "_cotejo.json"
    try:
        hechos = json.loads(ruta_cotejo.read_text(encoding="utf-8"))
    except (FileNotFoundError, ValueError):
        hechos = {}
    por_periodo = {}
    for rel, e in manifiesto.datos.items():
        if e.get("clase") in ("trimestral", "anual") and (raw / rel).exists():
            por_periodo.setdefault(_clave_periodo(e), []).append(rel)
    cambiado = False
    for clave, rels in sorted(por_periodo.items(), key=lambda kv: str(kv[0])):
        pdfs = [r for r in rels if manifiesto.get(r).get("formato") == "pdf"]
        xlss = [r for r in rels if manifiesto.get(r).get("formato") in ("xls", "xlsx")]
        for p in pdfs:
            if not xlss:
                continue
            fp = manifiesto.get(p).get("subido") or manifiesto.get(p).get("fecha_publicacion") or ""
            x = min(xlss, key=lambda r: abs(_segundos(manifiesto.get(r).get("subido")
                                                    or manifiesto.get(r).get("fecha_publicacion")) - _segundos(fp)))
            k = f"{manifiesto.get(x).get('sha256')}|{manifiesto.get(p).get('sha256')}"
            if k not in hechos:
                try:
                    rx, _ = leer_relacion(raw / x, manifiesto.get(x).get("formato"))
                    rp, _ = leer_relacion(raw / p, "pdf")
                except Exception as e:
                    resumen.fallidos.append(f"cotejo {x} / {p}: {type(e).__name__}: {e}")
                    continue
                hechos[k] = dict(cotejar(rx, rp), xls=x, pdf=p, periodo=_texto_periodo(clave), fecha=iso(ahora()))
                cambiado = True
            c = hechos[k]
            dist = ", ".join(f"{n} {v}" for n, v in c["distintos"].items() if v)
            ilegible = f"; fecha ilegible en el PDF: {c['fecha_ilegible_pdf']}" if c.get("fecha_ilegible_pdf") else ""
            resumen.cotejos.append(f"{c['periodo']}: {'igual' if c['ok'] else 'DISTINTO'} ({c['filas_xls']} filas "
                                   f"en el XLS, {c['filas_pdf']} en el PDF; distintos: {dist or 'ninguno'}{ilegible})")
            if not c["ok"]:
                resumen.avisos.append(f"cotejo {c['periodo']}: el XLS y el PDF no dicen lo mismo "
                                      f"({c['solo_xls']} filas solo en el XLS, {c['solo_pdf']} solo en el PDF, "
                                      f"distintos {c['distintos']}, p. ej. {c['ejemplos']}); el Parquet sale del XLS")
    if cambiado:
        raw.mkdir(parents=True, exist_ok=True)
        tmp = ruta_cotejo.with_name(".cotejo.json.tmp")
        tmp.write_text(json.dumps(hechos, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8")
        os.replace(tmp, ruta_cotejo)


def _segundos(texto_fecha):
    try:
        return datetime.fromisoformat(str(texto_fecha)[:19]).timestamp()
    except ValueError:
        return 0.0


def _texto_periodo(clave):
    clase, anio, trimestre = clave
    return f"{anio}-T{trimestre}" if clase == "trimestral" else f"{anio} (anual)"


def ordenar_columnas(df):
    """Columnas comunes (en su orden) y después las añadidas por el script."""
    propias = [c for c in ORDEN_METADATOS if c in df.columns]
    return df[[c for c in df.columns if c not in ORDEN_METADATOS] + propias]


def escribir_parquet(df, destino):
    """Parquet con todas las columnas como texto (comun.salida.escribir_parquet)."""
    return comun.salida.escribir_parquet(df, destino, ordenar_columnas=ordenar_columnas)


def acumular_fichero(actual, rel, anterior, metadatos, manifiesto, resumen):
    """Registros de un documento a lo largo de sus versiones (comun.salida.acumular_fichero)."""
    return comun.salida.acumular_fichero(actual, rel, anterior, metadatos, manifiesto, resumen,
                                         metadatos_origen=METADATOS_ORIGEN, leer_tabla=leer_tabla)


def generar_parquet(salida, manifiesto, resumen):
    raw = Path(salida) / "raw"
    ficheros = []
    for rel in para_el_parquet(manifiesto):
        e = manifiesto.get(rel)
        if not (raw / rel).exists():
            continue
        ficheros.append((raw / rel, rel, {
            "_fuente": e.get("url"), "_archivo_origen": rel, "_documento_id": e.get("id"), "_titulo": e.get("titulo"),
            "_publicado": e.get("subido") or e.get("fecha_publicacion"), "_relacion": e.get("clase"),
            "_anio": str(e.get("anio")), "_trimestre": None if e.get("trimestre") is None else str(e["trimestre"]),
            "_formato_origen": e.get("formato")}))
    destino = Path(salida) / PARQUET
    if ficheros or destino.exists():
        print("\n🧱 Generando el Parquet...")
        comun.salida.construir_parquet(destino, ficheros, raw, manifiesto, resumen,
                                       metadatos_origen=METADATOS_ORIGEN, acumular_fichero=acumular_fichero,
                                       escribir_parquet=escribir_parquet)
    else:
        resumen.fallidos.append("no hay ninguna relación en raw/: no se genera el Parquet")


def main(argv=None, cliente=None):
    p = argparse.ArgumentParser(description=TITULO)
    p.add_argument("--salida", type=Path, default=SALIDA, help=f"carpeta de salida (por defecto {SALIDA})")
    p.add_argument("--solo-parquet", action="store_true", help="no pide nada: rehace el Parquet desde raw/")
    p.add_argument("--comprobar-todo", action="store_true",
                   help="vuelve a pedir también los documentos ya descargados")
    p.add_argument("--max-descargas", type=int, default=TOPE_DESCARGAS,
                   help=f"documentos que se bajan como mucho en esta ejecución (por defecto {TOPE_DESCARGAS})")
    p.add_argument("--sin-cotejo", action="store_true", help="no coteja el XLS con el PDF")
    a = p.parse_args(argv)
    salida = Path(a.salida)
    raw = salida / "raw"
    resumen = Resumen(TITULO)
    manifiesto = Manifiesto(raw)
    print(f"🏛️  {TITULO}\n   salida: {salida}")
    if not a.solo_parquet:
        cliente = cliente or Cliente()
        descargar(salida, manifiesto, resumen, cliente, a.comprobar_todo, a.max_descargas)
        resumen.peticiones = cliente.peticiones
    clasificar_todo(salida, manifiesto, resumen)
    if not a.sin_cotejo:
        cotejar_todo(salida, manifiesto, resumen)
    generar_parquet(salida, manifiesto, resumen)
    return resumen.cerrar(raw)


if __name__ == "__main__":
    sys.exit(main())
